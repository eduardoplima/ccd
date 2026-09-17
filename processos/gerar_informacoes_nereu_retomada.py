"""Informações de retomada da marcha processual (lote Nereu / MS 0807247-93.2025.8.20.0000).

Em 16/07/2025 o Relator (Cons. Renato Costa Dias, GCREN) determinou nos sete processos abaixo o
"sobrestamento do feito até ulterior deliberação, nos termos do art. 36, inciso III, da Lei
Complementar nº 464/2012", por força da liminar do TJRN no MS 0807247-93. O STF suspendeu os
efeitos do acórdão do TJRN na SS 5.749 (decisão de 28/07/2026, trânsito em 27/08/2026), e a CCD
propõe a continuidade da marcha — é esse o despacho de processos/modelos/
modelo_analise_retomada_nereu.docx, escrito à mão para o 102455/2018 (Evento 99).

Conferido no banco e nos PDFs (15/09/2026): os sete têm a mesma marcha (acórdão da Sessão 10V de
06/06/2025, Termo de Ressalva 07/07/2025, sobrestamento 16/07/2025 como último evento), o mesmo
relator, estão na CCD, e o despacho de cada um contém literalmente a frase citada. Só variam o
número do processo e o número do Evento — o gerador reescreve esses dois runs do modelo com
python-docx e mantém todo o resto (o 102455/2018 sai idêntico ao modelo, o que valida o gerador).

Ao cadastrar: o texto cita "documento anexo" (decisão do STF) e "anexa certidão" (trânsito) —
os mesmos anexos do 102455/2018 acompanham cada informação.

Rodar: .venv/Scripts/python.exe processos/gerar_informacoes_nereu_retomada.py
"""
import re
import shutil
import unicodedata
from datetime import datetime
from pathlib import Path

import docx

from ccd.config import informacoes_dir
from ccd.db import run_query_df
from ccd.docs import docx_to_pdf
from ccd.pdf import extract_text_from_pdf

BASE = Path(__file__).parent
MODELO = BASE / "modelos" / "modelo_analise_retomada_nereu.docx"
DESTINO = BASE / "nereu_retomada"

MODELO_PROCESSO = "102455_2018"
MODELO_EVENTO = "99"
RELATOR = "RENATO COSTA DIAS"
TITULO_SOBRESTAMENTO = "DIP - SOBRESTAMENTO DECISÃO VANTAGEM TRANSITÓRIA"
FRASE = (r"sobrestamento do feito ate ulterior deliberacao, nos termos do art\. 36, inciso III, "
         r"da Lei Complementar n.{0,3}\s*464/2012")

PROCESSOS = ["022948/2016", "023030/2016", "101157/2018", "101158/2018", "102455/2018",
             "102521/2018", "022043/2016"]

SQL_EVENTOS = """
    SELECT RTRIM(v.numero_processo)+'/'+RTRIM(v.ano_processo) AS processo,
           RTRIM(v.setor) AS setor, RTRIM(v.Titulo_Modelo_informacao) AS titulo,
           v.ordem, ev.SequencialProcessoEvento AS evento
    FROM vw_ata_informacao v
    JOIN Pro_ProcessoEvento ev ON ev.idInformacao = v.idInformacao
    WHERE RTRIM(v.numero_processo)+'/'+RTRIM(v.ano_processo) IN ({chaves})
      AND v.Inativa IS NULL
"""
SQL_PROCESSOS = """
    SELECT RTRIM(p.numero_processo)+'/'+RTRIM(p.ano_processo) AS processo,
           RTRIM(p.setor_atual) AS setor, RTRIM(r.nome) AS relator
    FROM Processos p
    LEFT JOIN Relator r ON r.codigo = p.codigo_relator
    WHERE RTRIM(p.numero_processo)+'/'+RTRIM(p.ano_processo) IN ({chaves})
"""


def sem_acento(texto: str) -> str:
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", texto)


def carregar() -> list[dict]:
    chaves = {f"p{i}": p for i, p in enumerate(PROCESSOS)}
    marcadores = ",".join(f":{k}" for k in chaves)
    eventos = run_query_df(SQL_EVENTOS.format(chaves=marcadores), **chaves)
    processos = run_query_df(SQL_PROCESSOS.format(chaves=marcadores), **chaves).set_index("processo")

    itens = []
    for processo in PROCESSOS:
        proc = processos.loc[processo]
        assert proc.relator == RELATOR, f"{processo}: relator {proc.relator}"
        assert proc.setor == "CCD", f"{processo}: está em {proc.setor}, não na CCD"

        ev = eventos[eventos.processo == processo]
        desp = ev[(ev.setor == "GCREN") & (ev.titulo == TITULO_SOBRESTAMENTO)]
        assert len(desp) == 1, f"{processo}: {len(desp)} despachos de sobrestamento, esperava 1"
        desp = desp.iloc[0]
        assert desp.evento == ev.evento.max(), (
            f"{processo}: há evento posterior ao sobrestamento ({ev.evento.max()})")

        numero, ano = processo.split("/")
        pdf = informacoes_dir() / "GCREN" / f"GCREN_{numero}_{ano}_{int(desp.ordem):04d}.pdf"
        assert re.search(FRASE, sem_acento(extract_text_from_pdf(str(pdf))), re.I), (
            f"{processo}: o despacho {pdf.name} não contém a frase citada")

        itens.append({"processo": processo, "evento": int(desp.evento)})
    return itens


def substituir_run(doc, contem: str, run_antigo: str, novo: str) -> None:
    """Troca o texto de um run no parágrafo que contém `contem`; o run tem de existir sozinho."""
    par = next(p for p in doc.paragraphs if contem in p.text)
    runs = [r for r in par.runs if r.text == run_antigo]
    assert len(runs) == 1, f'run "{run_antigo}" não é único em "{contem}"'
    runs[0].text = novo


def gerar(item: dict) -> Path:
    numero, ano = item["processo"].split("/")
    pasta = DESTINO / f"{numero}_{ano}"
    pasta.mkdir(parents=True, exist_ok=True)
    out = pasta / f"{numero}_{ano}.docx"

    doc = docx.Document(str(MODELO))
    # cabeçalho: runs ['Processo nº ', '1', '02455', '_', '20', '1', '8', '-TC'] — o número está
    # picado pelo Word; junta-se tudo no primeiro pedaço e esvaziam-se os demais
    cab = next(p for p in doc.paragraphs if p.text.startswith("Processo nº "))
    assert MODELO_PROCESSO in cab.text
    pedacos = [r for r in cab.runs if r.text not in ("Processo nº ", "-TC")]
    assert "".join(r.text for r in pedacos) == MODELO_PROCESSO
    pedacos[0].text = f"{numero}_{ano}"
    for r in pedacos[1:]:
        r.text = ""
    substituir_run(doc, "Evento nº ", MODELO_EVENTO, str(item["evento"]))

    if out.exists():  # preserva a versão anterior (skill edicao-minima)
        bak = out.with_name(f"{out.stem}_{datetime.now():%Y%m%d_%H%M%S}{out.suffix}")
        shutil.copy2(out, bak)
        print(f"  versão anterior preservada em: {bak.name}")
    doc.save(str(out))
    docx_to_pdf(str(out), str(pasta))
    return out


if __name__ == "__main__":
    modelo = [p.text for p in docx.Document(str(MODELO)).paragraphs]
    itens = carregar()
    for item in itens:
        numero, ano = item["processo"].split("/")
        print(f'{item["processo"]} (sobrestamento no Evento {item["evento"]}):')
        out = gerar(item)

        # check: só o cabeçalho e o nº do evento mudam; todo o resto é o modelo, palavra a palavra
        saida = [p.text for p in docx.Document(str(out)).paragraphs]
        assert len(saida) == len(modelo)
        assert saida[0] == f"Processo nº {numero}_{ano}-TC", saida[0]
        esperado = modelo[:]
        esperado[0] = saida[0]
        i = next(i for i, t in enumerate(modelo) if "Evento nº " in t)
        esperado[i] = modelo[i].replace(f"Evento nº {MODELO_EVENTO})", f'Evento nº {item["evento"]})')
        assert saida == esperado, f'{item["processo"]}: texto divergente do modelo'
        assert f'(Evento nº {item["evento"]})' in saida[i]
        assert out.with_suffix(".pdf").is_file(), "PDF não gerado"
        print(f"  salvo: {out.relative_to(BASE)} (+ .pdf)")

    print(f"\n{len(itens)} informações geradas em {DESTINO.relative_to(BASE)}/")
