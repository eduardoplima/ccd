"""Informações de declínio de atribuição à CIP — Processos nº 001323/2024-TC e
001325/2024-TC (cadernos do TAG nº 04/2018, São Bento do Trairí; relator George
Montenegro Soares).

Ambos os cadernos vieram a esta Coordenadoria depois de apresentada a defesa, e as
informações de mérito já produzidas (`001323_2024/gerar_informacao.py` e
`001325_2024/gerar_informacao.py`, mantidas como registro da versão substituída)
analisaram as teses. O exame de processo de controle externo em sede de defesa,
porém, compete à Coordenadoria de Instrução Processual — CIP (art. 31 da Resolução
nº 042/2024-TCE), e não à CCD, cujas atribuições estão no art. 32 do mesmo
Regulamento.

Este script gera a informação curta de declínio, que substituirá a anterior. Não
entra no mérito nem discute se houve análise preliminar anterior pela unidade
técnica — por isso o art. 31 é invocado como um todo (incisos I a III), e não pelo
inciso I, cuja hipótese é condicionada àquela análise.

Rodar: .venv/Scripts/python.exe processos/gerar_informacao_cip.py
"""
import time
from copy import deepcopy
from datetime import datetime
from pathlib import Path

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docxtpl import DocxTemplate

from ccd.config import REPO_ROOT
from ccd.db import run_query_df
from ccd.docs import docx_to_pdf

TEMPLATE = str(REPO_ROOT / "scripts" / "automacao" / "templates" / "modelo_informacao.docx")

MESES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho",
         "agosto", "setembro", "outubro", "novembro", "dezembro"]
hoje = datetime.now()
DATA = f"Natal/RN, {hoje.day} de {MESES[hoje.month - 1]} de {hoje.year}."

RELATOR = "George Montenegro Soares"
COD_RELATOR = 19

CADERNOS = [
    # `responsavel` já vem com a preposição contraída ("diz respeito {responsavel}")
    dict(id_processo=596698, numero="001323", ano=2024,
         responsavel="ao Sr. José Aracleide de Araújo, então Prefeito Municipal de "
                     "São Bento do Trairí",  # o Assunto grafa "Aracieide" (erro de digitação)
         citado="Citado o responsável",
         ev_defesa=35, ev_remessa=39, data_remessa="11 de agosto de 2026"),
    dict(id_processo=596700, numero="001325", ano=2024,
         responsavel="à Sra. Juliana Patrícia de Oliveira Pessoa Dantas, então "
                     "Secretária Municipal de Educação de São Bento do Trairí",
         citado="Citada a responsável",
         ev_defesa=30, ev_remessa=34, data_remessa="7 de agosto de 2026"),
]


def paragrafos(c: dict) -> list[str]:
    return [
        'Trata-se de caderno autônomo aberto por força do Acórdão nº 293/2023-TC, '
        'proferido no Processo nº 006775/2018-TC, que determinou a apuração da '
        'responsabilidade dos signatários do Termo de Ajustamento de Gestão nº '
        '04/2018, celebrado com o Município de São Bento do Trairí. O presente '
        f'caderno diz respeito {c["responsavel"]}.',

        f'{c["citado"]}, sobreveio a defesa do Evento '
        f'{c["ev_defesa"]}, e o Despacho de {c["data_remessa"]} (Evento '
        f'{c["ev_remessa"]}) remeteu os autos a esta Coordenadoria de Controle de '
        'Decisões para análise técnica.',

        'Ocorre que o exame e a instrução dos processos de controle externo em sede '
        'de defesa competem à Coordenadoria de Instrução Processual — CIP, nos '
        'termos do art. 31 da Resolução nº 042/2024-TCE, que estabelece o '
        'regulamento da Secretaria de Controle Externo, seja para instruí-los '
        'diretamente, seja para encaminhá-los à unidade técnica de controle externo '
        'respectiva quando o Relator assim o determinar.',

        'A esta Coordenadoria de Controle de Decisões o art. 32 do mesmo '
        'Regulamento reserva o cadastro e o monitoramento das obrigações de fazer e '
        'de não fazer e das recomendações emitidas pelo Tribunal, as atividades e os '
        'controles inerentes à cobrança executiva e aos pagamentos decorrentes de '
        'deliberações e o apoio às medidas de indisponibilidade de bens, não se '
        'compreendendo entre elas a instrução processual em sede de defesa.',

        'Ante o exposto, retornam-se os autos ao Exmo. Cons. Relator, sugerindo-se o '
        'encaminhamento do processo à Coordenadoria de Instrução Processual, unidade '
        'a que compete o exame da defesa apresentada.',
    ]


def gerar(c: dict) -> Path:
    proc = run_query_df(
        "SELECT numero_processo, ano_processo, codigo_relator "
        "FROM processo.dbo.Processos WHERE IdProcesso = :p", p=c["id_processo"])
    assert len(proc) == 1, f"processo não encontrado: {c['id_processo']}"
    assert proc.numero_processo.iloc[0].strip() == c["numero"]
    assert int(proc.ano_processo.iloc[0]) == c["ano"]
    assert int(proc.codigo_relator.iloc[0]) == COD_RELATOR, "relator mudou — revisar texto"

    corpo = paragrafos(c)
    doc = DocxTemplate(TEMPLATE)
    doc.render({
        "processo": f'{c["numero"]}/{c["ano"]} - TC',
        "assunto": "APURAÇÃO DE RESPONSABILIDADE",
        "relator": RELATOR,
        "parágrafos": corpo,  # sem numerar: o modelo numera
    })

    # pós-processamento via doc.docx (get_docx() descartaria o render)
    d = doc.docx

    # data antes do bloco de assinatura (o modelo não traz)
    assinado = next(p for p in d.paragraphs if "assinado digitalmente" in p.text)
    for texto in (DATA, ""):
        novo = deepcopy(assinado._p)
        for r_ in list(novo):
            if r_.tag == qn("w:r"):
                novo.remove(r_)
        assinado._p.addprevious(novo)
        if texto:
            r_ = deepcopy(assinado.runs[0]._r)
            novo.append(r_)
            for t in r_.findall(qn("w:t")):
                t.text = texto

    # bloco data+assinatura inteiro na mesma página
    idx_data = next(i for i, p in enumerate(d.paragraphs) if p.text.startswith("Natal/RN"))
    for p in d.paragraphs[idx_data:-1]:
        p.paragraph_format.keep_with_next = True

    # assinatura em um único parágrafo, com quebras de linha (o modelo traz os três
    # em parágrafos separados)
    assinatura = next(p for p in d.paragraphs if "assinado digitalmente" in p.text)
    for _ in range(2):
        prox = assinatura._p.getnext()
        assinatura.runs[-1]._r.append(OxmlElement("w:br"))
        for run in prox.findall(qn("w:r")):
            assinatura._p.append(run)
        prox.getparent().remove(prox)

    # nome novo: não sobrescreve a informação de mérito substituída
    out = (Path(__file__).parent / f'{c["numero"]}_{c["ano"]}'
           / f'informacao_cip_{c["numero"]}_{c["ano"]}.docx')
    doc.save(str(out))

    # check (ponytail): o docx salvo é o declínio de atribuição, e não voltou ao mérito
    salvo = docx.Document(str(out))
    texto_final = "\n".join(p.text for p in salvo.paragraphs if p.text.strip())
    for trecho in (f'{c["numero"]}/{c["ano"]}', RELATOR, "Acórdão nº 293/2023-TC",
                   "006775/2018-TC", "Termo de Ajustamento de Gestão nº 04/2018",
                   c["responsavel"].split(", ")[0].split(" ", 2)[-1],  # o nome
                   f'Evento {c["ev_defesa"]}', f'Evento {c["ev_remessa"]}',
                   "art. 31", "art. 32", "Resolução nº 042/2024-TCE",
                   "Coordenadoria de Instrução Processual", DATA):
        assert trecho in texto_final, f"faltou: {trecho}"
    for proibido in ("tempestiv", "revel", "análise preliminar", "Cláusula Sétima",
                     "art. 37", "Ministério Público", "Aracieide"):
        assert proibido not in texto_final, f"o texto voltou ao mérito: {proibido!r}"
    assert len(salvo.tables) == 0, f"informação sem quadro, há {len(salvo.tables)}"
    assert len(corpo) == 5, "informação curta: 5 parágrafos"
    return out


def pdf(caminho: Path) -> Path:
    """docx_to_pdf chama word.Quit(); na conversão seguinte o Dispatch reencontra a
    instância morrendo e o SaveAs falha. Uma pausa e uma segunda tentativa bastam."""
    try:
        return docx_to_pdf(caminho, caminho.parent)
    except AttributeError:  # ponytail: só este script converte dois docx por execução
        time.sleep(5)
        return docx_to_pdf(caminho, caminho.parent)


for caderno in CADERNOS:
    caminho = gerar(caderno)
    print(f"Informação salva em: {caminho}")
    print(f"PDF gerado em: {pdf(caminho)}")
