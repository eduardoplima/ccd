"""Informações de arquivamento dos processos de NEREU BATISTA LINHARES.

Processos de execução do lote nereu_baixa em que o Relator decidiu sobre o saldo
residual do desconto em folha e a CCD executou as providências (baixa/retificação nos
sistemas e juntada da Certidão Declaratória de Quitação de Multa — art. 26, parágrafo
único, da Resolução nº 013/2015). Esta informação curta comunica o cumprimento e propõe
o envio à Diretoria de Expediente (DE) para arquivamento (art. 26, caput).

Lote 1 (000100/2023, 000106/2023, 001391/2023): decisão acolheu a informação de baixa e
determinou o cancelamento do resíduo com quitação — gerado, assinado e tramitado
CCD→DIP ("ENVIO A GCCTH") em 01/09/2026. Lote 2 (001417/2023, mesmo padrão; 003661/2022,
variante: a Decisão do Evento 124 levantou o sobrestamento — SS nº 5.749/STF —,
reconheceu a inexigibilidade do resíduo já cancelado e determinou exclusão do Cadastro
Informativo, certidão de quitação e retificação do motivo do cancelamento).

Os eventos citados são localizados ao vivo pelos títulos das informações
(vw_ata_informacao + Pro_ProcessoEvento): a decisão do gabinete ("Decisão."), a última
instrutiva da CCD anterior a ela (a informação de baixa acolhida), o documento de
cancelamento nos sistemas e a certidão de quitação. O script aborta se algum débito-filho
da cadeia ainda constar "Em Aberto" — não se afirma providência não concluída.

Modelo: scripts/automacao/templates/modelo_informacao.docx (parágrafos numerados pelo
modelo; a data entra por pós-processamento antes da assinatura).
Rodar: .venv/Scripts/python.exe processos/utils/gerar_informacoes_nereu_arquivamento.py
"""
import shutil
from copy import deepcopy
from datetime import datetime
from pathlib import Path

from docx.oxml.ns import qn
from docxtpl import DocxTemplate

from ccd.config import REPO_ROOT
from ccd.db import run_query_df
from ccd.docs import docx_to_pdf

BASE = Path(__file__).resolve().parents[1]  # processos/
TEMPLATE = str(REPO_ROOT / "scripts" / "automacao" / "templates" / "modelo_informacao.docx")
DESTINO = BASE / "projetos" / "nereu_arquivamento"

RESPONSAVEL = "Nereu Batista Linhares"

MESES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho",
         "agosto", "setembro", "outubro", "novembro", "dezembro"]
hoje = datetime.now()
DATA = f"Natal/RN, {hoje.day} de {MESES[hoje.month - 1]} de {hoje.year}."

# débitos-pai das execuções (mesmos ids do lote gerar_informacoes_nereu_baixa.py).
# Lote 1, já enviado em 01/09/2026: 23056 (000100/2023), 22859 (000106/2023),
# 22595 (001391/2023).
# 003661/2022: a certidão do Evento 127 cita a dívida 29.062 (do 001391/2023) em vez da
# 29.049 — erro formal avaliado pelo usuário como irrelevante à materialidade; mantido.
DEBITOS = [
    {"id": 23466},  # 001417/2023 — padrão do lote 1
    {"id": 22554,   # 003661/2022 — variante (quitado; Decisão Ev. 124)
     "titulo_cancelamento": "DESPACHO - genérico - DIP-DCC"},
]

# títulos dos eventos citados (conferidos no banco em 01/09/2026)
TITULO_DECISAO = "Decisão."  # GCCTH; "Decisão sobrestamento..." é outra, não conta
TITULO_INSTRUTIVA = "InformacaoInstrutiva"
TITULO_CANCELAMENTO = "Processo_Execucoes_Cancelamento_TCE.doc"
TITULO_CERTIDAO = "Certidão Declaratória de Quitação de Multa"


def carregar() -> list[dict]:
    ids = ",".join(str(d["id"]) for d in DEBITOS)
    pais = run_query_df(f"""
        SELECT e.IdDebito, s.DescricaoStatusDivida situacao,
               RTRIM(pe.numero_processo) numero, RTRIM(pe.ano_processo) ano,
               RTRIM(pe.assunto) assunto, RTRIM(pe.setor_atual) setor, re.nome relator
        FROM Exe_Debito e
        LEFT JOIN Exe_StatusDivida s ON s.CodigoStatusDivida = e.CodigoStatusDivida
        LEFT JOIN Processos pe ON pe.IdProcesso = e.IdProcessoExecucao
        LEFT JOIN Relator re ON re.codigo = pe.codigo_relator
        WHERE e.IdDebito IN ({ids})""").set_index("IdDebito")
    filhos = run_query_df(f"""
        SELECT f.IdDebitoAnterior pai, f.IdDebito, s.DescricaoStatusDivida situacao
        FROM Exe_Debito f
        LEFT JOIN Exe_StatusDivida s ON s.CodigoStatusDivida = f.CodigoStatusDivida
        WHERE f.IdDebitoAnterior IN ({ids})""")
    # baixa não concluída no banco de leitura? não gerar (conferir na Área Restrita:
    # o MSSQL de leitura atrasa — ver [[read-db-lags-area-restrita]])
    abertos = filhos[filhos.situacao == "Em Aberto"]
    assert abertos.empty, f"débito(s) ainda em aberto: {list(abertos.IdDebito)}"

    # `vw_ata_informacao.IdProcesso` vem NULL: filtrar por numero/ano; o "Evento" da
    # tela é o SequencialProcessoEvento (ver [[evento-e-sequencialprocessoevento]])
    chaves = {f"p{n}": f"{r.numero}/{r.ano}" for n, r in enumerate(pais.itertuples())}
    eventos = run_query_df(f"""
        SELECT RTRIM(v.numero_processo)+'/'+RTRIM(v.ano_processo) processo,
               RTRIM(v.setor) setor, RTRIM(v.Titulo_Modelo_informacao) titulo,
               ev.SequencialProcessoEvento evento
        FROM vw_ata_informacao v
        JOIN Pro_ProcessoEvento ev ON ev.idInformacao = v.idInformacao
        WHERE RTRIM(v.numero_processo)+'/'+RTRIM(v.ano_processo)
              IN ({",".join(f":{k}" for k in chaves)})""", **chaves)

    itens = []
    for deb in DEBITOS:
        id_debito = deb["id"]
        pai = pais.loc[id_debito]
        ev = eventos[eventos.processo == f"{pai.numero}/{pai.ano}"]
        decisao = int(ev[ev.titulo == TITULO_DECISAO].evento.max())
        antes = ev[(ev.titulo == TITULO_INSTRUTIVA) & (ev.evento < decisao)]
        depois = ev[ev.evento > decisao]
        titulo_canc = deb.get("titulo_cancelamento", TITULO_CANCELAMENTO)
        item = {"id": id_debito, "numero": pai.numero, "ano": pai.ano,
                "assunto": pai.assunto, "relator": pai.relator, "setor": pai.setor,
                "ev_informacao": int(antes.evento.max()),
                "ev_decisao": decisao,
                "ev_cancelamento": int(depois[depois.titulo == titulo_canc].evento.max()),
                "ev_certidao": int(depois[depois.titulo == TITULO_CERTIDAO].evento.max()),
                "filhos": [(int(r.IdDebito), r.situacao) for r in
                           filhos[filhos.pai == id_debito].itertuples()]}
        itens.append(item)
    return itens


FECHO = (
    "Ante o exposto, adotadas as providências determinadas e nada mais havendo a "
    "executar, sugere-se o envio dos autos à Diretoria de Expediente (DE) para "
    "arquivamento, nos termos do art. 26, caput, da Resolução nº 013/2015 – TCE/RN."
)


def montar_paragrafos(item: dict) -> list[str]:
    if item["id"] == 22554:  # 003661/2022 — dispositivo próprio da Decisão do Ev. 124
        return [
            (f"Trata-se de processo de execução de multa imputada ao Sr. {RESPONSAVEL}, "
             f"paga por desconto em folha, no qual o Exmo. Conselheiro Relator, por meio "
             f"da decisão exarada no Evento {item['ev_decisao']}, determinou o "
             "levantamento do sobrestamento do feito, reconheceu a inexigibilidade do "
             "saldo remanescente decorrente da correção monetária, autorizando o seu "
             "cancelamento, e determinou a exclusão do nome do responsável do Cadastro "
             "Informativo de Créditos não Quitados, com a expedição de certidão "
             "declaratória de quitação da dívida."),
            (f"Em cumprimento à decisão, procedeu-se ao cancelamento do saldo "
             f"remanescente nos sistemas deste Tribunal, com a retificação do respectivo "
             f"motivo (Evento {item['ev_cancelamento']}), e foi acostada aos autos a "
             f"Certidão Declaratória de Quitação de Multa (Evento {item['ev_certidao']})."),
            FECHO,
        ]
    return [
        (f"Trata-se de processo de execução de multa imputada ao Sr. {RESPONSAVEL}, "
         f"paga por desconto em folha, no qual o Exmo. Conselheiro Relator, "
         f"por meio da decisão exarada no Evento {item['ev_decisao']}, acolheu a sugestão "
         f"formulada por esta Coordenadoria (Evento {item['ev_informacao']}) e determinou "
         "o cancelamento do saldo residual da dívida, com a expedição de quitação ao "
         "responsável e o consequente arquivamento dos autos."),
        (f"Em cumprimento à decisão, procedeu-se à baixa do débito nos sistemas deste "
         f"Tribunal (Evento {item['ev_cancelamento']}) e foi acostada aos autos a Certidão "
         f"Declaratória de Quitação de Multa (Evento {item['ev_certidao']})."),
        FECHO,
    ]


def gerar(item: dict) -> Path:
    pasta = DESTINO / f"{item['numero']}_{item['ano']}"
    pasta.mkdir(parents=True, exist_ok=True)
    out = pasta / f"informacao_{item['numero']}_{item['ano']}.docx"

    doc = DocxTemplate(TEMPLATE)
    doc.render({
        "processo": f"{item['numero']}/{item['ano']} - TC",
        "assunto": item["assunto"],
        "relator": item["relator"].title(),
        "parágrafos": montar_paragrafos(item),  # sem numerar: o modelo numera
    })
    d = doc.docx  # pós-processamento (get_docx() descartaria o render)

    # data antes do bloco de assinatura (o modelo não traz); os parágrafos vazios que o
    # modelo deixa antes da assinatura saem para o conjunto caber em uma página
    assinado = next(p for p in d.paragraphs if "assinado digitalmente" in p.text)
    prev = assinado._p.getprevious()
    while prev is not None and not "".join(
            t.text or "" for t in prev.findall(f".//{qn('w:t')}")).strip():
        vazio, prev = prev, prev.getprevious()
        vazio.getparent().remove(vazio)
    novo = deepcopy(assinado._p)
    for r in list(novo):
        if r.tag == qn("w:r"):
            novo.remove(r)
    assinado._p.addprevious(novo)
    r = deepcopy(assinado.runs[0]._r)
    novo.append(r)
    for t in r.findall(qn("w:t")):
        t.text = DATA

    # bloco data+assinatura inteiro na mesma página
    idx_data = next(i for i, p in enumerate(d.paragraphs) if p.text.startswith("Natal/RN"))
    for p in d.paragraphs[idx_data:-1]:
        p.paragraph_format.keep_with_next = True

    # preserva a versão anterior antes de sobrescrever (skill edicao-minima)
    if out.exists():
        bak = out.with_name(f"{out.stem}_{datetime.now():%Y%m%d_%H%M%S}{out.suffix}")
        shutil.copy2(out, bak)
        print(f"  versão anterior preservada em: {bak.name}")
    doc.save(str(out))
    docx_to_pdf(str(out), str(pasta))
    return out


if __name__ == "__main__":
    import docx  # noqa: E402
    import pypdf  # noqa: E402

    for item in carregar():
        print(f"{item['numero']}/{item['ano']} (setor {item['setor']}, débito {item['id']}, "
              f"filhos {item['filhos']}, decisão Ev. {item['ev_decisao']}, "
              f"certidão Ev. {item['ev_certidao']})")
        # a cronologia dos eventos citados tem de fechar
        assert item["ev_informacao"] < item["ev_decisao"] < item["ev_cancelamento"], item
        assert item["ev_decisao"] < item["ev_certidao"], item
        out = gerar(item)

        # check (ponytail): nada de placeholder solto e os campos foram para o lugar certo
        texto = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
        assert "{{" not in texto and "}}" not in texto, "placeholder não substituído"
        for trecho in (f"{item['numero']}/{item['ano']} - TC", item["assunto"],
                       RESPONSAVEL, DATA, *montar_paragrafos(item)):
            assert trecho in texto, f"{item['numero']}: faltou \"{trecho[:40]}\""
        # a variante do 003661/2022 não cita a informação anterior (a Decisão não a acolheu)
        evs = ("ev_decisao", "ev_cancelamento", "ev_certidao") if item["id"] == 22554 \
            else ("ev_informacao", "ev_decisao", "ev_cancelamento", "ev_certidao")
        for ev in evs:
            assert f"Evento {item[ev]}" in texto, f"{item['numero']}: faltou {ev}"
        assert out.with_suffix(".pdf").is_file(), "PDF não gerado"
        paginas = len(pypdf.PdfReader(str(out.with_suffix(".pdf"))).pages)
        assert paginas == 1, f"{item['numero']}: {paginas} páginas (tem de caber em 1)"
        print(f"  salvo: {out.relative_to(BASE)} (+ .pdf)")
