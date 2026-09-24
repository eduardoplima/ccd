"""Informações de retomada da marcha processual — lote GCAED (Cons. Antônio Ed), marcador 6139.

Os processos com o marcador "Nereu- substituir informação" (Pro_Marcador 6139, posto pela CCD em
06/07/2026) têm todos a mesma marcha: despacho do GCAED em 11/12/2025 determinando "o sobrestamento dos
processos alcançados pela referida decisão [liminar do MS 0807247-93], até o desfecho judicial da
demanda" (art. 36, III, LC 464/2012 c/c art. 184, III, RI), peça da DIP_SOBR em 06/2026 e a instrutiva
da CCD de 17/07/2026 sugerindo suspensão até o trânsito do MS — hoje defasada: o STF suspendeu o
acórdão do TJRN na SS 5.749 (28/07/2026, trânsito 27/08/2026). A CCD propõe a continuidade da marcha,
com o despacho de processos/modelos/modelo_analise_retomada_nereu_gcaed.docx (derivado do modelo
GCREN, escrito para o 000098/2022 — Evento 72). A nova informação substitui a de 17/07 no cadastro.

Mesmo mecanismo de gerar_informacoes_nereu_retomada.py: só o número do processo e o do Evento mudam;
o 000098/2022 sai idêntico ao modelo, o que valida o gerador. Ao cadastrar, juntar os anexos
(decisão do STF + certidão) com nereu_retomada_gcaed/montar_envio.py.

Rodar: .venv/Scripts/python.exe processos/utils/gerar_informacoes_nereu_retomada_gcaed.py
"""
import re
import shutil
import unicodedata
from datetime import datetime
from pathlib import Path

import docx

from ccd.config import informacoes_dir
from ccd.db import run_query_df
from ccd.docs import docx_to_pdf
from ccd.pdf import extract_text_from_pdf

BASE = Path(__file__).resolve().parents[1]  # processos/
MODELO = BASE / "modelos" / "modelo_analise_retomada_nereu_gcaed.docx"
DESTINO = BASE / "projetos" / "nereu_retomada_gcaed"

MODELO_PROCESSO = "000098_2022"
MODELO_EVENTO = "72"
RELATOR = "ANTONIO ED SOUZA SANTANA"
SETOR_DESPACHO = "GCAED"
TITULO_SOBRESTAMENTO = "DESPACHO - SOBRESTAMENTO NA DIP - VT SESAP - EXECUÇÃO DA MULTA - APOSENTADORIA"
FRASE = (r"sobrestamento dos processos alcancados pela referida decisao, ate o desfecho judicial da "
         r"demanda, com espeque no art\. 36, inciso III da LC 464/2012")
MARCADOR = 6139
DATA_INSTRUTIVA = "2026-07-17"  # a informação da CCD que a nova vai substituir
ESPERADOS = 57

SQL_PROCESSOS = """
    SELECT RTRIM(p.numero_processo)+'/'+RTRIM(p.ano_processo) AS processo,
           RTRIM(p.setor_atual) AS setor, RTRIM(r.nome) AS relator
    FROM Pro_MarcadorProcesso mp
    JOIN Processos p ON p.IdProcesso = mp.IdProcesso
    LEFT JOIN Relator r ON r.codigo = p.codigo_relator
    WHERE mp.IdMarcador = :marcador AND mp.DataExclusao IS NULL
    ORDER BY p.ano_processo, p.numero_processo
"""
# vw_ata_informacao.IdProcesso vem NULL nas informações novas: filtrar por número/ano
SQL_EVENTOS = """
    SELECT RTRIM(v.numero_processo)+'/'+RTRIM(v.ano_processo) AS processo,
           RTRIM(v.setor) AS setor, RTRIM(v.Titulo_Modelo_informacao) AS titulo,
           CAST(v.data_resumo AS DATE) AS data, v.ordem, ev.SequencialProcessoEvento AS evento
    FROM vw_ata_informacao v
    JOIN Pro_ProcessoEvento ev ON ev.idInformacao = v.idInformacao
    WHERE RTRIM(v.numero_processo)+'/'+RTRIM(v.ano_processo) IN ({chaves})
      AND v.Inativa IS NULL
"""


def sem_acento(texto: str) -> str:
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", texto)


def carregar() -> list[dict]:
    processos = run_query_df(SQL_PROCESSOS, marcador=MARCADOR).set_index("processo")
    assert len(processos) == ESPERADOS, f"{len(processos)} processos com o marcador {MARCADOR}, esperava {ESPERADOS}"
    chaves = {f"p{i}": p for i, p in enumerate(processos.index)}
    eventos = run_query_df(SQL_EVENTOS.format(chaves=",".join(f":{k}" for k in chaves)), **chaves)

    itens = []
    for processo, proc in processos.iterrows():
        assert proc.relator == RELATOR, f"{processo}: relator {proc.relator}"
        assert proc.setor == "CCD", f"{processo}: está em {proc.setor}, não na CCD"

        ev = eventos[eventos.processo == processo]
        desp = ev[(ev.setor == SETOR_DESPACHO) & (ev.titulo == TITULO_SOBRESTAMENTO)]
        assert len(desp) == 1, f"{processo}: {len(desp)} despachos de sobrestamento, esperava 1"
        desp = desp.iloc[0]
        ultimo = ev.loc[ev.evento.idxmax()]
        assert ultimo.setor == "CCD" and str(ultimo.data) == DATA_INSTRUTIVA, (
            f"{processo}: último evento ativo é {ultimo.setor} de {ultimo.data}, não a instrutiva de {DATA_INSTRUTIVA}")

        numero, ano = processo.split("/")
        pdf = informacoes_dir() / SETOR_DESPACHO / f"{SETOR_DESPACHO}_{numero}_{ano}_{int(desp.ordem):04d}.pdf"
        assert re.search(FRASE, sem_acento(extract_text_from_pdf(str(pdf))), re.I), (
            f"{processo}: o despacho {pdf.name} não contém a frase citada")

        itens.append({"processo": processo, "evento": int(desp.evento)})
    return itens


def substituir_run(doc, contem: str, run_antigo: str, novo: str) -> None:
    """Troca o texto de um run no parágrafo que contém `contem`; o run tem de existir sozinho."""
    par = next(p for p in doc.paragraphs if contem in p.text)
    runs = [r for r in par.runs if r.text == run_antigo]
    assert len(runs) == 1, f'run "{run_antigo}" não é único em "{contem}"'
    runs[0].text = novo


def gerar(item: dict) -> Path:
    numero, ano = item["processo"].split("/")
    pasta = DESTINO / f"{numero}_{ano}"
    pasta.mkdir(parents=True, exist_ok=True)
    out = pasta / f"{numero}_{ano}.docx"

    doc = docx.Document(str(MODELO))
    cab = next(p for p in doc.paragraphs if p.text.startswith("Processo nº "))
    assert MODELO_PROCESSO in cab.text
    pedacos = [r for r in cab.runs if r.text not in ("Processo nº ", "-TC")]
    assert "".join(r.text for r in pedacos) == MODELO_PROCESSO
    pedacos[0].text = f"{numero}_{ano}"
    for r in pedacos[1:]:
        r.text = ""
    substituir_run(doc, "Evento nº ", MODELO_EVENTO, str(item["evento"]))

    if out.exists():  # preserva a versão anterior (skill edicao-minima)
        bak = out.with_name(f"{out.stem}_{datetime.now():%Y%m%d_%H%M%S}{out.suffix}")
        shutil.copy2(out, bak)
        print(f"  versão anterior preservada em: {bak.name}")
    doc.save(str(out))
    docx_to_pdf(str(out), str(pasta))
    return out


if __name__ == "__main__":
    modelo = [p.text for p in docx.Document(str(MODELO)).paragraphs]
    itens = carregar()
    for item in itens:
        numero, ano = item["processo"].split("/")
        print(f'{item["processo"]} (sobrestamento no Evento {item["evento"]}):')
        out = gerar(item)

        # check: só o cabeçalho e o nº do evento mudam; todo o resto é o modelo, palavra a palavra
        saida = [p.text for p in docx.Document(str(out)).paragraphs]
        assert len(saida) == len(modelo)
        assert saida[0] == f"Processo nº {numero}_{ano}-TC", saida[0]
        esperado = modelo[:]
        esperado[0] = saida[0]
        i = next(i for i, t in enumerate(modelo) if "Evento nº " in t)
        esperado[i] = modelo[i].replace(f"Evento nº {MODELO_EVENTO})", f'Evento nº {item["evento"]})')
        assert saida == esperado, f'{item["processo"]}: texto divergente do modelo'
        assert f'(Evento nº {item["evento"]})' in saida[i]
        assert out.with_suffix(".pdf").is_file(), "PDF não gerado"
        print(f"  salvo: {out.relative_to(BASE)} (+ .pdf)")

    print(f"\n{len(itens)} informações geradas em {DESTINO.relative_to(BASE)}/")
