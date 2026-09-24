"""Informações que submetem ao Relator a análise da prescrição.

Um documento por processo com a flag "prescrito" na aba *Risco de prescrição* da fila da
CCD: débito em aberto cuja data-base (última citação C05 ou, na falta dela, o trânsito em
julgado) já tem 5 anos ou mais. A regra canônica é a de
`web/backend/app/ccd/service.py::listar_prescricao` — aqui ela é reproduzida em SQL mais
o mesmo corte de 5*365 dias, porque a webapp roda em outro venv (workspace uv) e não é
importável da árvore raiz.

Duas diferenças deliberadas em relação à aba, que só sinaliza:
  * débitos já quitados (status 2/9) ficam de fora — a aba os mostra (é a ressalva
    `ponytail:` de service.py:202), mas não há o que prescrever neles;
  * quando o mesmo débito aparece em dois processos da CCD (origem e execução), a
    informação sai só na execução.

A informação NÃO afirma que houve prescrição: expõe os fatos (trânsito, última citação,
último ato processual, prazo decorrido), a base legal (arts. 328/329/330/332 do Regimento
Interno e o Tema 899/STF para o ressarcimento) e submete a questão ao Relator.

Modelo: scripts/automacao/templates/modelo_informacao.docx (parágrafos numerados pelo
modelo; a data entra por pós-processamento antes da assinatura).
Rodar: .venv/Scripts/python.exe processos/utils/gerar_informacoes_prescritos.py [001126/2015 ...]
"""
import shutil
import sys
from copy import deepcopy
from datetime import datetime
from pathlib import Path

import pandas as pd
from docx.oxml.ns import qn
from docxtpl import DocxTemplate

from ccd.config import REPO_ROOT
from ccd.db import get_connection, run_query_df
from ccd.docs import docx_to_pdf

BASE = Path(__file__).resolve().parents[1]  # processos/
TEMPLATE = str(REPO_ROOT / "scripts" / "automacao" / "templates" / "modelo_informacao.docx")
DESTINO = BASE / "projetos" / "prescritos"

PRAZO_DIAS = 5 * 365  # igual a service.py::_PRAZO_PRESCRICIONAL_DIAS (Tema 899/STF)
PARALISACAO_DIAS = 3 * 365  # art. 328 do RITCE (prescrição intercorrente)
PAGOS = (2, 9)  # Exe_StatusDivida: pago integralmente / pago aguardando compensação

# marcadores de permanência da CCD (MARCADORES_PERMANENCIA, service.py:23): parcelamento,
# desconto em folha, sobrestamento, decisão judicial, protesto, pagamento integral
MARCADORES_PERMANENCIA = (5020, 5032, 5040, 5041, 5469, 5591, 5593, 5684,
                          5712, 5790, 5797, 5846, 5966, 6127)

RELATORAS = {"ANA PAULA GOMES DE FREITAS"}  # o tratamento no fecho é flexionado

MESES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho",
         "agosto", "setembro", "outubro", "novembro", "dezembro"]
hoje = datetime.now()
HOJE = pd.Timestamp(hoje.date())
DATA = f"Natal/RN, {hoje.day} de {MESES[hoje.month - 1]} de {hoje.year}."

# um registro por débito aberto de processo hoje na CCD, com as duas datas que podem
# servir de marco prescricional. A escolha da base e o corte dos 5 anos ficam no pandas.
SQL_DEBITOS = f"""
WITH deb AS (
    -- débitos vigentes (folha da cadeia) e não quitados de processos hoje na CCD;
    -- vínculo pelos DOIS papéis (origem OU execução), nunca COALESCE
    SELECT p.IdProcesso AS id_ccd, e.IdDebito, e.IdProcessoOrigem, e.IdProcessoExecucao,
           e.valorOriginalDebito AS valor, e.dataTransito, e.CodigoTipoDebito,
           e.CodigoStatusDivida AS status_cod
    FROM dbo.Processos p
    JOIN dbo.Exe_Debito e
      ON e.IdProcessoOrigem = p.IdProcesso OR e.IdProcessoExecucao = p.IdProcesso
    JOIN dbo.Exe_StatusDivida sd ON sd.CodigoStatusDivida = e.CodigoStatusDivida
    WHERE p.setor_atual = 'CCD' AND p.IdProcessoApensador IS NULL
      AND NOT EXISTS (SELECT 1 FROM dbo.Exe_Debito g WHERE g.IdDebitoAnterior = e.IdDebito)
      AND sd.StatusCancelamento IS NULL
      AND NOT EXISTS (
          SELECT 1 FROM dbo.Pro_MarcadorProcesso mpx
          WHERE mpx.IdProcesso = p.IdProcesso AND mpx.DataExclusao IS NULL
            AND mpx.IdMarcador IN ({",".join(str(m) for m in MARCADORES_PERMANENCIA)}))
),
cit AS (
    -- citação C05 mais recente por processo (interrompe a prescrição);
    -- data via informação vinculada (Data_envio_AR é sempre vazia)
    SELECT c.IdProcesso, MAX(COALESCE(inf.DataPublicacao, inf.data_ultima_atualizacao,
                                      c.DataInclusao)) AS data_citacao
    FROM dbo.Cit_Citacoes c
    LEFT JOIN dbo.vw_ata_informacao inf ON inf.idInformacao = c.IdInformacao
    WHERE (c.Tipo = 'C05' OR (c.Tipo = 'C' AND c.Prazo = 5)) AND c.DataExclusao IS NULL
    GROUP BY c.IdProcesso
),
tj AS (
    -- trânsito por processo (fallback quando o débito não tem dataTransito)
    SELECT p2.IdProcesso, MAX(t.datatransito) AS data_transito
    FROM dbo.Processo_TransitoJulgado t
    JOIN dbo.Processos p2 ON p2.numero_processo = t.numero_processo
                         AND p2.ano_processo = t.ano_processo
    WHERE t.inativo = 0
    GROUP BY p2.IdProcesso
)
SELECT d.id_ccd, d.IdDebito, d.valor, d.status_cod, RTRIM(td.Descricao) AS tipo,
       RTRIM(p.numero_processo) + '/' + RTRIM(p.ano_processo) AS processo,
       RTRIM(p.numero_processo) AS numero, RTRIM(p.ano_processo) AS ano,
       RTRIM(p.assunto) AS assunto, RTRIM(r.nome) AS relator,
       RTRIM(po.numero_processo) + '/' + RTRIM(po.ano_processo) AS origem,
       RTRIM(px.numero_processo) + '/' + RTRIM(px.ano_processo) AS execucao,
       (SELECT MAX(v) FROM (VALUES (co.data_citacao), (ce.data_citacao)) x(v)) AS data_citacao,
       COALESCE(d.dataTransito, tjo.data_transito, tje.data_transito) AS data_transito
FROM deb d
JOIN dbo.Processos p     ON p.IdProcesso = d.id_ccd
LEFT JOIN dbo.Relator r  ON r.codigo = p.codigo_relator
LEFT JOIN dbo.Exe_TipoDebito td ON td.CodigoTipoDebito = d.CodigoTipoDebito
LEFT JOIN dbo.Processos po ON po.IdProcesso = d.IdProcessoOrigem
LEFT JOIN dbo.Processos px ON px.IdProcesso = d.IdProcessoExecucao
LEFT JOIN cit co ON co.IdProcesso = d.IdProcessoOrigem
LEFT JOIN cit ce ON ce.IdProcesso = d.IdProcessoExecucao
LEFT JOIN tj tjo ON tjo.IdProcesso = d.IdProcessoOrigem
LEFT JOIN tj tje ON tje.IdProcesso = d.IdProcessoExecucao
"""

SQL_RESPONSAVEIS = """
SELECT DISTINCT dp.IDDebito AS IdDebito, RTRIM(gp.Nome) AS nome
FROM dbo.Exe_DebitoPessoa dp
JOIN dbo.GenPessoa gp ON gp.IdPessoa = dp.IDPessoa
WHERE dp.IDDebito IN :debitos
"""

# citação C05 com o Evento em que a informação foi juntada (a citação costuma estar nos
# autos de origem, não na execução que está na CCD)
SQL_CITACOES = """
SELECT RTRIM(p.numero_processo) + '/' + RTRIM(p.ano_processo) AS processo,
       COALESCE(inf.DataPublicacao, inf.data_ultima_atualizacao, c.DataInclusao) AS data_cit,
       ev.SequencialProcessoEvento AS evento
FROM dbo.Cit_Citacoes c
JOIN dbo.Processos p ON p.IdProcesso = c.IdProcesso
LEFT JOIN dbo.vw_ata_informacao inf ON inf.idInformacao = c.IdInformacao
LEFT JOIN dbo.Pro_ProcessoEvento ev ON ev.IdInformacao = c.IdInformacao
                                   AND ev.IdProcesso = c.IdProcesso
WHERE (c.Tipo = 'C05' OR (c.Tipo = 'C' AND c.Prazo = 5)) AND c.DataExclusao IS NULL
  AND RTRIM(p.numero_processo) + '/' + RTRIM(p.ano_processo) IN :processos
"""

# último ato processual: o evento mais recente que tem informação com título (os eventos
# sem título são tramitações e movimentos internos, que não servem de referência)
SQL_ULTIMO_ATO = """
SELECT processo, evento, DataInclusao AS data_ato, titulo, setor FROM (
    SELECT RTRIM(p.numero_processo) + '/' + RTRIM(p.ano_processo) AS processo,
           ev.SequencialProcessoEvento AS evento, ev.DataInclusao,
           RTRIM(v.Titulo_Modelo_informacao) AS titulo, RTRIM(v.setor) AS setor,
           ROW_NUMBER() OVER (PARTITION BY p.IdProcesso
                              ORDER BY ev.SequencialProcessoEvento DESC) AS rn
    FROM dbo.Processos p
    JOIN dbo.Pro_ProcessoEvento ev ON ev.IdProcesso = p.IdProcesso
    JOIN dbo.vw_ata_informacao v ON v.idInformacao = ev.IdInformacao
    WHERE RTRIM(v.Titulo_Modelo_informacao) <> ''
      AND RTRIM(p.numero_processo) + '/' + RTRIM(p.ano_processo) IN :processos
) q WHERE rn = 1
"""

# a decisão condenatória (nº/ano) só existe em Processo_TransitoJulgado para parte dos
# processos — quando falta, a informação cita apenas a data do trânsito
SQL_DECISOES = """
SELECT processo, numerodecisao, anodecisao, tipodecisao, datatransito FROM (
    SELECT RTRIM(t.numero_processo) + '/' + RTRIM(t.ano_processo) AS processo,
           RTRIM(t.numerodecisao) AS numerodecisao, RTRIM(t.anodecisao) AS anodecisao,
           RTRIM(t.tipodecisao) AS tipodecisao, t.datatransito,
           ROW_NUMBER() OVER (PARTITION BY t.numero_processo, t.ano_processo
                              ORDER BY t.datatransito DESC) AS rn
    FROM dbo.Processo_TransitoJulgado t
    WHERE t.inativo = 0
      AND RTRIM(t.numero_processo) + '/' + RTRIM(t.ano_processo) IN :processos
) q WHERE rn = 1
"""

TIPOS_DECISAO = {"A": "Acórdão", "D": "Decisão", "R": "Resolução"}
# Exe_TipoDebito.Descricao -> como o débito é nomeado no texto (multa percentual é multa)
TIPOS_DEBITO = {"Multa": "multa", "Multa Percentual": "multa",
                "Multa Cominatória": "multa cominatória", "Ressarcimento": "ressarcimento",
                "Remanejamento": "remanejamento"}
CONECTIVOS = {"de", "da", "do", "das", "dos", "e"}


def nome_proprio(nome: str) -> str:
    return " ".join(w if (w := p.title().strip()).lower() not in CONECTIVOS else w.lower()
                    for p in nome.split())


def enumerar(itens: list[str]) -> str:
    return " e ".join(filter(None, [", ".join(itens[:-1]), itens[-1]])) if itens else ""


def brl(v: float) -> str:
    return f"R$ {v:,.2f}".replace(",", "@").replace(".", ",").replace("@", ".")


def dt(d) -> str:
    return pd.Timestamp(d).strftime("%d/%m/%Y")


def periodo(dias: int) -> str:
    anos, resto = divmod(int(dias), 365)
    if not anos:
        return f"{resto} dias"
    return f"{anos} ano{'s' if anos > 1 else ''}" + (f" e {resto} dias" if resto else "")


def carregar(filtro: set[str] | None = None) -> tuple[list[dict], list[str]]:
    eng = get_connection()
    deb = run_query_df(SQL_DEBITOS, eng)
    for col in ("data_citacao", "data_transito"):
        deb[col] = pd.to_datetime(deb[col])
    deb["data_base"] = deb.data_citacao.fillna(deb.data_transito)
    deb = deb[deb.data_base.notna()]
    deb["dias"] = (HOJE - deb.data_base.dt.normalize()).dt.days
    # o processo prescreve pelo débito de base mais antiga (o primeiro a prescrever)
    prescrito = deb.groupby("id_ccd").dias.transform("max") >= PRAZO_DIAS

    pulados: list[str] = []
    # a aba mostra débito quitado como "aberto" (só exclui status de cancelamento):
    # sem saldo, não há pretensão executória a prescrever
    aberto = ~deb.status_cod.isin(PAGOS)
    prescritos = deb[prescrito & aberto]
    for p in sorted(set(deb[prescrito].processo) - set(prescritos.processo)):
        pulados.append(f"{p}: todos os débitos quitados (a aba os conta como abertos)")
    # mesmo débito em dois processos da CCD (origem e execução): fica a execução.
    # Um processo só é descartado se TODOS os seus débitos correm em outro.
    dono = {}
    for id_debito, g in prescritos.groupby("IdDebito"):
        if len(g) > 1:
            dono[id_debito] = g.execucao.iloc[0]
    descartar = set()
    for processo, g in prescritos.groupby("processo"):
        alheios = {dono.get(i, processo) for i in g.IdDebito}
        if processo not in alheios:
            descartar.add(processo)
            pulados.append(f"{processo}: débito(s) {sorted(g.IdDebito)} corre(m) em "
                           f"{', '.join(sorted(alheios))} (execução)")

    resp = run_query_df(SQL_RESPONSAVEIS, eng,
                        debitos=tuple(int(i) for i in prescritos.IdDebito))
    responsaveis = resp.groupby("IdDebito").nome.apply(lambda s: sorted(set(s))).to_dict()

    envolvidos = tuple(sorted({*prescritos.processo, *prescritos.origem.dropna(),
                               *prescritos.execucao.dropna()}))
    citacoes = run_query_df(SQL_CITACOES, eng, processos=envolvidos)
    citacoes["data_cit"] = pd.to_datetime(citacoes.data_cit)
    atos = run_query_df(SQL_ULTIMO_ATO, eng, processos=envolvidos).set_index("processo")
    decisoes = run_query_df(SQL_DECISOES, eng, processos=envolvidos).set_index("processo")

    itens = []
    for processo, g in prescritos.groupby("processo", sort=False):
        g = g.sort_values("dias", ascending=False)
        primeiro = g.iloc[0]  # débito que prescreve primeiro: define a data-base
        if processo in descartar or (filtro and processo not in filtro):
            continue
        outros = [p for p in (primeiro.origem, primeiro.execucao)
                  if p and p != processo]
        # citação e trânsito costumam estar nos autos de origem
        cits = citacoes[citacoes.processo.isin([processo, *outros])]
        cit = cits.sort_values("data_cit").iloc[-1] if len(cits) else None
        dec = next((decisoes.loc[p] for p in (primeiro.origem, primeiro.execucao, processo)
                    if p in decisoes.index), None)
        ato = atos.loc[processo] if processo in atos.index else None
        nomes = sorted({n for i in g.IdDebito for n in responsaveis.get(i, [])})
        itens.append({
            "processo": processo, "numero": primeiro.numero, "ano": primeiro.ano,
            "assunto": primeiro.assunto, "relator": primeiro.relator,
            "responsaveis": nomes, "tipos": sorted(set(g.tipo.dropna())),
            "qtd": len(g), "valor": float(g.valor.sum()),
            "origem": primeiro.origem if primeiro.origem != processo else None,
            "data_base": primeiro.data_base, "dias": int(primeiro.dias),
            "fonte": "citação" if pd.notna(primeiro.data_citacao) else "trânsito",
            "data_transito": primeiro.data_transito,
            "decisao": None if dec is None else dict(dec),
            "citacao": None if cit is None else dict(cit),
            "ato": None if ato is None else dict(ato, processo=processo),
        })
    return itens, pulados


def montar_paragrafos(item: dict) -> list[str]:
    tipos = enumerar(sorted({TIPOS_DEBITO.get(t, t.lower()) for t in item["tipos"]}))
    nomes = enumerar([nome_proprio(n) for n in item["responsaveis"]]) or "responsável não identificado"
    dec = item["decisao"]
    origem = f", oriundo do Processo nº {item['origem']}" if item["origem"] else ""
    titulo = ""
    if dec and dec["numerodecisao"]:
        tipo_dec = TIPOS_DECISAO.get((dec["tipodecisao"] or "").strip(), "decisão")
        titulo = f", decorrente do {tipo_dec} nº {dec['numerodecisao']}/{dec['anodecisao']}"
    debitos = (f"{item['qtd']} débitos em aberto, no valor originário total de "
               f"{brl(item['valor'])}" if item["qtd"] > 1 else
               f"débito em aberto, no valor originário de {brl(item['valor'])}")
    p1 = (f"Trata-se de processo de execução de {tipos} em desfavor de {nomes}{titulo}"
          f"{origem}, com trânsito em julgado em {dt(item['data_transito'])}, em cobrança "
          f"nesta Coordenadoria, remanescendo {debitos}.")

    cit = item["citacao"]
    if cit is not None:
        onde = f"nos autos do Processo nº {cit['processo']}" if cit["processo"] != item["processo"] \
            else "nestes autos"
        evento = f", Evento {int(cit['evento'])}" if pd.notna(cit["evento"]) else ""
        p2 = (f"O último ato interruptivo do prazo prescricional registrado é a citação do "
              f"responsável, com prazo de 5 (cinco) dias, realizada em {dt(cit['data_cit'])} "
              f"({onde}{evento}), na forma do art. 329, inciso I, e do art. 332, parágrafo "
              f"único, do Regimento Interno. Desde então decorreram "
              f"{periodo(item['dias'])}, sem que conste dos autos nova causa de interrupção.")
    else:
        p2 = (f"Não consta dos autos citação do responsável na fase de execução, de modo que "
              f"o prazo prescricional corre do trânsito em julgado da decisão condenatória, "
              f"ocorrido há {periodo(item['dias'])}, sem causa de interrupção registrada.")

    ato = item["ato"]
    if ato is not None:
        dias_ato = (HOJE - pd.Timestamp(ato["data_ato"]).normalize()).days
        p3 = (f"O último ato processual praticado é \"{ato['titulo']}\" ({ato['setor']}), de "
              f"{dt(ato['data_ato'])} (Evento {int(ato['evento'])})")
        p3 += (f", estando o feito sem impulso há {periodo(dias_ato)}, o que atrai o exame do "
               f"art. 328 do Regimento Interno." if dias_ato > PARALISACAO_DIAS else ".")
    else:
        p3 = "Não há nos autos ato processual posterior com título registrado no sistema."

    p4 = ("Dispõe o art. 332 do Regimento Interno que, após o trânsito em julgado da decisão "
          "condenatória, prescreve em cinco anos a pretensão executória relativa a crédito "
          "decorrente da aplicação de multa, interrompendo-se o prazo pela citação da parte, "
          "inclusive por edital, e suspendendo-se pelo período de cumprimento do parcelamento. "
          "Suspendem ainda o prazo as hipóteses do art. 330, e o art. 328 cuida da prescrição "
          "intercorrente no processo paralisado por mais de três anos.")
    if any(t.lower().startswith("ressarcimento") for t in item["tipos"]):
        p4 += (" Quanto ao débito de ressarcimento, o Supremo Tribunal Federal, no julgamento "
               "do Tema 899 da repercussão geral (RE 636.886), firmou a prescritibilidade da "
               "pretensão de ressarcimento ao erário fundada em decisão de Tribunal de Contas.")

    tratamento = "Exma. Sra. Conselheira Relatora" if (item["relator"] or "").upper() \
        in RELATORAS else "Exmo. Sr. Conselheiro Relator"
    p5 = (f"Ante o exposto, e por ser o reconhecimento da prescrição matéria afeta ao juízo do "
          f"Relator, sugere-se o encaminhamento dos autos ao {tratamento}, para que se pronuncie "
          f"acerca da ocorrência ou não da prescrição da pretensão executória e delibere sobre o "
          f"prosseguimento ou a extinção da cobrança.")
    return [p1, p2, p3, p4, p5]


def gerar(item: dict) -> Path:
    pasta = DESTINO / f"{item['numero']}_{item['ano']}"
    pasta.mkdir(parents=True, exist_ok=True)
    out = pasta / f"informacao_{item['numero']}_{item['ano']}.docx"

    doc = DocxTemplate(TEMPLATE)
    doc.render({
        "processo": f"{item['numero']}/{item['ano']} - TC",
        "assunto": item["assunto"],
        "relator": (item["relator"] or "").title(),
        "parágrafos": montar_paragrafos(item),  # sem numerar: o modelo numera
    })
    d = doc.docx  # pós-processamento (get_docx() descartaria o render)

    # data antes do bloco de assinatura (o modelo não traz); os parágrafos vazios que o
    # modelo deixa antes da assinatura saem para o conjunto caber em uma página
    assinado = next(p for p in d.paragraphs if "assinado digitalmente" in p.text)
    prev = assinado._p.getprevious()
    while prev is not None and not "".join(
            t.text or "" for t in prev.findall(f".//{qn('w:t')}")).strip():
        vazio, prev = prev, prev.getprevious()
        vazio.getparent().remove(vazio)
    novo = deepcopy(assinado._p)
    for r in list(novo):
        if r.tag == qn("w:r"):
            novo.remove(r)
    assinado._p.addprevious(novo)
    r = deepcopy(assinado.runs[0]._r)
    novo.append(r)
    for t in r.findall(qn("w:t")):
        t.text = DATA

    # bloco data+assinatura inteiro na mesma página
    idx_data = next(i for i, p in enumerate(d.paragraphs) if p.text.startswith("Natal/RN"))
    for p in d.paragraphs[idx_data:-1]:
        p.paragraph_format.keep_with_next = True

    # preserva a versão anterior antes de sobrescrever (skill edicao-minima)
    if out.exists():
        bak = out.with_name(f"{out.stem}_{datetime.now():%Y%m%d_%H%M%S}{out.suffix}")
        shutil.copy2(out, bak)
        print(f"  versão anterior preservada em: {bak.name}")
    doc.save(str(out))
    docx_to_pdf(str(out), str(pasta))
    return out


if __name__ == "__main__":
    import docx  # noqa: E402

    itens, pulados = carregar(set(sys.argv[1:]) or None)
    for p in pulados:
        print(f"PULADO  {p}")

    for item in itens:
        cit = item["citacao"]
        print(f"{item['processo']} ({item['relator']}, {item['qtd']} débito(s), "
              f"{brl(item['valor'])}, base {item['fonte']} {dt(item['data_base'])}, "
              f"{item['dias']} dias)")
        # os fatos citados têm de fechar cronologicamente
        assert item["dias"] >= PRAZO_DIAS, item["processo"]
        assert pd.notna(item["data_transito"]), f"{item['processo']}: sem trânsito"
        if cit is not None:
            assert pd.Timestamp(cit["data_cit"]) <= HOJE, item["processo"]
        if item["ato"] is not None:
            assert pd.Timestamp(item["ato"]["data_ato"]) <= HOJE, item["processo"]

        out = gerar(item)

        # check (ponytail): nada de placeholder solto e os fatos foram para o documento
        texto = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
        assert "{{" not in texto and "}}" not in texto, "placeholder não substituído"
        for trecho in (f"{item['numero']}/{item['ano']} - TC", item["assunto"], DATA,
                       dt(item["data_transito"]), *montar_paragrafos(item)):
            assert trecho in texto, f"{item['processo']}: faltou \"{trecho[:40]}\""
        assert out.with_suffix(".pdf").is_file(), "PDF não gerado"
        print(f"  salvo: {out.relative_to(BASE)} (+ .pdf)")

    print(f"\n{len(itens)} informações geradas em {DESTINO.relative_to(REPO_ROOT)}")
