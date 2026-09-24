"""Informações de remessa ao MPC das execuções de NEREU BATISTA LINHARES.

Em 25 e 26/05/2026 o Conselheiro Marco Antônio de Moraes Rêgo Montenegro, em substituição
legal no gabinete do Conselheiro Antônio Ed (GCAED), exarou um despacho em lote nas execuções
contra o responsável. O item 23 determina que a DIP (alínea a) certifique a inviabilidade do
desconto em folha e (alínea b) encaminhe os autos ao Ministério Público de Contas, sem retorno
ao gabinete, para as providências junto à PGE. Esta é a informação que cumpre esse item.

O lote foi apurado sobre os 77 processos de docs/notas/PROCESSOS_MS_0807247_CCD.md: lendo o PDF de cada
peça de gabinete cadastrada a partir de 2025 e procurando a assinatura do Conselheiro
substituto, passam 19 processos — exatamente os 19 que aquele levantamento classificara como
"sem padrão". Os outros 58 só têm o despacho de sobrestamento de 11/12/2025, que não é dele.

O texto é o do 002551/2024, que saiu à mão e serve de modelo — o gerador o reproduz palavra
por palavra a partir dos dados do banco, e é isso que valida os outros 18.

O texto é o mesmo do modelo, com quatro pontos variáveis:

- o acórdão exequendo e o processo de origem (parágrafo 1);
- "multa cominatória" ou "multa", conforme o `CodigoTipoDebito` do débito vigente (5 ou 2).
  Oito dos 18 são multa comum. A cláusula "pelo descumprimento de determinação desta Corte"
  vale para ambos: é como o próprio despacho anterior da CCD descreve esses autos;
- o evento do despacho do Conselheiro substituto (parágrafos 2, 4 e 5);
- o evento em que a CCD já havia certificado a inviabilidade (parágrafo 4).

O que é constante nos 19, conferido peça a peça e não presumido: o Memorando é sempre o
nº 000040/2026-DIP, citado no item 19; a determinação é sempre o item 23, alíneas a e b no
mesmo teor. Os asserts de `carregar` param o script se algum despacho fugir do lote.

Duas ressalvas do levantamento:

- 002548/2024 — o despacho diz que a origem é 013501/2017, mas esse processo é uma averbação
  de tempo de serviço arquivada. O débito aponta para 013501/2016 ("Apreciação de concessão de
  aposentadoria"), que é o que consta do assunto da execução. Vale o banco (ver `ORIGEM_ERRADA`).
- 000099/2022 e 001414/2022 têm mais de um débito na folha da cadeia, com os demais cancelados
  por erro de cadastro. Vale o que está Em Aberto.

Base: scripts/automacao/templates/modelo_informacao.docx (o mesmo do modelo 002551/2024).
Rodar: .venv/Scripts/python.exe processos/utils/gerar_informacoes_nereu_mpc.py
"""
import re
import shutil
import unicodedata
from copy import deepcopy
from datetime import datetime
from pathlib import Path

from docx.enum.text import WD_LINE_SPACING
from docx.shared import Pt
from docxtpl import DocxTemplate

from ccd.config import REPO_ROOT, informacoes_dir
from ccd.db import run_query_df
from ccd.docs import docx_to_pdf
from ccd.pdf import extract_text_from_pdf

BASE = Path(__file__).resolve().parents[1]  # processos/
TEMPLATE = str(REPO_ROOT / "scripts" / "automacao" / "templates" / "modelo_informacao.docx")
DESTINO = BASE / "projetos" / "nereu_mpc_antonio_ed"

RESPONSAVEL = "Nereu Batista Linhares"
CONSELHEIRO = "Marco Antônio de Moraes Rêgo Montenegro"
MEMORANDO = "000040/2026-DIP"
ITEM_MEMORANDO = 19  # item do despacho que cita o memorando
ITEM_DETERMINACAO = 23  # item que determina a certificação (a) e a remessa ao MPC (b)

PROCESSOS = [
    "000099/2022", "000124/2022", "000137/2022", "000146/2022", "000147/2022",
    "000151/2022", "001414/2022", "001436/2022", "002708/2022",
    "002029/2024", "002038/2024", "002047/2024", "002548/2024", "002549/2024",
    "002550/2024", "002551/2024", "002552/2024", "002553/2024", "002555/2024",
]

# o despacho do gabinete errou o ano da origem; o débito e o assunto da execução dizem 2016
ORIGEM_ERRADA = {"002548/2024": "013501/2017"}

TIPO_COMINATORIA = 5  # Exe_Debito.CodigoTipoDebito; 2 = multa comum

ASSINATURA = ("[assinado digitalmente]", "Eduardo Pereira Lima",
              "Auditor de Controle Externo", "Coordenador de Controle de Decisões")

ABERTURA = (
    "Trata-se de processo de execução do Acórdão nº {acordao}, proferido nos autos do processo "
    "nº {origem}-TC, relativo à cobrança de {multa} imputada ao Sr. {responsavel} pelo "
    "descumprimento de determinação desta Corte."
)
RETORNO = (
    "Os autos retornaram à DIP por determinação do Conselheiro {conselheiro}, em substituição "
    "legal (evento {ev_despacho}), para certificação da inviabilidade do desconto em folha e "
    "posterior encaminhamento ao Ministério Público de Contas, sem retorno ao gabinete."
)
INVIABILIDADE = (
    "No tocante à inviabilidade, cumpre informar que o responsável já se encontra submetido a "
    "descontos em folha decorrentes de outras execuções desta Corte, cujo somatório consome "
    "integralmente a margem consignável, esgotando o limite legal a que se refere o art. 339, "
    "II, do Regimento Interno — que condiciona o desconto aos limites previstos na legislação "
    "aplicável —, de modo que não há espaço para a implantação de novo desconto relativo a "
    "estes autos."
)
REITERACAO = (
    "A inviabilidade já havia sido certificada por esta unidade no evento {ev_certidao}, ora "
    "reiterada, no mesmo sentido do Memorando nº {memorando}, citado no parágrafo "
    "{item_memorando} do despacho no evento {ev_despacho}, que aponta a necessidade de "
    "encaminhamento ao MPC dos processos sem viabilidade de desconto, para providências "
    "junto à PGE."
)
ENCAMINHAMENTO = (
    "Ante o exposto, em cumprimento ao item b do despacho no evento {ev_despacho}, remetem-se "
    "os autos ao Ministério Público de Contas, sem retorno ao gabinete, para providências "
    "junto à Procuradoria-Geral do Estado destinadas à inscrição em dívida ativa e à cobrança "
    "judicial, nos termos do art. 339, III, do Regimento Interno, com redação dada pela "
    "Resolução nº 013/2015-TCE."
)

SQL_EVENTOS = """
    SELECT RTRIM(v.numero_processo)+'/'+RTRIM(v.ano_processo) AS processo,
           RTRIM(v.setor) AS setor, RTRIM(v.Titulo_Modelo_informacao) AS titulo,
           v.ordem, v.DataInclusao, ev.SequencialProcessoEvento AS evento
    FROM vw_ata_informacao v
    JOIN Pro_ProcessoEvento ev ON ev.idInformacao = v.idInformacao
    WHERE RTRIM(v.numero_processo)+'/'+RTRIM(v.ano_processo) IN ({chaves})
"""
SQL_PROCESSOS = """
    SELECT RTRIM(p.numero_processo)+'/'+RTRIM(p.ano_processo) AS processo,
           RTRIM(p.assunto) AS assunto, RTRIM(p.setor_atual) AS setor, r.nome AS relator
    FROM Processos p
    LEFT JOIN Relator r ON r.codigo = p.codigo_relator
    WHERE RTRIM(p.numero_processo)+'/'+RTRIM(p.ano_processo) IN ({chaves})
"""
# débito vigente = folha da cadeia ([[exe-debito-cadeia-folha-vigente]]); o processo entra pelos
# dois papéis, e a origem sai do IdProcessoOrigem desse mesmo débito
SQL_DEBITOS = """
    SELECT RTRIM(p.numero_processo)+'/'+RTRIM(p.ano_processo) AS processo,
           e.IdDebito, e.CodigoTipoDebito AS tipo,
           s.DescricaoStatusDivida AS situacao,
           RTRIM(o.numero_processo)+'/'+RTRIM(o.ano_processo) AS origem
    FROM Exe_Debito e
    LEFT JOIN Exe_StatusDivida s ON s.CodigoStatusDivida = e.CodigoStatusDivida
    LEFT JOIN Processos o ON o.IdProcesso = e.IdProcessoOrigem
    JOIN Processos p ON p.IdProcesso IN (e.IdProcessoOrigem, e.IdProcessoExecucao)
    WHERE RTRIM(p.numero_processo)+'/'+RTRIM(p.ano_processo) IN ({chaves})
      AND NOT EXISTS (SELECT 1 FROM Exe_Debito g WHERE g.IdDebitoAnterior = e.IdDebito)
"""


def sem_acento(texto: str) -> str:
    """Texto do PDF sem diacríticos e com espaços colapsados.

    O extrator quebra palavras e números no meio ("Ant onio", "2 268/2020") e o `nº` sai com
    grafias variadas — sem normalizar, nenhum casamento literal funciona."""
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", texto)


def so_alfanum(texto: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "", sem_acento(texto)).upper()


def caminho_pdf(processo: str, setor: str, ordem: int) -> Path:
    numero, ano = processo.split("/")
    return informacoes_dir() / setor / f"{setor}_{numero}_{ano}_{ordem:04d}.pdf"


def do_conselheiro(texto: str) -> bool:
    marcas = so_alfanum(texto)
    return so_alfanum(CONSELHEIRO) in marcas and "EMSUBSTITUICAOLEGAL" in marcas


def conferir_despacho(processo: str, texto: str) -> None:
    """O texto do parágrafo cita item, alínea e memorando como constantes: se o despacho deste
    processo não for o do lote, é melhor parar do que emitir uma referência falsa."""
    t = sem_acento(texto)
    assert re.search(rf"Memorando n.{{0,3}}\s*{MEMORANDO[:6]}\s*/\s*2026\s*-?\s*DIP", t, re.I), (
        f"{processo}: despacho não cita o Memorando nº {MEMORANDO}")
    assert re.search(rf"{ITEM_MEMORANDO}\.\s+Em reforco a esse cenario, o Memorando", t), (
        f"{processo}: o Memorando não está no item {ITEM_MEMORANDO}")
    assert re.search(rf"{ITEM_DETERMINACAO}\.\s+Diante do exposto, determino", t), (
        f"{processo}: a determinação não está no item {ITEM_DETERMINACAO}")
    assert "a) Proceda a certificacao da inviabilidade" in t, f"{processo}: alínea a diferente"
    assert "b) Apos, sem necessidade de retornar os autos ao meu Gabinete" in t, (
        f"{processo}: alínea b diferente")


def acordao_do_despacho(processo: str, texto: str) -> str:
    """Nº do acórdão exequendo, do item 1 do despacho. Tolera espaços dentro do número."""
    m = re.search(r"Acordao\s*n[o°.\s]*((?:\d\s*){1,6})/\s*(\d{4})\s*-?\s*TC",
                  sem_acento(texto), re.I)
    assert m, f"{processo}: acórdão não localizado no despacho"
    return f'{re.sub(r"\s+", "", m.group(1))}/{m.group(2)}-TC'


def carregar() -> list[dict]:
    chaves = {f"p{i}": p for i, p in enumerate(PROCESSOS)}
    marcadores = ",".join(f":{k}" for k in chaves)
    eventos = run_query_df(SQL_EVENTOS.format(chaves=marcadores), **chaves)
    processos = run_query_df(SQL_PROCESSOS.format(chaves=marcadores), **chaves).set_index("processo")
    debitos = run_query_df(SQL_DEBITOS.format(chaves=marcadores), **chaves)

    itens = []
    for processo in PROCESSOS:
        ev = eventos[eventos.processo == processo]

        # o despacho é a peça de gabinete assinada pelo Conselheiro substituto
        candidatas = ev[ev.setor.str.startswith("GC") & (ev.DataInclusao >= "2025-01-01")]
        despachos = [(int(r.evento), r.setor, int(r.ordem),
                      extract_text_from_pdf(str(caminho_pdf(processo, r.setor, int(r.ordem)))))
                     for r in candidatas.itertuples()]
        despachos = [d for d in despachos if do_conselheiro(d[3])]
        assert len(despachos) == 1, (
            f"{processo}: {len(despachos)} despachos do Conselheiro substituto, esperava 1")
        ev_despacho, _, _, texto = despachos[0]
        conferir_despacho(processo, texto)

        # a certificação anterior da CCD é o despacho que já delineara a inviabilidade
        anteriores = ev[(ev.setor == "CCD") & (ev.evento < ev_despacho)
                        & (ev.titulo.str.strip() == "Complementar Processo")]
        assert not anteriores.empty, f"{processo}: certificação anterior da CCD não localizada"
        ev_certidao = int(anteriores.evento.max())

        # entre os débitos vigentes vale o que não foi cancelado
        deb = debitos[(debitos.processo == processo)
                      & ~debitos.situacao.str.startswith("Cancelada")]
        assert len(deb) == 1, f"{processo}: {len(deb)} débitos vigentes não cancelados, esperava 1"
        deb = deb.iloc[0]

        origem_despacho = re.search(r"processo n[o°.\s]*((?:\d\s*){1,6})/\s*(\d{4})\s*-?\s*TC",
                                    sem_acento(texto), re.I)
        origem_despacho = (f'{re.sub(r"\s+", "", origem_despacho.group(1)).zfill(6)}'
                           f'/{origem_despacho.group(2)}') if origem_despacho else None
        assert origem_despacho in (deb.origem, ORIGEM_ERRADA.get(processo)), (
            f"{processo}: origem {deb.origem} no débito e {origem_despacho} no despacho")

        proc = processos.loc[processo]
        itens.append({
            "processo": processo, "assunto": proc.assunto, "relator": proc.relator.title(),
            "setor": proc.setor, "id_debito": int(deb.IdDebito), "tipo": int(deb.tipo),
            "origem": deb.origem, "acordao": acordao_do_despacho(processo, texto),
            "ev_despacho": ev_despacho, "ev_certidao": ev_certidao,
        })
    return itens


def contexto(item: dict) -> dict:
    return {
        "processo": item["processo"],
        "assunto": item["assunto"],
        "relator": item["relator"],
        "parágrafos": [
            ABERTURA.format(acordao=item["acordao"], origem=item["origem"],
                            responsavel=RESPONSAVEL,
                            multa="multa cominatória" if item["tipo"] == TIPO_COMINATORIA
                            else "multa"),
            RETORNO.format(conselheiro=CONSELHEIRO, ev_despacho=item["ev_despacho"]),
            INVIABILIDADE,
            REITERACAO.format(ev_certidao=item["ev_certidao"], memorando=MEMORANDO,
                              item_memorando=ITEM_MEMORANDO, ev_despacho=item["ev_despacho"]),
            ENCAMINHAMENTO.format(ev_despacho=item["ev_despacho"]),
        ],
    }


def bloco_assinatura(doc) -> None:
    """Monta o bloco de assinatura: insere o cargo de auditor e aperta o espaçamento.

    A linha "Auditor de Controle Externo" está no modelo 002551/2024 e não no template —
    clona-se o parágrafo do nome para herdar centralização e fonte, reescrevendo o texto do
    primeiro run (idioma de processos/informacoes/001454_2023/gerar_informacao.py).

    As quatro linhas saem em espaço simples e sem espaço depois: o padrão do documento é 1,15
    com 10 pt após cada parágrafo, o que esparramaria a assinatura. Só o bloco é alterado —
    o corpo mantém o espaçamento do template."""
    nome = next(p for p in doc.docx.paragraphs if p.text.strip() == "Eduardo Pereira Lima")
    linha = deepcopy(nome._p)
    for run in linha.findall(f'{{{nome._p.nsmap["w"]}}}r')[1:]:
        linha.remove(run)
    linha.findall(f'{{{nome._p.nsmap["w"]}}}r')[0].find(
        f'{{{nome._p.nsmap["w"]}}}t').text = "Auditor de Controle Externo"
    nome._p.addnext(linha)

    for paragrafo in doc.docx.paragraphs:
        if paragrafo.text.strip() in ASSINATURA:
            paragrafo.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            paragrafo.paragraph_format.space_after = Pt(0)


def gerar(item: dict) -> Path:
    numero, ano = item["processo"].split("/")
    pasta = DESTINO / f"{numero}_{ano}"
    pasta.mkdir(parents=True, exist_ok=True)
    out = pasta / f"{numero}_{ano}.docx"
    doc = DocxTemplate(TEMPLATE)
    doc.render(contexto(item))
    bloco_assinatura(doc)  # depois do render: get_docx() antes descartaria a renderização
    # preserva a versão anterior antes de sobrescrever (skill edicao-minima)
    if out.exists():
        bak = out.with_name(f"{out.stem}_{datetime.now():%Y%m%d_%H%M%S}{out.suffix}")
        shutil.copy2(out, bak)
        print(f"  versão anterior preservada em: {bak.name}")
    doc.save(str(out))
    docx_to_pdf(str(out), str(pasta))
    return out


if __name__ == "__main__":
    import docx  # noqa: E402

    itens = carregar()

    for item in itens:
        ctx = contexto(item)
        multa = "cominatória" if item["tipo"] == TIPO_COMINATORIA else "multa comum"
        print(f'{item["processo"]} (setor {item["setor"]}, débito {item["id_debito"]}, {multa}, '
              f'Acórdão {item["acordao"]}, origem {item["origem"]}, '
              f'eventos {item["ev_certidao"]}/{item["ev_despacho"]}):')
        out = gerar(item)

        # check (ponytail): o documento salvo tem o cabeçalho, os cinco parágrafos e a assinatura
        salvo = docx.Document(str(out))
        paragrafos = [p.text for p in salvo.paragraphs]
        texto = "\n".join(paragrafos)
        assert "{{" not in texto and "{%" not in texto, "placeholder não substituído"
        for trecho in (item["processo"], item["assunto"], item["relator"], *ASSINATURA):
            assert trecho in texto, f'{item["processo"]}: faltou "{trecho}"'
        # as quatro linhas da assinatura em espaço simples, sem espaço depois
        assinatura = [p for p in salvo.paragraphs if p.text.strip() in ASSINATURA]
        assert len(assinatura) == len(ASSINATURA), (
            f'{item["processo"]}: {len(assinatura)} linhas de assinatura, esperava '
            f'{len(ASSINATURA)}')
        for p in assinatura:
            pf = p.paragraph_format
            assert pf.line_spacing_rule == WD_LINE_SPACING.SINGLE and pf.space_after == Pt(0), (
                f'{item["processo"]}: "{p.text.strip()}" sem espaço simples')
        corpo = [p for p in paragrafos if p.strip() and p in ctx["parágrafos"]]
        assert corpo == ctx["parágrafos"], (
            f'{item["processo"]}: corpo saiu com {len(corpo)} dos 5 parágrafos, fora de ordem '
            f'ou alterado')
        # as referências que o texto faz aos autos têm de ser as apuradas
        assert f'evento {item["ev_despacho"]}' in ctx["parágrafos"][1]
        assert f'evento {item["ev_certidao"]}' in ctx["parágrafos"][3]
        assert f'parágrafo {ITEM_MEMORANDO} do despacho no evento {item["ev_despacho"]}' in (
            ctx["parágrafos"][3])
        assert f'item b do despacho no evento {item["ev_despacho"]}' in ctx["parágrafos"][4]
        # "cominatória" aparece se e somente se o débito vigente for do tipo 5
        assert ("multa cominatória" in ctx["parágrafos"][0]) == (item["tipo"] == TIPO_COMINATORIA)
        assert out.with_suffix(".pdf").is_file(), "PDF não gerado"
        print(f"  salvo: {out.relative_to(BASE)} (+ .pdf)")

    print(f"\n{len(itens)} informações geradas em {DESTINO.relative_to(BASE)}/")
