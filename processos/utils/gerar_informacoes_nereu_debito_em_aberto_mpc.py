"""Informações do lote GCAED/Nereu: débito de "suspenso" para "em aberto" + remessa ao MPC.

Ato Ordinatório do GCAED de 21/10/2026 (mesmo texto em todos): adequar a situação do débito de
"suspenso" para "em aberto" (Quota Ministerial) e remeter os autos ao MPC (Despacho anterior).
Mesmo texto da informação do 002073/2024 (processos/informacoes/002073_2024, versão editada pelo
usuário em 26/10/2026), trocando eventos, decisão e processo de origem. A alteração da situação
do débito é feita pela CCD antes de cadastrar; o script só avisa se o banco ainda mostra Suspenso.

Decisões conferidas nas informações v2 (processos/projetos/nereu_retomada_gcaed_v2/<n_a>/).
Modelo: scripts/automacao/templates/modelo_informacao.docx (numera os parágrafos).
Rodar: .venv/Scripts/python.exe processos/utils/gerar_informacoes_nereu_debito_em_aberto_mpc.py
"""
import re
import shutil
from copy import deepcopy
from datetime import datetime
from pathlib import Path

import docx
import pypdf
from docx.oxml.ns import qn
from docxtpl import DocxTemplate

from ccd.config import REPO_ROOT
from ccd.db import run_query_df
from ccd.docs import docx_to_pdf
from ccd.pdf import extract_text_from_pdf
from ccd.processo import get_info_file_path

BASE = Path(__file__).resolve().parents[1]
DESTINO = BASE / "projetos" / "nereu_debito_em_aberto_mpc"
TEMPLATE = str(REPO_ROOT / "scripts" / "automacao" / "templates" / "modelo_informacao.docx")
RELATOR = "Antônio Ed Souza Santana"  # Conselheiro titular
COD_RELATOR = 18

# processo -> (decisão, processo de origem)
PROCESSOS = {
    "002074/2024": ("Decisão nº 730/2020-TC", "006920/2016"),
    "002075/2024": ("Decisão nº 731/2020-TC", "006953/2016"),
    "002543/2024": ("Decisão nº 3182/2020-TC", "002468/2017"),
    "002544/2024": ("Decisão nº 2424/2020-TC", "014008/2016"),
    "002546/2024": ("Decisão nº 3073/2020-TC", "018094/2015"),
    "002563/2024": ("Decisão nº 701/2019-TC", "006249/2016"),
    "002566/2024": ("Decisão nº 3185/2020-TC", "018873/2016"),
    "002733/2022": ("Decisão nº 2153/2020-TC", "011078/2017"),
    "002734/2022": ("Decisão nº 733/2020-TC", "007059/2016"),
    "003008/2022": ("Decisão nº 728/2020-TC", "005726/2016"),
    "003053/2022": ("Decisão nº 2975/2020-TC", "013441/2017"),
}

MESES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho",
         "agosto", "setembro", "outubro", "novembro", "dezembro"]
hoje = datetime.now()
DATA = f"Natal/RN, {hoje.day} de {MESES[hoje.month - 1]} de {hoje.year}."


def conferir(processo: str, origem: str) -> dict:
    """Fatos do banco/PDF que entram no texto; assert em qualquer divergência."""
    n, a = processo.split("/")
    proc = run_query_df(
        "SELECT IdProcesso, RTRIM(setor_atual) setor, codigo_relator, RTRIM(assunto) assunto "
        "FROM Processos WHERE numero_processo = :n AND ano_processo = :a", n=n, a=int(a))
    assert proc.setor.iloc[0] == "CCD", f"{processo}: saiu da CCD ({proc.setor.iloc[0]!r})"
    assert int(proc.codigo_relator.iloc[0]) == COD_RELATOR, f"{processo}: relator mudou"

    ev = run_query_df(
        "SELECT ev.SequencialProcessoEvento evento, RTRIM(v.setor) setor, v.ordem, "
        "CONVERT(date, v.data_resumo) data "
        "FROM vw_ata_informacao v JOIN Pro_ProcessoEvento ev ON ev.idInformacao = v.idInformacao "
        "WHERE RTRIM(v.numero_processo) = :n AND RTRIM(v.ano_processo) = :a", n=n, a=a)
    ato = ev.sort_values("evento").iloc[-1]  # última peça tem de ser o ato do gabinete
    assert ato.setor == "GCAED", f"{processo}: última peça é do setor {ato.setor} (Evento {ato.evento})"
    assert str(ato.data) == "2026-10-21", f"{processo}: ato de {ato.data}"
    texto = re.sub(r"\s+", " ", extract_text_from_pdf(get_info_file_path(
        {"setor": "GCAED", "arquivo": f"GCAED_{n}_{a}_{int(ato.ordem):04d}.pdf"})))
    assert "suspenso" in texto and "em aberto" in texto, f"{processo}: ato sem a determinação"
    quota = int(re.search(r"Quota Ministerial de evento (\d+)", texto).group(1))
    despacho = int(re.search(r"Despacho de evento (\d+)", texto).group(1))
    setores = dict(zip(ev.evento, ev.setor, strict=True))
    assert setores.get(quota) == "PROC_EXE", f"{processo}: Evento {quota} não é a quota do MPC"
    assert setores.get(despacho) == "GCAED", f"{processo}: Evento {despacho} não é o despacho"

    deb = run_query_df(
        "SELECT d.IdDebito, d.CodigoStatusDivida st, "
        "RTRIM(o.numero_processo)+'/'+RTRIM(o.ano_processo) origem "
        "FROM Exe_Debito d JOIN Processos o ON o.IdProcesso = d.IdProcessoOrigem "
        "WHERE d.IdProcessoExecucao = :id AND d.DataCancelamento IS NULL "
        "AND NOT EXISTS (SELECT 1 FROM Exe_Debito f WHERE f.IdDebitoAnterior = d.IdDebito)",
        id=int(proc.IdProcesso.iloc[0]))
    assert len(deb) == 1, f"{processo}: débitos inesperados {deb.to_dict('records')}"
    assert deb.origem.iloc[0] == origem, f"{processo}: origem {deb.origem.iloc[0]} ≠ {origem}"
    if int(deb.st.iloc[0]) != 1:
        # ponytail: só aviso — o banco atrasa em relação ao e-Contas, que é o autoritativo
        print(f"  AVISO: débito {int(deb.IdDebito.iloc[0])} ainda com status {int(deb.st.iloc[0])} "
              "(20 = Suspenso); confirmar 'Em Aberto' antes de cadastrar.")
    return {"assunto": proc.assunto.iloc[0], "ato": int(ato.evento), "quota": quota, "despacho": despacho}


def paragrafos(c: dict) -> list[str]:
    return [
        f"Por Ato Ordinatório de 21 de outubro de 2026 (Evento {c['ato']}), o Gabinete do Conselheiro "
        "Relator encaminhou os autos a esta Diretoria para a adequação da situação cadastral do "
        "débito imputado ao Sr. Nereu Batista Linhares, de “suspenso” para “em aberto”, nos termos "
        f"da Quota Ministerial do Evento {c['quota']}, com posterior remessa dos autos ao Ministério "
        f"Público de Contas, em cumprimento ao Despacho do Evento {c['despacho']}.",

        "Informa-se que a providência foi realizada: a situação cadastral do débito decorrente da "
        f"{c['decisao']} (Processo nº {c['origem']}-TC) foi alterada de “suspenso” para "
        "“em aberto” no sistema desta Corte, o que viabiliza o seu encaminhamento à "
        "Procuradoria-Geral do Estado por meio do SICODE.",

        "Ante o exposto, remetem-se os autos ao Ministério Público de Contas, conforme determinado "
        f"no Evento {c['ato']}.",
    ]


def texto_xml(el) -> str:
    return "".join(t.text or "" for t in el.findall(f".//{qn('w:t')}")).strip()


def gerar(processo: str, c: dict) -> Path:
    out = DESTINO / processo.replace("/", "_") / f"informacao_{processo.replace('/', '_')}.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    ps = paragrafos(c)
    doc = DocxTemplate(TEMPLATE)
    doc.render({"processo": f"{processo} - TC", "assunto": c["assunto"], "relator": RELATOR,
                "parágrafos": ps})  # sem numerar: o modelo numera
    d = doc.docx  # pós-processamento (get_docx() descartaria o render)

    # o modelo deixa duas linhas em branco após o título; o espaçamento do parágrafo basta
    titulo = next(p for p in d.paragraphs if p.text == "INFORMAÇÃO INSTRUTIVA")
    for _ in range(2):
        nxt = titulo._p.getnext()
        if texto_xml(nxt):
            break
        nxt.getparent().remove(nxt)

    # data antes do bloco de assinatura; parágrafos vazios antes dela saem (cabe em 1 página)
    assinado = next(p for p in d.paragraphs if "assinado digitalmente" in p.text)
    prev = assinado._p.getprevious()
    while prev is not None and not texto_xml(prev):
        vazio, prev = prev, prev.getprevious()
        vazio.getparent().remove(vazio)
    novo = deepcopy(assinado._p)
    for r in list(novo):
        if r.tag == qn("w:r"):
            novo.remove(r)
    assinado._p.addprevious(novo)
    r = deepcopy(assinado.runs[0]._r)
    novo.append(r)
    for t in r.findall(qn("w:t")):
        t.text = DATA
    idx_data = next(i for i, p in enumerate(d.paragraphs) if p.text.startswith("Natal/RN"))
    for p in d.paragraphs[idx_data:-1]:
        p.paragraph_format.keep_with_next = True

    # preserva a versão anterior antes de sobrescrever (skill edicao-minima)
    if out.exists():
        bak = out.with_name(f"{out.stem}_{datetime.now():%Y%m%d_%H%M%S}{out.suffix}")
        shutil.copy2(out, bak)
        print(f"  versão anterior preservada em: {bak.name}")
    doc.save(str(out))
    docx_to_pdf(str(out), str(out.parent))

    # check (ponytail): texto completo, sem placeholder, sem termo proibido, 1 página
    texto = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    assert "{{" not in texto and "}}" not in texto, f"{processo}: placeholder não substituído"
    for trecho in (*ps, DATA, f"{processo} - TC", RELATOR):
        assert trecho in texto, f"{processo}: faltou {trecho[:50]}"
    for proibido in ("É a informação.", "Substituto", "Exe_Debito", "vw_ata"):
        assert proibido not in texto, f"{processo}: termo proibido {proibido}"
    paginas = len(pypdf.PdfReader(str(out.with_suffix(".pdf"))).pages)
    assert paginas == 1, f"{processo}: {paginas} páginas (tem de caber em 1)"
    return out


if __name__ == "__main__":
    for processo, (decisao, origem) in PROCESSOS.items():
        print(processo, flush=True)
        c = {**conferir(processo, origem), "decisao": decisao, "origem": origem}
        print(f"  Ato {c['ato']}, Quota {c['quota']}, Despacho {c['despacho']} -> {gerar(processo, c)}")
