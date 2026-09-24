"""Informações CCD — 000130/2023 e 000107/2023: resposta da SEAD não confirmada.

A SEAD respondeu às Notificações nº 001981/2024-DAE (130/2023) e nº 001976/2024-DAE
(107/2023) dizendo que somou as duas multas do Sr. Nereu Batista Linhares (R$ 1.281,37
cada = R$ 2.562,74) e implantou o desconto no contracheque de novembro/2025 (tela ERGON:
rubrica 644 FRAP TC, 1 parcela, vínculo 3 – IPERN). Nem o SIAI Pessoal mostra a retenção
em 11/2025, nem o extrato do FRAP mostra o crédito. Uma informação por processo com:
resumo, resposta do órgão, consultas ao SIAI Pessoal e ao FRAP (dois quadros) e sugestão
de notificação do órgão; a do 107/2023 contextualiza o sobrestamento (MS 0807247-93) e a
do 130/2023 registra que não é alcançada pelo MS nem teve o débito suspenso.
Texto uniformizado em 22/09/2026 a partir das edições manuais dos dois .docx (versão
enxuta: sem SEI do anexo, alcance do SIAI/extrato, "sob pena" e "É a informação").

Modelo: scripts/automacao/templates/modelo_informacao.docx (quadros no padrão de
processos/projetos/verificacao_frap/gerar_informacoes.py). Saída: processos/projetos/nereu_resposta_sead/.

Rodar: .venv/Scripts/python.exe processos/utils/gerar_informacoes_nereu_resposta_sead.py
"""
import shutil
import sys
from copy import deepcopy
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor
from docxtpl import DocxTemplate

from ccd.config import REPO_ROOT, cpf
from ccd.db import get_connection, run_query_df
from ccd.docs import docx_to_pdf

TEMPLATE = str(REPO_ROOT / "scripts" / "automacao" / "templates" / "modelo_informacao.docx")
DESTINO = Path(__file__).resolve().parents[1] / "projetos" / "nereu_resposta_sead"
CPF = cpf("NEREU")
VALOR_NOTIFICADO = 1281.37
VALOR_SOMADO = 2562.74
COMPETENCIA = (2025, 11)

MESES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho",
         "agosto", "setembro", "outubro", "novembro", "dezembro"]
hoje = datetime.now()
DATA = f"Natal/RN, {hoje.day} de {MESES[hoje.month - 1]} de {hoje.year}."

# Fatos conferidos nos PDFs do share (ofícios da SEAD, despachos da COPAG, tela ERGON)
# e no banco processo em 17/09/2026. "ordens" = vw_ata_informacao.ordem; o nº do
# Evento exibido vem de Pro_ProcessoEvento.SequencialProcessoEvento (consultado ao vivo).
CASOS = {
    "000130/2023": dict(
        origem="012540/2017", decisao="2614/2020", acordao="141/2022",
        notificacao="001981/2024", par="000107/2023",
        apenso="300046/2025", oficio="10364/2024", sei_processo="00110012.003264/2024-60",
        sei_despacho="31107685",
        sobrestado=False,
        ordens=dict(despacho=65, notificacao=67, certidao=68, apensamento=71, relator=74, quota=75),
        eventos_esperados=dict(despacho=75, notificacao=78, certidao=80, apensamento=84,
                               relator=90, quota=93),
    ),
    "000107/2023": dict(
        origem="006602/2016", decisao="856/2020", acordao="78/2022",
        notificacao="001976/2024", par="000130/2023",
        apenso="300045/2025", oficio="10365/2024", sei_processo="00110012.003263/2024-15",
        sei_despacho="31107262",
        sobrestado=True,
        ordens=dict(despacho=71, notificacao=73, certidao=74, apensamento=75, relator=78,
                    suspensao=79, sugestao=80, decisao_sobr=81, ccd_2026=82, mantem=83),
        eventos_esperados=dict(despacho=81, notificacao=84, certidao=86, apensamento=87, relator=93,
                               suspensao=96, sugestao=97, decisao_sobr=100, mantem=106),
    ),
}

SQL_BASE = """
SELECT RTRIM(p.assunto) AS assunto, RTRIM(r.nome) AS relator, RTRIM(p.setor_atual) AS setor,
       d.IdDebito, d.valorOriginalDebito AS valor_original,
       processo.dbo.fn_Exe_RetornaValorAtualizado(d.IdDebito) AS valor,
       RTRIM(sd.DescricaoStatusDivida) AS situacao, RTRIM(gp.Nome) AS nome, gp.Documento AS documento
FROM dbo.Processos p
LEFT JOIN dbo.Relator r ON r.codigo = p.codigo_relator
JOIN dbo.Exe_Debito d ON d.IdProcessoExecucao = p.IdProcesso
     AND d.DataCancelamento IS NULL AND d.CodigoTipoDebito = 2
JOIN dbo.Exe_StatusDivida sd ON sd.CodigoStatusDivida = d.CodigoStatusDivida
JOIN dbo.Exe_DebitoPessoa dp ON dp.IDDebito = d.IdDebito
JOIN dbo.GenPessoa gp ON gp.IdPessoa = dp.IDPessoa
WHERE p.numero_processo = :num AND p.ano_processo = :ano
"""

# vw_ata_informacao.IdProcesso é NULL — filtrar por numero/ano.
SQL_EVENTOS = """
SELECT v.ordem, ev.SequencialProcessoEvento AS evento
FROM dbo.vw_ata_informacao v
JOIN dbo.Pro_ProcessoEvento ev ON ev.IdInformacao = v.idInformacao
WHERE v.numero_processo = :num AND v.ano_processo = :ano
"""

# SIAI Pessoal pelas tabelas-base (web/backend/app/ccd/siai_pessoal/service.py): um
# contracheque por (ano, mês, órgão, tipo de folha), última remessa; competência da FOLHA.
CTE_SIAI = """
WITH cc AS (
    SELECT CC.IdContraCheque, F.IdTipoFolhaPagamento, ER.IdOrgao,
           TRY_CAST(LTRIM(RTRIM(F.Ano)) AS INT) AS ano,
           TRY_CAST(LTRIM(RTRIM(F.Mes)) AS INT) AS mes,
           FU.Matricula,
           ROW_NUMBER() OVER (
               PARTITION BY TRY_CAST(LTRIM(RTRIM(F.Ano)) AS INT),
                            TRY_CAST(LTRIM(RTRIM(F.Mes)) AS INT),
                            ER.IdOrgao, F.IdTipoFolhaPagamento
               ORDER BY ER.IdEnvioRemessa DESC, CC.IdContraCheque DESC) AS rn
    FROM BDSiaiPessoal.dbo.Comum_Pessoa P
    JOIN BDSiaiPessoal.dbo.SiaiDp_Funcionario FU ON FU.IdPessoa = P.IdPessoa
    JOIN BDSiaiPessoal.dbo.SiaiDp_ContraCheque CC ON CC.IdFuncionario = FU.IdFuncionario
    JOIN BDSiaiPessoal.dbo.SiaiDp_FolhaPagamento F ON F.IdFolhaPagamento = CC.IdFolhaPagamento
    JOIN BDSiaiPessoal.dbo.SiaiDp_RemessaFolhaPagamento RFP
      ON RFP.IdRemessaFolhaPagamento = F.IdRemessaFolhaPagamento
    JOIN BDSiaiPessoal.dbo.Envio_Remessa ER ON ER.IdEnvioRemessa = RFP.IdEnvioRemessa
    WHERE P.CPF = :cpf
      AND ER.IdTipoEnvioRemessa = 1
      AND ER.IdSituacaoRemessa IN (4, 7)
      AND ER.IdEnvioRemessaVinculada IS NOT NULL
      AND TRY_CAST(LTRIM(RTRIM(F.Ano)) AS INT) >= :ano_min
)
"""
SQL_SIAI_ITENS = CTE_SIAI + """
SELECT cc.ano, cc.mes, cc.IdOrgao AS id_orgao, RTRIM(O.NomeOrgao) AS orgao,
       RTRIM(cc.Matricula) AS matricula, RTRIM(R.Codigo) AS codigo,
       RTRIM(R.Descricao) AS rubrica, R.Tipo AS tipo, CAST(CCI.Valor AS DECIMAL(18, 2)) AS valor
FROM cc
LEFT JOIN Bdc.dbo.vw_Gen_Orgao O ON O.IdOrgao = cc.IdOrgao
JOIN BDSiaiPessoal.dbo.SiaiDp_ContraChequeItem CCI ON CCI.IdContraCheque = cc.IdContraCheque
JOIN BDSiaiPessoal.dbo.SiaiDp_Rubrica R ON R.IdRubrica = CCI.IdRubrica
WHERE cc.rn = 1
"""

SQL_FRAP_EXATOS = """
SELECT L.IdLancamento, L.DtMovimento, L.Documento, L.Valor
FROM dbo.FRAPLancamento L
WHERE L.ValorDC = 'C' AND L.IdCategoria IN (1, 3) AND L.DtMovimento >= :desde
  AND (ABS(L.Valor - :v1) < 0.005 OR ABS(L.Valor - :v2) < 0.005)
"""
SQL_FRAP_CREDITOS = """
SELECT L.IdLancamento, L.DtMovimento, L.Documento, L.Valor, RTRIM(c.Conta) AS conta
FROM dbo.FRAPLancamento L JOIN dbo.FRAPConta c ON c.IdConta = L.IdConta
WHERE L.ValorDC = 'C' AND L.IdCategoria IN (1, 3) AND L.DtMovimento >= :desde
ORDER BY L.DtMovimento
"""


def brl(v: float) -> str:
    return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def fmt_cpf(c: str) -> str:
    return f"{c[:3]}.{c[3:6]}.{c[6:9]}-{c[9:]}"


def comp(ano: int, mes: int) -> str:
    return f"{MESES[mes - 1]}/{ano}"


def rubrica_tce(tipo, nome) -> bool:
    n = str(nome or "").upper()
    return str(tipo or "").strip() == "2" and any(k in n for k in ("TCE", "FRAP", "TRIBUNAL DE CONTAS"))


# --- identidade visual TCE/RN (ref. processos/projetos/verificacao_frap/gerar_informacoes.py) ---
TCE = {"verde_principal": "2E5B3C", "verde_escuro": "1A3D28",
       "cinza_claro": "F2F2F2", "cinza_borda": "CCCCCC", "cinza_texto": "333333"}


def _shade(cell, hex_cor: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_cor)
    cell._tc.get_or_add_tcPr().append(shd)


def _tabela_apos(d, ancora_texto: str, legenda: str, cabecalho: list[str],
                 linhas: list[list[str]]) -> None:
    """Insere legenda "Quadro N – título" + tabela após o parágrafo-âncora."""
    ancora = next(p for p in d.paragraphs if ancora_texto in p.text)
    cap = d.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_cap = cap.add_run(legenda)
    run_cap.bold = True
    run_cap.font.color.rgb = RGBColor.from_string(TCE["verde_escuro"])
    ancora._p.addnext(cap._p)
    t = d.add_table(rows=0, cols=len(cabecalho))
    borders = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"),
               TCE["cinza_borda"] if lado.startswith("inside") else TCE["verde_principal"])
        borders.append(el)
    t._tbl.tblPr.append(borders)
    for i, linha in enumerate([cabecalho] + linhas):
        total = linha[0].startswith("Total")
        cells = t.add_row().cells
        for cell, texto in zip(cells, linha, strict=True):
            if i == 0:
                _shade(cell, TCE["verde_principal"])
            elif total:
                _shade(cell, TCE["verde_escuro"])
            elif i % 2 == 1:
                _shade(cell, TCE["cinza_claro"])
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = par.add_run(texto)
            run.bold = i == 0 or total
            run.font.color.rgb = RGBColor.from_string(
                "FFFFFF" if (i == 0 or total) else TCE["cinza_texto"])
    cap._p.addnext(t._tbl)


def carregar_conciliacao() -> dict:
    """SIAI Pessoal (contracheque 11/2025 + retenções TCE/FRAP) e FRAP, uma vez para o CPF."""
    eng = get_connection("BdDIP")
    itens = run_query_df(SQL_SIAI_ITENS, eng, cpf=CPF, ano_min=2024)
    assert not itens.empty, "SIAI Pessoal sem folhas do CPF"
    itens["valor"] = itens.valor.astype(float)
    ano, mes = COMPETENCIA
    contracheque = itens[(itens.ano == ano) & (itens.mes == mes)].sort_values(
        ["id_orgao", "tipo", "codigo"])
    assert not contracheque.empty, f"sem folha {mes:02d}/{ano} no SIAI"
    assert not contracheque.apply(lambda r: rubrica_tce(r.tipo, r.rubrica), axis=1).any(), (
        f"há rubrica TCE/FRAP em {mes:02d}/{ano} — a premissa da informação caiu")
    ret = itens[itens.apply(lambda r: rubrica_tce(r.tipo, r.rubrica), axis=1)]
    ret = ret.groupby(["ano", "mes", "orgao", "matricula", "codigo", "rubrica"], as_index=False).valor.sum()
    ret = ret.sort_values(["ano", "mes"])

    desde = date(2024, 11, 1)
    exatos = run_query_df(SQL_FRAP_EXATOS, eng, desde=desde, v1=VALOR_SOMADO, v2=VALOR_NOTIFICADO)
    assert exatos.empty, f"há crédito no FRAP com o valor informado: {exatos.to_dict('records')}"
    creditos = run_query_df(SQL_FRAP_CREDITOS, eng, desde=desde)
    creditos["Valor"] = creditos.Valor.astype(float)
    creditos["DtMovimento"] = pd.to_datetime(creditos.DtMovimento)

    # padrão IPERN: retenção em M -> OB de valor idêntico no fim de M ou em M+1 (até M+2);
    # cada OB casa uma vez
    usados: set[int] = set()
    casadas = []
    for r in ret.itertuples():
        ini = pd.Timestamp(int(r.ano), int(r.mes), 1)
        fim = ini + pd.offsets.MonthEnd(3)
        cand = creditos[(creditos.DtMovimento >= ini) & (creditos.DtMovimento <= fim)
                        & ((creditos.Valor - r.valor).abs() < 0.005)
                        & ~creditos.IdLancamento.isin(usados)]
        ob = cand.iloc[0] if not cand.empty else None
        if ob is not None:
            usados.add(int(ob.IdLancamento))
        casadas.append((r, ob))
    return dict(contracheque=contracheque, retencoes=casadas)


def carregar_processo(proc: str) -> dict:
    num, ano = proc.split("/")
    base = run_query_df(SQL_BASE, None, num=num, ano=ano)
    assert len(base) == 1, f"{proc}: {len(base)} débitos de multa vigentes"
    b = base.iloc[0]
    assert str(b.documento).strip().zfill(11) == CPF, f"{proc}: responsável não é o Nereu"
    ev = run_query_df(SQL_EVENTOS, None, num=num, ano=ano).set_index("ordem").evento.to_dict()
    caso = CASOS[proc]
    eventos = {k: int(ev[o]) for k, o in caso["ordens"].items()}
    for k, esperado in caso["eventos_esperados"].items():
        assert eventos[k] == esperado, f"{proc}: evento {k} = {eventos[k]}, esperado {esperado}"
    return dict(processo=proc, assunto=b.assunto, relator=b.relator, setor=b.setor,
                id_debito=int(b.IdDebito), valor_original=float(b.valor_original),
                valor=float(b.valor), situacao=b.situacao, nome=b.nome, eventos=eventos, **caso)


def montar_paragrafos(it: dict, conc: dict) -> tuple[list[str], list[tuple]]:
    ev = it["eventos"]
    n_par = it["par"]
    ano_c, mes_c = COMPETENCIA
    comp_txt = f"{MESES[mes_c - 1]} de {ano_c}"
    sobrestado = it["sobrestado"]

    pars = [
        f'Trata-se de execução de multa imputada ao Sr. {it["nome"].title()} (CPF nº {fmt_cpf(CPF)}), '
        f'então Presidente do IPERN, pelo Acórdão nº {it["acordao"]}-TC (Processo nº {it["origem"]}-TC), '
        f'por descumprimento da Decisão nº {it["decisao"]}-TC (débito nº {it["id_debito"]}, valor '
        f'atualizado de R$ {brl(it["valor"])}). Diante da inércia do responsável, a Diretoria de Atos e '
        f'Execuções – DAE determinou o desconto em folha de R$ {brl(VALOR_NOTIFICADO)} (Evento '
        f'{ev["despacho"]}), expedindo-se a Notificação nº {it["notificacao"]}-DAE à SEARH (atual SEAD) em '
        f'30/10/2024 (Eventos {ev["notificacao"]} e {ev["certidao"]}).'
    ]
    if sobrestado:
        pars[0] += (
            " Respondida a notificação, os autos vieram à DIP para acompanhamento da execução "
            f'(Evento {ev["relator"]}).'
        )
    else:
        pars[0] += (
            " Respondida a notificação, o Ministério Público de Contas sugeriu o acompanhamento dos "
            f'descontos pela DIP (Eventos {ev["relator"]} e {ev["quota"]}).'
        )

    pars.append(
        f'A resposta da SEAD (Ofício nº {it["oficio"]}/SEAD, SEI nº {it["sei_processo"]}) foi autuada '
        f'como Processo nº {it["apenso"]}-TC e apensada a estes autos (Evento {ev["apensamento"]}). '
        f'Segundo Despacho da COPAG/SEAD de 26/12/2024 (SEI nº {it["sei_despacho"]}), esta multa foi '
        f'somada à do Processo nº {n_par}-TC, "totalizando um montante de '
        f'R$ {brl(VALOR_SOMADO)}, desconto este que fora implantado na integralidade através do '
        f'contracheque do mês de {MESES[mes_c - 1]}/{ano_c}". O Anexo II '
        f"reproduz tela do ERGON com a consignação: rubrica 644 – FRAP TC, 1 parcela em {mes_c:02d}/{ano_c}, "
        f"R$ {brl(VALOR_SOMADO)}, funcionário 1732889, vínculo 3 – PR/IPERN. O despacho ressalva que o "
        'repasse é "de inteira responsabilidade do órgão de lotação", o IPERN.'
    )

    # --- SIAI Pessoal ---
    cc = conc["contracheque"]
    orgaos = cc.drop_duplicates(["id_orgao", "matricula"])
    desc_orgaos = " e ".join(
        f"o da matrícula {o.matricula} no órgão {o.orgao}" for o in orgaos.itertuples())
    pars.append(
        "Para verificar a efetivação do desconto informado, esta Coordenadoria consultou as folhas de "
        "pagamento remetidas ao SIAI Pessoal vinculadas ao CPF do servidor, considerando a última "
        f"remessa de cada competência. Em {comp_txt} constam dois contracheques: {desc_orgaos} "
        "– a matrícula 17328893 corresponde ao funcionário 1732889, vínculo 3, indicado na tela do "
        "ERGON. Em nenhum deles há a rubrica 644 – FRAP TC ou qualquer outro desconto em favor deste "
        "Tribunal, conforme o quadro a seguir:"
    )
    linhas1 = [[f"{r.orgao} (matrícula {r.matricula})", f"{r.codigo} – {r.rubrica}",
                "Desconto" if str(r.tipo).strip() == "2" else "Vantagem", brl(r.valor)]
               for r in cc.itertuples()]
    quadros = [("conforme o quadro a seguir",
                f"Quadro 1 – Contracheque de {comp_txt} do servidor no SIAI Pessoal",
                ["Órgão / matrícula", "Rubrica", "Tipo", "Valor (R$)"], linhas1)]

    ret = conc["retencoes"]
    pars.append(
        "A rubrica 644 – FRAP TC aparece nas folhas do servidor em outras competências – de dezembro "
        "de 2024 a maio de 2025 (IPERN) e a partir de abril de 2026 (Procuradoria-Geral do Estado) –, "
        f"sempre em valores distintos de R$ {brl(VALOR_SOMADO)} e de R$ {brl(VALOR_NOTIFICADO)}, "
        "relacionados a outros processos de execução em curso contra o mesmo responsável."
    )

    # --- FRAP ---
    pars.append(
        "No extrato bancário das contas do FRAP no Banco do Brasil não há crédito de "
        f"R$ {brl(VALOR_SOMADO)} nem de R$ {brl(VALOR_NOTIFICADO)}. Observa-se "
        "que os repasses do IPERN relativos a este servidor seguem padrão regular: cada retenção lançada "
        "na folha de uma competência ingressa na conta nº 700000-6 como Ordem Bancária de valor idêntico "
        f"no mês seguinte, como demonstra o quadro abaixo. Para a competência de {comp_txt} não há "
        "retenção nem crédito correspondente. As retenções lançadas pela Procuradoria-Geral do "
        "Estado a partir de abril de 2026, relativas a outros processos, têm um único crédito "
        "identificado."
    )
    linha_nov = [comp(*COMPETENCIA), "IPERN (informado pela SEAD)",
                 f"não localizada (informado: {brl(VALOR_SOMADO)})", "—"]
    linhas2 = []
    inserida = False
    for r, ob in ret:
        chave = (int(r.ano), int(r.mes))
        if not inserida and chave > COMPETENCIA:
            linhas2.append(linha_nov)
            inserida = True
        cred = (f"{ob.DtMovimento:%d/%m/%Y} – OB {ob.Documento} – R$ {brl(ob.Valor)}"
                if ob is not None else "—")
        linhas2.append([comp(*chave), r.orgao, brl(r.valor), cred])
    if not inserida:
        linhas2.append(linha_nov)
    quadros.append(("como demonstra o quadro abaixo",
                    "Quadro 2 – Retenções da rubrica FRAP TC no SIAI Pessoal e créditos "
                    "correspondentes na conta do FRAP",
                    ["Competência", "Órgão", "Retenção no SIAI (R$)", "Crédito na conta do FRAP"],
                    linhas2))

    if sobrestado:
        pars.append(
            "Cumpre contextualizar que este processo está sobrestado. Em 02/10/2025 esta Coordenadoria "
            f'registrou a suspensão da exigibilidade da multa (Evento {ev["suspensao"]}), em razão da liminar '
            "deferida no Mandado de Segurança nº 0807247-93.2025.8.20.0000, impetrado pelo responsável "
            "perante o Tribunal de Justiça do Estado, que determinou a sustação da imposição de multas "
            "decorrentes da matéria discutida no Processo Judicial nº 0811707-97.2021.8.20.5001; e, "
            f'acolhendo sugestão desta unidade (Evento {ev["sugestao"]}), o Conselheiro Relator determinou o '
            f'sobrestamento do feito (Evento {ev["decisao_sobr"]}).'
        )
        pars.append(
            "Em 23/01/2026 o Pleno do TJ/RN denegou a segurança e revogou a liminar, mas, diante dos "
            "embargos de declaração opostos pelo impetrante e dos recursos especial e extraordinário "
            "interpostos na Apelação Cível nº 0811707-97.2021.8.20.5001, o Relator manteve o "
            f'sobrestamento em 23/06/2026 (Evento {ev["mantem"]}), determinando o acompanhamento do feito '
            "por esta Coordenadoria."
        )
        pars.append(
            "A resposta da SEAD, contudo, é anterior ao sobrestamento e é comum "
            f"aos Processos nº {it['processo']}-TC e nº {n_par}-TC, de modo que a apuração da veracidade "
            "do desconto informado não implica retomada da execução: se houve retenção em "
            f"{comp_txt} – já sob a vigência da liminar –, o valor retido precisa ser identificado e "
            "destinado; se não houve, a informação prestada ao Tribunal não corresponde aos fatos."
        )
    else:
        pars.append(
            "Registre-se que este processo não é alcançado pelo Mandado de Segurança "
            "nº 0807247-93.2025.8.20.0000, impetrado pelo responsável perante o Tribunal de "
            "Justiça do Estado, nem teve suspensa a exigibilidade do débito, prosseguindo "
            "regularmente a execução."
        )

    pars.append(
        f"Diante do exposto, o desconto de R$ {brl(VALOR_SOMADO)} informado pela SEAD não está "
        "comprovado: não há retenção na folha de pagamento do servidor na competência indicada, não há "
        "crédito correspondente na conta do FRAP e não foi juntado aos autos o comprovante de crédito "
        "exigido pelo art. 25, §§ 3º e 4º, da Resolução nº 013/2015-TCE."
    )
    ressalva = ", sem prejuízo do sobrestamento do feito," if sobrestado else ""
    pars.append(
        f"Assim, sugere-se ao Conselheiro Relator{ressalva} a notificação da Secretaria de Estado da "
        "Administração – SEAD, por sua Coordenadoria da Folha de Pagamento, com ciência ao IPERN, para "
        f"que, no prazo de 15 (quinze) dias, esclareça se o desconto de R$ {brl(VALOR_SOMADO)} foi "
        f"efetivamente realizado no contracheque de {comp_txt}, juntando o contracheque e o comprovante "
        "de crédito à conta do FRAP, ou, em caso negativo, informe as razões da não efetivação e a data "
        "em que o desconto será implantado."
    )
    return pars, quadros


def gerar(it: dict, conc: dict) -> Path:
    num, ano = it["processo"].split("/")
    pasta = DESTINO / f"{num}_{ano}"
    pasta.mkdir(parents=True, exist_ok=True)
    out = pasta / f"informacao_{num}_{ano}.docx"

    pars, quadros = montar_paragrafos(it, conc)
    doc = DocxTemplate(TEMPLATE)
    doc.render({"processo": f'{it["processo"]} - TC', "assunto": it["assunto"],
                "relator": it["relator"], "parágrafos": pars})
    d = doc.docx  # pós-processamento (get_docx() descartaria o render)
    for ancora, legenda, cab, linhas in quadros:
        _tabela_apos(d, ancora, legenda, cab, linhas)

    # data antes do bloco de assinatura (o modelo não traz)
    assinado = next(p for p in d.paragraphs if "assinado digitalmente" in p.text)
    for texto in (DATA, ""):
        novo = deepcopy(assinado._p)
        for r in list(novo):
            if r.tag == qn("w:r"):
                novo.remove(r)
        assinado._p.addprevious(novo)
        if texto:
            r = deepcopy(assinado.runs[0]._r)
            novo.append(r)
            for t in r.findall(qn("w:t")):
                t.text = texto

    if out.exists():  # preserva a versão anterior (skill edicao-minima)
        bak = out.with_name(f"{out.stem}_{datetime.now():%Y%m%d_%H%M%S}{out.suffix}")
        shutil.copy2(out, bak)
        print(f"  versão anterior preservada em: {bak.name}")
    doc.save(str(out))
    docx_to_pdf(str(out), str(pasta))
    return out


if __name__ == "__main__":
    import docx

    conc = carregar_conciliacao()
    print(f"retenções FRAP TC: {len(conc['retencoes'])}, casadas com OB: "
          f"{sum(ob is not None for _, ob in conc['retencoes'])}")
    for proc in (sys.argv[1:] or CASOS):  # opcional: só os processos passados na linha de comando
        it = carregar_processo(proc)
        out = gerar(it, conc)
        d = docx.Document(str(out))
        texto = "\n".join([p.text for p in d.paragraphs]
                          + [c.text for t in d.tables for r in t.rows for c in r.cells])
        assert "{{" not in texto and "}}" not in texto, "placeholder solto"
        for s in (brl(VALOR_SOMADO), brl(VALOR_NOTIFICADO), "novembro/2025", it["apenso"],
                  it["notificacao"], f'Evento {it["eventos"]["despacho"]}',
                  f'Eventos {it["eventos"]["notificacao"]} e', DATA,
                  "têm um único crédito identificado."):
            assert s in texto, f"{proc}: falta '{s}'"
        for s in ("sob pena", "de igual valor", "É a informação", "SEI nº 31107703", "SEI nº 31107535",
                  "desde novembro de 2024", "alcançam a competência"):
            assert s not in texto, f"{proc}: sobrou '{s}'"
        assert ("sobrestado" in texto) == it["sobrestado"], f"{proc}: sobrestamento"
        assert len(d.tables) == 2, f"{proc}: {len(d.tables)} quadros"
        assert "644" not in "\n".join(c.text for r in d.tables[0].rows for c in r.cells)
        print(f"{proc}: {out.name} | relator {it['relator']} | setor {it['setor']} | "
              f"débito {it['id_debito']} R$ {brl(it['valor'])} ({it['situacao']}) | eventos {it['eventos']}")
