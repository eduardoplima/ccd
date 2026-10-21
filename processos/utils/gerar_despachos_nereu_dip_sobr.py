"""Despacho que substitui a peça DIP_SOBR (envio à DAP, 06/2026) nos 57 processos do lote GCAED (marcador 6139).

Um parágrafo só: tramitação à CCD para análise de fato superveniente. Modelo da informação instrutiva
(scripts/automacao/templates/modelo_informacao.docx) com o título trocado para "DESPACHO".
Imprime a ordem da peça DIP_SOBR a substituir.

Rodar: .venv/Scripts/python.exe processos/utils/gerar_despachos_nereu_dip_sobr.py
"""
import shutil
import sys
from datetime import datetime
from pathlib import Path

import docx
from docxtpl import DocxTemplate

from ccd.db import run_query_df
from ccd.docs import docx_to_pdf

sys.path.insert(0, str(Path(__file__).parent))
from gerar_informacoes_nereu_retomada_gcaed import (  # noqa: E402
    BASE,
    ESPERADOS,
    MARCADOR,
)

MODELO = BASE.parent / "scripts" / "automacao" / "templates" / "modelo_informacao.docx"
DESTINO = BASE / "projetos" / "nereu_retomada_gcaed_dip_sobr"
RELATOR = "Conselheiro ANTÔNIO ED SOUZA SANTANA"
PARAGRAFO = ("Em razão de fato superveniente relativo ao Mandado de Segurança Cível nº 0807247-93.2025.8.20.0000, "
             "que fundamentou o sobrestamento do feito — a suspensão, pelo Supremo Tribunal Federal, nos autos da "
             "Suspensão de Segurança nº 5.749, dos efeitos do acórdão que concedeu a segurança, com trânsito em "
             "julgado em 27/08/2026 —, tramitem-se os presentes autos à Coordenadoria de Controle de Decisões "
             "(CCD) para análise.")

SQL_DIP_SOBR = """
    SELECT RTRIM(v.numero_processo)+'/'+RTRIM(v.ano_processo) AS processo, v.ordem,
           CAST(v.data_resumo AS DATE) AS data, RTRIM(p.assunto) AS assunto
    FROM vw_ata_informacao v
    JOIN Processos p ON p.numero_processo = v.numero_processo AND p.ano_processo = v.ano_processo
    JOIN Pro_MarcadorProcesso mp ON mp.IdProcesso = p.IdProcesso
    WHERE mp.IdMarcador = :marcador  -- marcador excluído dos 57 em 09/2026: não filtrar DataExclusao
      AND RTRIM(v.setor) = 'DIP_SOBR' AND v.Inativa IS NULL
"""

if __name__ == "__main__":
    pecas = run_query_df(SQL_DIP_SOBR, marcador=MARCADOR).set_index("processo")
    processos = sorted(pecas.index, key=lambda p: p[-4:] + p)
    assert len(processos) == ESPERADOS and pecas.index.is_unique, "esperava 1 peça DIP_SOBR por processo"

    for p in processos:
        numero, ano = p.split("/")
        pasta = DESTINO / f"{numero}_{ano}"
        pasta.mkdir(parents=True, exist_ok=True)
        out = pasta / f"{numero}_{ano}.docx"
        if out.exists():  # preserva a versão anterior (skill edicao-minima)
            shutil.copy2(out, out.with_name(f"{out.stem}_{datetime.now():%Y%m%d_%H%M%S}{out.suffix}"))
        doc = DocxTemplate(str(MODELO))
        doc.render({"processo": f"{p}-TC", "assunto": pecas.loc[p].assunto, "relator": RELATOR,
                    "parágrafos": [PARAGRAFO]})
        titulo = next(par for par in doc.docx.paragraphs if par.text == "INFORMAÇÃO INSTRUTIVA")
        titulo.runs[0].text = "DESPACHO"  # doc.docx, não get_docx(): este descartaria o render
        doc.save(str(out))
        docx_to_pdf(str(out), str(pasta))

        texto = "\n".join(par.text for par in docx.Document(str(out)).paragraphs)
        assert "{{" not in texto and PARAGRAFO in texto and f"{p}-TC" in texto and "\nDESPACHO\n" in texto, p
        assert out.with_suffix(".pdf").is_file(), f"{p}: PDF não gerado"
        peca = pecas.loc[p]
        print(f"{p}: substitui DIP_SOBR ordem {int(peca.ordem)} ({peca.data:%d/%m/%Y})")

    print(f"\n{len(processos)} despachos em {DESTINO.relative_to(BASE)}/")
