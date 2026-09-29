"""Informações de retomada — lote GCAED (marcador 6139), versão 2: modelo de INFORMAÇÃO (parágrafos
numerados), histórico do MS 0807247-93 até a SS 5.749 no STF, teto da multa diária e excesso de
notificações de desconto em folha / envio à dívida ativa.

Mesmos 57 processos e mesmas checagens do v1 (gerar_informacoes_nereu_retomada_gcaed.py). Por processo:
- decisão que fixou a multa diária "limitad[a] ao teto previsto no art. 323, II, f, do RI" (1º evento com
  a cláusula e o nº da Decisão) e o despacho do Relator que fixa o teto em 50% do valor máximo da multa;
- débito vivo: é cominatória no teto quando o valor é 50% de um dos valores máximos anuais da multa e bate
  com `LimiteValor` (tipo 5) ou com o "limite máximo estabelecido (R$ X)" da DE, quando existirem.
  Só o 004898/2024 (multa do art. 107, II, f, pelo Acórdão nº 104/2024-TC) fica sem o parágrafo do teto.
- notificações de desconto em folha: texto genérico ("volume expressivo"), sem quantidade nem valores.

Modelo: scripts/automacao/templates/modelo_informacao.docx (numera os parágrafos); as ementas e o
dispositivo do STF entram por pós-processamento como citação recuada, sem número.
Os textos dos PDFs (~40 min no share) ficam em cache em output/automacao/; apagar o .pkl para reler.
Rodar: .venv/Scripts/python.exe processos/utils/gerar_informacoes_nereu_retomada_gcaed_v2.py
"""
import pickle
import re
import shutil
import sys
from copy import deepcopy
from datetime import datetime
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docxtpl import DocxTemplate

from ccd.config import OUTPUT_DIR, REPO_ROOT
from ccd.db import run_query_df
from ccd.docs import docx_to_pdf
from ccd.processo import get_informacoes_processo

sys.path.insert(0, str(Path(__file__).parent))
from gerar_informacoes_nereu_retomada_gcaed import (  # noqa: E402
    BASE,
    MARCADOR,
    carregar,
    sem_acento,
)

TEMPLATE = REPO_ROOT / "scripts" / "automacao" / "templates" / "modelo_informacao.docx"
DESTINO = BASE / "projetos" / "nereu_retomada_gcaed_v2"
CACHE = OUTPUT_DIR / "automacao" / "nereu_retomada_gcaed_v2_textos.pkl"
RELATOR = "Antônio Ed Souza Santana"  # Conselheiro titular
MS = "Mandado de Segurança Cível nº 0807247-93.2025.8.20.0000"
DATA = "Natal/RN, data da assinatura eletrônica."
# multa isolada: (processo) -> (acórdão, evento, valor), conferidos no texto do acórdão
MULTA_ISOLADA = {"004898/2024": ("Acórdão nº 104/2024-TC", 66, 1000.00)}
# valor máximo da multa (art. 323 do RI, atualizado por portaria da Presidência) usado pela DE no teto
VALORES_MAXIMOS = {16054.81, 17728.31, 19659.84}

CLAUSULA = re.compile(r"multa di ?a ?ria.{0,400}?limitad[oa] ao ?teto previsto no ?a ?rt\. ?323", re.I)
DECISAO = re.compile(r"DECIS ?A ?O\s*N[.o]?\s*[:.]?\s*(\d{1,4})\s*/\s*(\d{4})", re.I)
DESPACHO_TETO = re.compile(r"50 ?% ?\(cinquenta por cento\) do valor m ?a ?ximo", re.I)
LIMITE_DE = re.compile(r"limite m ?a ?ximo es ?tabelecido \(R\$ ?([\d.,]+)\)", re.I)

SQL_DEBITOS = """
    SELECT RTRIM(p.numero_processo)+'/'+RTRIM(p.ano_processo) AS processo, RTRIM(p.assunto) AS assunto,
           d.IdDebito, d.valorOriginalDebito AS valor, mc.LimiteValor AS limite
    FROM Pro_MarcadorProcesso mp
    JOIN Processos p ON p.IdProcesso = mp.IdProcesso
    JOIN Exe_Debito d ON d.IdProcessoExecucao = p.IdProcesso AND d.DataCancelamento IS NULL
    LEFT JOIN Exe_Debito_MultaCominatoria mc ON mc.IdDebito = d.IdDebito
    WHERE mp.IdMarcador = :marcador AND mp.DataExclusao IS NULL
      AND NOT EXISTS (SELECT 1 FROM Exe_Debito f WHERE f.IdDebitoAnterior = d.IdDebito)
"""

EMENTA_DENEGA = (
    "“MANDADO DE SEGURANÇA. MULTAS APLICADAS PELO TRIBUNAL DE CONTAS. ALEGAÇÃO DE ILEGALIDADE E "
    "ABUSIVIDADE. NECESSIDADE DE PROVA PRÉ-CONSTITUÍDA DO DIREITO LÍQUIDO E CERTO. INEXISTÊNCIA DE "
    "DEMONSTRAÇÃO DE QUE AS PENALIDADES IMPUGNADAS FORAM APLICADAS EM PERÍODOS COBERTOS POR DECISÃO "
    "JUDICIAL. ATUAÇÃO DO TCE/RN FUNDADA EM COMPETÊNCIA CONSTITUCIONAL. IMPOSSIBILIDADE DE DILAÇÃO "
    "PROBATÓRIA E DE REEXAME DO MÉRITO ADMINISTRATIVO NA VIA MANDAMENTAL. SEGURANÇA DENEGADA.”")
EMENTA_CONCEDE = (
    "“EMBARGOS DE DECLARAÇÃO EM MANDADO DE SEGURANÇA. SANÇÕES APLICADAS PELO TRIBUNAL DE CONTAS. "
    "DESCUMPRIMENTO DE DETERMINAÇÕES RELATIVAS A VANTAGENS REMUNERATÓRIAS. OMISSÃO CONFIGURADA. "
    "EXISTÊNCIA DE DECISÕES JUDICIAIS CONFLITANTES. AUSÊNCIA DE DEMONSTRAÇÃO DE DOLO OU CULPA DO GESTOR. "
    "EFEITOS INFRINGENTES. SEGURANÇA CONCEDIDA.”")
DISPOSITIVO_STF = [
    "“(...)",
    "Diante do exposto, DEFIRO O PEDIDO E SUSPENDO OS EFEITOS do Acórdão do Tribunal de Justiça do Rio "
    "Grande do Norte, publicado em 29 de junho de 2026, nos autos do Mandado de Segurança "
    "nº 0807247-93.2025.8.20.0000.",
    "Comunique-se COM URGÊNCIA o Tribunal de Justiça do Estado do Rio Grande do Norte e o Tribunal de "
    "Contas do Estado do Rio Grande do Norte.",
    "Ciência à Procuradoria-Geral da República.",
    "Publique-se.",
    "Brasília, 28 de julho de 2026.",
    "Ministro ALEXANDRE DE MORAES",
    "Vice-Presidente no exercício da Presidência”",
]


def no_teto(valor: float) -> bool:
    """valor = 50% de um valor máximo (a DE trunca o centavo: 16.054,81 → 8.027,40)."""
    return any(abs(valor - v / 2) < 0.01 for v in VALORES_MAXIMOS)


def brl(v: float) -> str:
    return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def textos(processos: list[str]) -> dict[str, list[dict]]:
    cache = pickle.loads(CACHE.read_bytes()) if CACHE.exists() else {}
    for p in processos:
        if p not in cache:
            print(f"  lendo PDFs de {p}...", flush=True)
            df = get_informacoes_processo(p)
            cache[p] = df[["evento", "setor", "ordem", "texto"]].to_dict("records")
            CACHE.parent.mkdir(parents=True, exist_ok=True)
            CACHE.write_bytes(pickle.dumps(cache))
    return {p: [{**r, "texto": sem_acento(r["texto"] or "")} for r in cache[p]] for p in processos}


def multa(processo: str, rows: list[dict], debito) -> dict:
    """Origem da multa e, quando diária, a prova de que o débito já está no teto."""
    if processo in MULTA_ISOLADA:
        acordao, evento, valor = MULTA_ISOLADA[processo]
        texto = next(r["texto"] for r in rows if r["evento"] == evento)
        n = re.search(r"(\d+/\d{4})", acordao).group(1)
        assert re.search(rf"ACORDAO N\w*\.? {n}", texto, re.I), f"{processo}: Evento {evento} não é o {acordao}"
        assert float(debito.valor) == valor and not no_teto(debito.valor), f"{processo}: débito {debito.valor}"
        return {"cominatoria": False, "decisao": acordao, "evento_decisao": evento, "valor": brl(valor)}
    rows = sorted(rows, key=lambda r: r["evento"])
    decisao = next(r for r in rows if CLAUSULA.search(r["texto"]) and DECISAO.search(r["texto"]))
    despacho = [r for r in rows if DESPACHO_TETO.search(r["texto"])]
    assert len(despacho) == 1, f"{processo}: {len(despacho)} despachos do Relator com o teto de 50%"
    assert no_teto(debito.valor), f"{processo}: {debito.valor} não é 50% do valor máximo"
    limites = {float(v.replace(".", "").replace(",", ".")) for r in rows for v in LIMITE_DE.findall(r["texto"])}
    if debito.limite == debito.limite:  # não-NaN: tipo 5
        limites.add(float(debito.limite))
    assert limites <= {float(debito.valor)}, f"{processo}: limite {limites} ≠ débito {debito.valor}"
    n, ano = DECISAO.search(decisao["texto"]).groups()
    return {"cominatoria": True, "decisao": f"Decisão nº {n}/{ano}-TC", "evento_decisao": decisao["evento"],
            "evento_despacho": despacho[0]["evento"], "valor": brl(debito.valor)}


def paragrafos(c: dict) -> list[tuple[str, list[str]]]:
    """(parágrafo numerado, citações recuadas que vêm logo depois)."""
    nereu = ("ao Sr. Nereu Batista Linhares, na qualidade de gestor do Instituto de Previdência dos "
             "Servidores do Estado do Rio Grande do Norte (IPERN)")
    if c["cominatoria"]:
        abertura = (
            f"Trata-se de Processo de Execução instaurado para a cobrança de débito imputado "
            f"{nereu}, em razão do descumprimento da {c['decisao']} (Evento nº {c['evento_decisao']}), "
            f"proferida nos autos do Processo nº {c['origem']}-TC, que fixou multa diária de R$ 50,00 "
            "(cinquenta reais) por dia de atraso, limitada ao teto previsto no art. 323, inciso II, "
            "alínea “f”, do Regimento Interno.")
    else:
        abertura = (
            f"Trata-se de Processo de Execução instaurado para a cobrança de débito imputado {nereu}, "
            f"decorrente da multa de R$ {c['valor']} aplicada pelo {c['decisao']} (Evento nº {c['evento_decisao']}), proferido nos autos "
            f"do Processo nº {c['origem']}-TC, com fundamento no art. 107, inciso II, alínea “f”, da Lei "
            "Complementar Estadual nº 464/2012, em razão do descumprimento de decisão desta Corte.")
    ps: list[tuple[str, list[str]]] = [
        (abertura, []),
        (f"Por Despacho de {c['data_sobrestamento']} (Evento nº {c['evento_sobrestamento']}), o Conselheiro "
         f"Relator determinou o sobrestamento dos processos alcançados pela decisão proferida no {MS}, até o "
         "desfecho judicial da demanda, com fundamento no art. 36, inciso III, da Lei Complementar Estadual "
         "nº 464/2012 c/c o art. 184, inciso III, do Regimento Interno, e o envio do feito a esta Diretoria, "
         "onde deveria permanecer até novo comando ou superveniência de alteração fática relevante.", []),
        (f"Em 29/01/2026, o Tribunal de Justiça do Estado do Rio Grande do Norte, nos autos do {MS}, "
         "denegou a segurança, em acórdão assim ementado:", [EMENTA_DENEGA]),
        ("Em 29/06/2026, ao apreciar embargos de declaração opostos naquele feito, a mesma Corte de Justiça "
         "atribuiu-lhes efeitos infringentes e, reconsiderando o entendimento anteriormente adotado, "
         "concedeu a segurança, nos termos da seguinte ementa:", [EMENTA_CONCEDE]),
        ("Em 28/07/2026, nos autos da Suspensão de Segurança nº 5.749, em resposta ao pedido formulado por "
         "este Tribunal de Contas, o Ministro Alexandre de Moraes, Vice-Presidente no exercício da Presidência "
         "do Supremo Tribunal Federal (STF), deferiu o pedido e suspendeu os efeitos do acórdão concessivo "
         "(documento anexo), nos seguintes termos:", DISPOSITIVO_STF),
        ("Em 27/08/2026, a referida decisão transitou em julgado, consoante certidão emitida pela Secretaria "
         "Judiciária do STF (anexa). Suspensos os efeitos do acórdão que concedeu a segurança, subsiste o "
         "acórdão denegatório e cessa o fundamento do sobrestamento, impondo-se a retomada da marcha "
         "processual.", []),
    ]
    if c["cominatoria"]:
        ps.append((
            f"No que concerne ao montante devido, registre-se que a multa diária já atingiu o limite máximo "
            f"fixado. Nos termos da {c['decisao']} (Evento nº {c['evento_decisao']}) e do Despacho da Relatoria "
            f"(Evento nº {c['evento_despacho']}), que estabeleceu corresponder o teto a 50% (cinquenta por "
            "cento) do valor máximo da multa vigente, o débito encontra-se cadastrado em "
            f"R$ {c['valor']}, valor correspondente ao referido limite. Não há, portanto, período adicional "
            "de mora a ser apurado, estando o crédito integralmente quantificado.", []))
    ps += [
        ("No tocante às medidas de cobrança, observa-se que o responsável já é destinatário de volume "
         "expressivo de notificações para desconto em folha expedidas em processos desta Corte, dirigidas ao "
         "IPERN e à Secretaria de Estado da Administração (SEAD). Considerando que o desconto "
         "previsto no art. 118, inciso I, da Lei Complementar Estadual nº 464/2012 deve observar o limite de "
         "30% (trinta por cento) dos vencimentos, subsídios ou proventos do responsável, na linha da "
         "jurisprudência do Superior Tribunal de Justiça, a expedição de novas notificações revela-se de "
         "reduzida eficácia. Mostra-se oportuno, assim, o "
         "encaminhamento do débito à Procuradoria-Geral do Estado para inscrição em dívida ativa e cobrança, "
         "nos termos do art. 118, inciso II e § 1º, da mesma Lei Complementar.", []),
        ("Ante o exposto, esta Coordenadoria sugere o encaminhamento dos autos ao Ministério Público junto "
         "ao Tribunal de Contas (MPC), para manifestação acerca da remessa do débito à Procuradoria-Geral do "
         "Estado para inscrição em dívida ativa.", []),
    ]
    return ps


def citacao_apos(par, textos_citacao: list[str]) -> None:
    """Parágrafos recuados (4 cm), sem numeração, fonte 10, logo após `par`."""
    ancora = par._p
    for t in textos_citacao:
        novo = deepcopy(par._p)
        ppr = novo.find(qn("w:pPr"))
        if ppr is not None and ppr.find(qn("w:numPr")) is not None:
            ppr.remove(ppr.find(qn("w:numPr")))
        ancora.addnext(novo)
        p = docx.text.paragraph.Paragraph(novo, par._parent)
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
        p.runs[0].text = t
        p.runs[0].font.size = Pt(10)
        p.paragraph_format.left_indent = Cm(4)
        p.paragraph_format.first_line_indent = Cm(0)
        ancora = novo


def gerar(c: dict) -> tuple[Path, list[tuple[str, list[str]]]]:
    pasta = DESTINO / c["arquivo"]
    out = pasta / f'{c["arquivo"]}.docx'
    ps = paragrafos(c)
    doc = DocxTemplate(str(TEMPLATE))
    doc.render({"processo": f'{c["processo"]} - TC', "assunto": c["assunto"], "relator": RELATOR,
                "parágrafos": [t for t, _ in ps]})  # sem numerar: o modelo numera
    d = doc.docx  # pós-processamento (get_docx() descartaria o render)

    for texto, cit in ps:
        if cit:
            citacao_apos(next(p for p in d.paragraphs if p.text == texto), cit)

    # o modelo deixa duas linhas em branco após o título; o espaçamento do parágrafo basta
    titulo = next(p for p in d.paragraphs if p.text == "INFORMAÇÃO INSTRUTIVA")
    for _ in range(2):
        titulo._p.getnext().getparent().remove(titulo._p.getnext())
    # data antes do bloco de assinatura, no lugar das linhas em branco
    assinado = next(p for p in d.paragraphs if "assinado digitalmente" in p.text)
    prev = assinado._p.getprevious()
    while prev is not None and not "".join(t.text or "" for t in prev.findall(f".//{qn('w:t')}")).strip():
        vazio, prev = prev, prev.getprevious()
        vazio.getparent().remove(vazio)
    data = deepcopy(assinado._p)
    for r in data.findall(qn("w:r"))[1:]:
        data.remove(r)
    assinado._p.addprevious(data)
    p_data = docx.text.paragraph.Paragraph(data, assinado._parent)
    p_data.runs[0].text = DATA
    p_data.runs[0].italic = False
    p_data.paragraph_format.space_before = Pt(12)

    pasta.mkdir(parents=True, exist_ok=True)
    if out.exists():  # preserva a versão anterior (skill edicao-minima)
        bak = out.with_name(f"{out.stem}_{datetime.now():%Y%m%d_%H%M%S}{out.suffix}")
        shutil.copy2(out, bak)
        print(f"  versão anterior preservada em: {bak.name}")
    doc.save(str(out))
    docx_to_pdf(str(out), str(pasta))
    return out, ps


if __name__ == "__main__":
    itens = carregar()
    debitos = run_query_df(SQL_DEBITOS, marcador=MARCADOR).set_index("processo")
    assert debitos.index.is_unique, "processo com mais de um débito vivo"
    pdfs = textos([i["processo"] for i in itens])

    com_teto = 0
    for item in itens:
        p = item["processo"]
        numero, ano = p.split("/")
        deb = debitos.loc[p]
        origem = re.search(r"PROCESSO N\S* (\d{6}/\d{4})", deb.assunto)
        assert origem, f"{p}: assunto sem processo de origem: {deb.assunto!r}"
        c = {"processo": p, "arquivo": f"{numero}_{ano}", "assunto": deb.assunto, "origem": origem.group(1),
             "evento_sobrestamento": item["evento"], "data_sobrestamento": item["data"].strftime("%d/%m/%Y"),
             **multa(p, pdfs[p], deb)}
        out, ps = gerar(c)

        texto = "\n".join(par.text for par in docx.Document(str(out)).paragraphs)
        assert "{{" not in texto and "{%" not in texto, f"{p}: tag Jinja sobrando"
        for t, cit in ps:
            assert t in texto and all(x in texto for x in cit), f"{p}: faltou {t[:40]}"
        assert texto.count("Trata-se de") == 1 and DATA in texto
        assert ("limite máximo fixado" in texto) == c["cominatoria"], p
        assert out.with_suffix(".pdf").is_file(), "PDF não gerado"
        com_teto += c["cominatoria"]
        print(f'{p}: sobrestamento Ev. {item["evento"]}, {c["decisao"]} Ev. {c["evento_decisao"]}'
              + (f', despacho Ev. {c["evento_despacho"]}, R$ {c["valor"]} (teto)' if c["cominatoria"]
                 else f', R$ {c["valor"]} (multa isolada)'))

    assert com_teto == len(itens) - len(MULTA_ISOLADA), com_teto
    print(f"\n{len(itens)} informações em {DESTINO.relative_to(BASE)}/ ({com_teto} com teto)")
