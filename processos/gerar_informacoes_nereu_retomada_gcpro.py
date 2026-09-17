"""Informações de desconto em folha — lote GCPRO (Cons. Paulo Roberto) pós-SS 5.749/RN.

Seis processos na CCD com débito de NEREU BATISTA LINHARES em que o relator, após a Suspensão
de Segurança nº 5.749/RN (STF, 28/07/2026), determinou o retorno dos autos à DIP para execução
(despachos de 14 e 17/08/2026). Em 001368/2022 e 001843/2025 o relator ouviu antes a CONJU
(Parecer nº 411/2026-CJTC) e o acolheu em 19/08/2026 — a informação cita parecer e despacho.

Mesmo modelo do lote GCREN (scripts/automacao/templates/desconto_folha.docx); a diferença é o
parágrafo introdutório com a referência ao despacho, inserido via python-docx entre render() e
save() (docxtpl: get_docx() descarta o render — usar doc.docx).

Rodar (de processos/): ../.venv/Scripts/python.exe gerar_informacoes_nereu_retomada_gcpro.py
"""
import copy
import shutil
from datetime import datetime
from pathlib import Path

from docxtpl import DocxTemplate
from gerar_informacoes_nereu_desconto_folha import ORGAO, TEMPLATE, brl, por_extenso

from ccd.config import cpf
from ccd.db import run_query_df
from ccd.docs import docx_to_pdf

DESTINO = Path(__file__).parent / "nereu_retomada_gcpro"
CPF = cpf("NEREU")
RELATOR = "PAULO ROBERTO CHAVES ALVES"

# processo -> (evento do despacho que devolve à DIP, evento do parecer da CONJU ou None)
PROCESSOS = {
    "003515/2017": (106, None),
    "003790/2017": (142, None),
    "008943/2017": (123, None),
    "019033/2017": (142, None),
    "001368/2022": (79, 77),
    "001843/2025": (90, 88),
}

INTRO = (
    "O Conselheiro Relator, por meio do despacho constante do Evento nº {ev}, considerando a "
    "Suspensão de Segurança nº 5.749/RN, proferida em 28 de julho de 2026 pelo Ministro Alexandre "
    "de Moraes, do Supremo Tribunal Federal, que suspendeu os efeitos do acórdão do Tribunal de "
    "Justiça do Estado do Rio Grande do Norte no Mandado de Segurança nº 0807247-93.2025.8.20.0000, "
    "determinou o retorno dos autos a esta Diretoria para a adoção das medidas pertinentes à "
    "execução do decisum, nos termos do art. 118 da Lei Complementar Estadual nº 464/2012."
)
INTRO_CONJU = (
    "O Conselheiro Relator, por meio do despacho constante do Evento nº {ev}, acolheu o Parecer "
    "nº 411/2026-CJTC, da Consultoria Jurídica deste Tribunal (Evento nº {ev_conju}), que concluiu "
    "pela possibilidade atual de execução das multas aplicadas ao gestor do IPERN, Sr. Nereu "
    "Batista Linhares, com a reimplantação das medidas correlatas (desconto em folha, protesto e "
    "CADINQ), à luz da eficácia restabelecida pela Suspensão de Segurança nº 5.749/RN, proferida "
    "em 28 de julho de 2026 pelo Supremo Tribunal Federal, e determinou o retorno dos autos a esta "
    "Diretoria para a continuação do feito."
)


def carregar() -> list[dict]:
    """Débitos vigentes do responsável ligados ao processo (execução OU origem) + checagens."""
    ph = {f"p{i}": p for i, p in enumerate(PROCESSOS)}
    inc = ", ".join(f":{k}" for k in ph)
    df = run_query_df(f"""
        SELECT CONCAT(RIGHT('000000'+CAST(pro.numero_processo AS varchar),6),'/',
                      pro.ano_processo) AS processo,
               RTRIM(pro.assunto) AS assunto, RTRIM(pro.setor_atual) AS setor,
               RTRIM(r.nome) AS relator,
               ed.IdDebito, RTRIM(s.DescricaoStatusDivida) AS situacao, gp.Nome AS nome,
               processo.dbo.fn_Exe_RetornaValorAtualizado(ed.IdDebito) AS valor,
               (SELECT MAX(e.SequencialProcessoEvento) FROM processo.dbo.Pro_ProcessoEvento e
                 WHERE e.IdProcesso = pro.IdProcesso) AS ultimo_evento
        FROM processo.dbo.Processos pro
        JOIN processo.dbo.Relator r ON r.codigo = pro.codigo_relator
        JOIN processo.dbo.Exe_Debito ed
             ON ed.IdProcessoExecucao = pro.IdProcesso OR ed.IdProcessoOrigem = pro.IdProcesso
        JOIN processo.dbo.Exe_StatusDivida s ON s.CodigoStatusDivida = ed.CodigoStatusDivida
        JOIN processo.dbo.Exe_DebitoPessoa edp ON edp.IDDebito = ed.IdDebito
        JOIN processo.dbo.GenPessoa gp ON gp.IdPessoa = edp.IDPessoa
        WHERE gp.Documento = :cpf AND ed.DataCancelamento IS NULL
          AND CONCAT(RIGHT('000000'+CAST(pro.numero_processo AS varchar),6),'/',
                     pro.ano_processo) IN ({inc})""", cpf=CPF, **ph)

    itens = []
    for proc, (ev, ev_conju) in PROCESSOS.items():
        g = df[df.processo == proc]
        assert not g.empty, f"{proc}: nenhum débito vigente do responsável"
        assert g.setor.iloc[0] == "CCD", f"{proc}: está em {g.setor.iloc[0]}"
        assert g.relator.iloc[0] == RELATOR, f"{proc}: relator {g.relator.iloc[0]}"
        assert ev <= g.ultimo_evento.iloc[0], f"{proc}: evento {ev} > último {g.ultimo_evento.iloc[0]}"
        itens.append({
            "processo": proc, "assunto": g.assunto.iloc[0], "nome": g.nome.iloc[0],
            "valor": float(g.valor.sum()),
            "debitos": [(int(r.IdDebito), r.situacao) for r in g.itertuples()],
            "intro": (INTRO_CONJU if ev_conju else INTRO).format(ev=ev, ev_conju=ev_conju),
        })
    return itens


def contexto(item: dict) -> dict:
    return {
        "processo": item["processo"], "assunto": item["assunto"],
        "nome": f'{item["nome"]} (CPF: {CPF})', "orgao": ORGAO,
        "valor": f'{brl(item["valor"])} ({por_extenso(item["valor"])})',
    }


def inserir_intro(doc, texto: str) -> None:
    """Clona o parágrafo 'Em cumprimento...' (mesma formatação) e o põe antes, com o texto novo."""
    alvo = next(p for p in doc.paragraphs if p.text.startswith("Em cumprimento"))
    novo_p = copy.deepcopy(alvo._p)
    alvo._p.addprevious(novo_p)
    novo = next(p for p in doc.paragraphs if p._p is novo_p)
    novo.runs[0].text = texto
    for r in novo.runs[1:]:
        r._r.getparent().remove(r._r)


def gerar(item: dict) -> Path:
    num, ano = item["processo"].split("/")
    pasta = DESTINO / f"{num}_{ano}"
    pasta.mkdir(parents=True, exist_ok=True)
    out = pasta / f"informacao_{num}_{ano}.docx"
    doc = DocxTemplate(TEMPLATE)
    doc.render(contexto(item))
    inserir_intro(doc.docx, item["intro"])
    if out.exists():  # preserva a versão anterior (skill edicao-minima)
        bak = out.with_name(f"{out.stem}_{datetime.now():%Y%m%d_%H%M%S}{out.suffix}")
        shutil.copy2(out, bak)
        print(f"  versão anterior preservada em: {bak.name}")
    doc.save(str(out))
    docx_to_pdf(str(out), str(pasta))
    return out


if __name__ == "__main__":
    import docx

    for item in carregar():
        ctx = contexto(item)
        deb = ", ".join(f"{i} ({s})" for i, s in item["debitos"])
        print(f'{item["processo"]}: {brl(item["valor"])} | débitos {deb}')
        out = gerar(item)

        # check: placeholders substituídos, campos no lugar e intro imediatamente antes do corpo
        d = docx.Document(str(out))
        pars = [p.text for p in d.paragraphs]
        texto = "\n".join(pars + [c.text for t in d.tables for r in t.rows for c in r.cells])
        assert "{{" not in texto and "}}" not in texto, "placeholder não substituído"
        for campo, valor in ctx.items():
            assert valor in texto, f'{item["processo"]}: faltou {campo} ("{valor[:40]}")'
        i = pars.index(item["intro"])
        assert pars[i + 1].startswith("Em cumprimento"), "intro fora do lugar"
        assert out.with_suffix(".pdf").is_file(), "PDF não gerado"
    print(f"\n{len(PROCESSOS)} informações em {DESTINO}")
