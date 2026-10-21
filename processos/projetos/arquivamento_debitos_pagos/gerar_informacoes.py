"""Informações de arquivamento de processos na CCD com todos os débitos quitados.

Lote levantado em 06/10/2026 (débitos vigentes — folha da cadeia — todos "Pago
Integralmente" ou cancelados, processo hoje na CCD). O 011064/2016 tem gerador próprio
em `011064_2016/`; o 021694/2016 já tem informação de arquivamento (Evento 83) e só
falta tramitar, por isso fica fora daqui.

- 000924/2020, 007671/2008, 701134/2012: parcelamentos quitados no fim de 09/2026; a
  certidão declaratória de quitação (art. 26, p.ú., Res. 013/2015) ainda não está nos
  autos — emitir e preencher CERTIDOES no topo antes de cadastrar.
- 000633/2026, 000625/2026: houve notificação para desconto em folha (SEAD / ALRN); além
  do arquivamento, sugere-se comunicar a quitação ao órgão (a ALRN pediu esclarecimento
  no Ofício nº 590/2026-GP, Evento 94).

O script confere no banco a situação e o valor pago de cada débito antes de gerar.
Rodar: .venv/Scripts/python.exe processos/projetos/arquivamento_debitos_pagos/gerar_informacoes.py
"""
import shutil
from copy import deepcopy
from datetime import datetime
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor
from docxtpl import DocxTemplate

from ccd.config import REPO_ROOT
from ccd.db import run_query_df

TEMPLATE = str(REPO_ROOT / "scripts" / "automacao" / "templates" / "modelo_informacao.docx")
DESTINO = Path(__file__).resolve().parent

MESES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho",
         "agosto", "setembro", "outubro", "novembro", "dezembro"]
hoje = datetime.now()
DATA = f"Natal/RN, {hoje.day} de {MESES[hoje.month - 1]} de {hoje.year}."

# preencher após emitir as certidões de quitação (ato de ofício da CCD)
CERTIDOES = {
    "000924/2020": ("____/2026", "__"),  # (nº da certidão, evento)
    "007671/2008": ("____/2026", "__"),
    "701134/2012": ("____/2026", "__"),
}

ARQUIVAMENTO = ("nos termos do art. 26 da Resolução nº 013/2015-TCE e do art. 209, inciso "
                "III, do Regimento Interno desta Corte de Contas.")


def par_certidao(proc: str, codigo: str, nome: str) -> str:
    num, ev = CERTIDOES[proc]
    return ("Nos termos do art. 26 da Resolução nº 013/2015-TCE, comprovado o recolhimento "
            "do valor devido, procede-se à exclusão do nome do responsável do Cadastro "
            "Informativo de Créditos Não Quitados e ao posterior arquivamento do processo de "
            "execução, devendo constar nos autos certidão declaratória da quitação "
            "(parágrafo único). Em cumprimento ao referido dispositivo, esta Coordenadoria "
            f"emitiu a Certidão Declaratória de Quitação nº {num}, relativa à multa "
            f"identificada pelo código {codigo}, em nome do Senhor {nome} (Evento {ev}).")


PROCESSOS = [
    {
        "processo": "000924/2020",
        "assunto": "APURAÇÃO DE RESPONSABILIDADE - PORTAL DA TRANSPARÊNCIA 2019",
        "relator": "Antônio Gilberto de Oliveira Jales",
        "debitos": {19786: 7360.92},  # IdDebito: ValorPago esperado
        "paragrafos": [
            "Trata-se de processo de apuração de responsabilidade, instaurado em face da "
            "Câmara Municipal de Lagoa Salgada, no qual o Acórdão nº 108/2021-TC, de 15 de "
            "abril de 2021, impôs multa ao Senhor Ozivaldo Nascimento Queiroz, no valor de "
            "R$ 5.000,00 (cinco mil reais). Após o desprovimento do agravo interposto contra "
            "a decisão que inadmitiu o pedido de reconsideração (Acórdão nº 106/2023-TC, "
            "Evento 103), o Acórdão nº 108/2021-TC transitou em julgado em 14 de fevereiro "
            "de 2023 (Evento 116). Os autos foram encaminhados a esta Diretoria para o "
            "acompanhamento da cobrança pelo Despacho do Conselheiro Relator de 10 de "
            "fevereiro de 2025 (Evento 144).",

            "O responsável requereu o parcelamento da multa em 33 (trinta e três) parcelas, "
            "na forma do art. 20 da Resolução nº 013/2015-TCE, e efetuou o pagamento da "
            "primeira parcela em 31 de janeiro de 2024 e da última em 25 de setembro de "
            "2026, totalizando R$ 7.360,92 (sete mil, trezentos e sessenta reais e noventa "
            "e dois centavos), conforme o quadro a seguir:",

            par_certidao("000924/2020", "19.786", "Ozivaldo Nascimento Queiroz"),

            "Ante o exposto, cumprida a obrigação imposta no Acórdão nº 108/2021-TC, "
            f"sugere-se o arquivamento do feito, {ARQUIVAMENTO}",
        ],
        "quadro": ["19.786", "Ozivaldo Nascimento Queiroz", "R$ 5.000,00", "R$ 7.360,92",
                   "33 parcelas, de 31/01/2024 a 25/09/2026"],
    },
    {
        "processo": "007671/2008",
        "assunto": "PRESTAÇÃO DE CONTAS DE ACORDO COM A RESOLUÇÃO 012/2007 REF. AO BIMESTRE: 01/2008",
        "relator": "Paulo Roberto Chaves Alves",
        "debitos": {18769: 47936.45},
        "paragrafos": [
            "Trata-se de prestação de contas da Prefeitura Municipal de Caiçara do Rio do "
            "Vento, referente ao 1º bimestre de 2008, na qual o Acórdão nº 234/2020-TC, de "
            "29 de setembro de 2020, imputou multa ao Senhor Felipe Eloi Muller, registrada "
            "sob o código 18.769, no valor de R$ 35.300,00 (trinta e cinco mil e trezentos "
            "reais). O pedido de reconsideração interposto foi desprovido pelo Acórdão "
            "nº 452/2023-TC (Evento 36), e o Acórdão nº 234/2020-TC transitou em julgado em "
            "25 de agosto de 2023 (Evento 87).",

            "Por Despacho de 25 de fevereiro de 2025 (Evento 83), o Conselheiro Relator "
            "registrou que o responsável havia aderido ao parcelamento da multa e determinou "
            "o encaminhamento dos autos a esta Diretoria para o acompanhamento do pagamento "
            "das parcelas. O responsável efetuou o pagamento das 36 (trinta e seis) parcelas, "
            "a primeira em 30 de outubro de 2023 e a última em 29 de setembro de 2026, "
            "totalizando R$ 47.936,45 (quarenta e sete mil, novecentos e trinta e seis reais "
            "e quarenta e cinco centavos), conforme o quadro a seguir:",

            par_certidao("007671/2008", "18.769", "Felipe Eloi Muller"),

            "Ante o exposto, cumprida a obrigação imposta no Acórdão nº 234/2020-TC, "
            f"sugere-se o arquivamento do feito, {ARQUIVAMENTO}",
        ],
        "quadro": ["18.769", "Felipe Eloi Muller", "R$ 35.300,00", "R$ 47.936,45",
                   "36 parcelas, de 30/10/2023 a 29/09/2026"],
    },
    {
        "processo": "701134/2012",
        "assunto": "PRESTAÇÃO DE CONTAS DE ACORDO COM A RESOLUÇÃO 022/2011 REF. AO BIMESTRE: 01/2012",
        "relator": "Marco Antônio de Moraes Rêgo Montenegro",
        "debitos": {27706: 430.47, 28584: 2333.36},
        "status": {27706: 3, 28584: 2},  # o 27706 fica "pago parcialmente": saldo = 28584
        "paragrafos": [
            "Trata-se de prestação de contas da Câmara Municipal de Ipanguaçu, referente ao "
            "1º bimestre de 2012, na qual foram impostas multas ao Senhor Tunefis da Silva "
            "Moraes pelo Acórdão nº 199/2016-1ª Câmara. Em sede de pedido de reconsideração, "
            "o Acórdão nº 266/2024-TC (Evento 16) excluiu a multa do item \"a\" e manteve os "
            "demais dispositivos, remanescendo a multa de R$ 1.600,00 (mil e seiscentos "
            "reais). Os embargos de declaração opostos foram desprovidos pelo Acórdão "
            "nº 198/2025-TC (Evento 38), que transitou em julgado em 31 de março de 2025 "
            "(Evento 56).",

            "Citado para o pagamento, o responsável parcelou a multa em 12 (doze) parcelas e "
            "recolheu as duas primeiras, em 29 de agosto e 30 de setembro de 2025, no total "
            "de R$ 430,47. Diante da falta de pagamento da 3ª parcela, o parcelamento foi "
            "cancelado em 1º de dezembro de 2025, com o vencimento antecipado das parcelas "
            "restantes (Evento 68), nos termos do art. 20, § 3º, da Resolução nº "
            "013/2015-TCE. O saldo devedor, registrado sob o código 28.584, foi objeto de "
            "novo parcelamento, em 10 (dez) parcelas, pagas entre 11 de dezembro de 2025 e "
            "28 de setembro de 2026, no total de R$ 2.333,36, conforme o quadro a seguir:",

            par_certidao("701134/2012", "28.584", "Tunefis da Silva Moraes"),

            "Ante o exposto, cumprida a obrigação imposta no Acórdão nº 199/2016-1ª Câmara, "
            f"com a redação dada pelo Acórdão nº 266/2024-TC, sugere-se o arquivamento do "
            f"feito, {ARQUIVAMENTO}",
        ],
        "quadros": [
            ["27.706", "Tunefis da Silva Moraes", "R$ 1.600,00", "R$ 430,47",
             "Parcelas 1 e 2/12 (29/08 e 30/09/2025); parcelamento cancelado em 01/12/2025"],
            ["28.584", "Tunefis da Silva Moraes", "R$ 1.333,34 (saldo)", "R$ 2.333,36",
             "10 parcelas, de 11/12/2025 a 28/09/2026"],
        ],
    },
    {
        "processo": "000633/2026",
        "assunto": "EXECUÇÃO DA DECISÃO PROFERIDA NOS AUTOS DO PROCESSO Nº 200030/2023 - TC",
        "relator": "Antonio Ed Souza Santana",
        "debitos": {28682: 2760.61},
        "paragrafos": [
            "Trata-se de processo de execução do Acórdão nº 740/2025-TC, de 17 de novembro "
            "de 2025, proferido nos autos do Processo nº 200030/2023-TC e transitado em "
            "julgado em 27 de janeiro de 2026, que impôs multa ao Senhor Gustavo Montenegro "
            "Soares, no valor de R$ 2.600,00 (dois mil e seiscentos reais), inscrita no "
            "Cadastro Informativo de Créditos Não Quitados em 3 de julho de 2026 (Certidão "
            "nº 600760/2026, Evento 44).",

            "Em seguida, esta Coordenadoria propôs a notificação da Secretaria de Estado da "
            "Administração (SEAD) para o desconto da multa nos vencimentos do responsável "
            "(Evento 45), nos termos do art. 25, inciso I, da Resolução nº 013/2015-TCE, "
            "tendo sido expedida a Notificação nº 001749/2026 (Evento 49), recebida "
            "tacitamente em 23 de julho de 2026 (Evento 50).",

            "Antes da efetivação do desconto, em 29 de julho de 2026, o responsável recolheu "
            "integralmente a multa, no valor atualizado de R$ 2.760,61 (dois mil, setecentos "
            "e sessenta reais e sessenta e um centavos), conforme o quadro a seguir. Foi "
            "emitida a Certidão Declaratória de Quitação de Multa nº 502303/2026 (Evento 52), "
            "com a exclusão do nome do responsável do Cadastro Informativo de Créditos Não "
            "Quitados, e a Diretoria de Expediente certificou o recolhimento total dos "
            "valores (Evento 53).",

            "Ante o exposto, cumprida a obrigação imposta no Acórdão nº 740/2025-TC, "
            "sugere-se:",
        ],
        "itens": [
            "a) comunicação à SEAD da quitação integral da multa, a fim de que se abstenha de "
            "efetuar o desconto determinado na Notificação nº 001749/2026; e",
            f"b) arquivamento do feito, {ARQUIVAMENTO}",
        ],
        "quadro": ["28.682", "Gustavo Montenegro Soares", "R$ 2.600,00", "R$ 2.760,61",
                   "Pagamento integral em 29/07/2026 (guia nº 039933)"],
    },
    {
        "processo": "000625/2026",
        "assunto": "EXECUÇÃO DA DECISÃO PROFERIDA NOS AUTOS DO PROCESSO Nº 000912/2020 - TC",
        "relator": "Antônio Gilberto de Oliveira Jales",
        "debitos": {20785: 5477.19},
        "paragrafos": [
            "Trata-se de processo de execução do Acórdão nº 369/2021-TC, de 24 de agosto de "
            "2021, proferido nos autos do Processo nº 000912/2020-TC e transitado em julgado "
            "em 29 de julho de 2025, que impôs multa ao Senhor Leonardo Nunes Rego, no valor "
            "de R$ 4.000,00 (quatro mil reais), inscrita no Cadastro Informativo de Créditos "
            "Não Quitados em 17 de junho de 2026 (Certidão nº 600709/2026, Evento 71).",

            "Em seguida, esta Coordenadoria propôs a notificação da Assembleia Legislativa do "
            "Estado do Rio Grande do Norte para o desconto da multa nos vencimentos do "
            "responsável (Evento 72), nos termos do art. 25, inciso I, da Resolução nº "
            "013/2015-TCE, tendo sido expedida a Notificação nº 001582/2026. O responsável, "
            "por sua vez, requereu a não efetivação do desconto e a liberação do parcelamento "
            "(Eventos 78 a 80), pedido acolhido pelo Conselheiro Relator por Despacho de 4 de "
            "agosto de 2026 (Evento 87).",

            "Liberada a funcionalidade, o responsável parcelou a multa em 2 (duas) parcelas, "
            "com deferimento automático pelo pagamento da primeira (Evento 91), nos termos do "
            "art. 20, § 2º, da Resolução nº 013/2015-TCE, e quitou a segunda em 11 de "
            "setembro de 2026, totalizando R$ 5.477,19 (cinco mil, quatrocentos e setenta e "
            "sete reais e dezenove centavos), conforme o quadro a seguir. Foi emitida a "
            "Certidão Declaratória de Quitação de Multa nº 02338/2026 (Evento 92), com a "
            "exclusão do nome do responsável do Cadastro Informativo de Créditos Não "
            "Quitados.",

            "Por meio do Ofício nº 590/2026-GP-ALRN, de 15 de setembro de 2026 (Evento 94), a "
            "Assembleia Legislativa solicitou esclarecimentos quanto à necessidade de "
            "cumprimento da determinação de desconto, tendo em vista o parcelamento do débito "
            "perante esta Corte. Com a quitação integral da multa, não há valor a ser "
            "descontado.",

            "Ante o exposto, cumprida a obrigação imposta no Acórdão nº 369/2021-TC, "
            "sugere-se:",
        ],
        "itens": [
            "a) comunicação à Assembleia Legislativa do Estado do Rio Grande do Norte da "
            "quitação integral da multa, em resposta ao Ofício nº 590/2026-GP-ALRN, a fim de "
            "que se abstenha de efetuar o desconto determinado na Notificação nº 001582/2026; e",
            f"b) arquivamento do feito, {ARQUIVAMENTO}",
        ],
        "quadro": ["20.785", "Leonardo Nunes Rego", "R$ 4.000,00", "R$ 5.477,19",
                   "2 parcelas (10/08/2026 e 11/09/2026)"],
    },
]

# identidade visual TCE/RN (skill tce-rn-identity) — mesmo padrão de 011064_2016
TCE = {"verde_principal": "2E5B3C", "verde_escuro": "1A3D28", "cinza_claro": "F2F2F2",
       "cinza_borda": "CCCCCC", "cinza_texto": "333333"}
CABECALHO = ["Código", "Responsável", "Valor original", "Valor recolhido", "Pagamento"]


def conferir(item: dict) -> None:
    """Aborta se o banco não confirmar a situação/valor pago usados no texto."""
    ids = ",".join(str(i) for i in item["debitos"])
    deb = run_query_df(f"SELECT IdDebito, CodigoStatusDivida, ValorPago FROM Exe_Debito "
                       f"WHERE IdDebito IN ({ids})").set_index("IdDebito")
    for id_debito, pago in item["debitos"].items():
        r = deb.loc[id_debito]
        assert int(r.CodigoStatusDivida) == item.get("status", {}).get(id_debito, 2), (id_debito, r)
        assert round(float(r.ValorPago), 2) == pago, (id_debito, r)


def _times(run) -> None:
    """Times New Roman como o modelo; sem isso a tabela herda Calibri do tema."""
    run.font.name = "Times New Roman"
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), "Times New Roman")


def _shade(cell, hex_cor: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_cor)
    cell._tc.get_or_add_tcPr().append(shd)


def _tabela_apos(d, ancora_texto: str, legenda: str, linhas: list[list[str]]) -> None:
    ancora = next(p for p in d.paragraphs if ancora_texto in p.text)
    cap = d.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_cap = cap.add_run(legenda)
    run_cap.bold = True
    run_cap.font.color.rgb = RGBColor.from_string(TCE["verde_escuro"])
    _times(run_cap)
    cap.paragraph_format.keep_with_next = True  # legenda não fica órfã da tabela
    ancora._p.addnext(cap._p)
    t = d.add_table(rows=0, cols=len(CABECALHO))
    borders = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"),
               TCE["cinza_borda"] if lado.startswith("inside") else TCE["verde_principal"])
        borders.append(el)
    t._tbl.tblPr.append(borders)
    for i, linha in enumerate([CABECALHO] + linhas):
        for cell, texto in zip(t.add_row().cells, linha, strict=True):
            if i == 0:
                _shade(cell, TCE["verde_principal"])
            elif i % 2 == 1:
                _shade(cell, TCE["cinza_claro"])
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = par.add_run(texto)
            run.bold = i == 0
            run.font.color.rgb = RGBColor.from_string("FFFFFF" if i == 0 else TCE["cinza_texto"])
            _times(run)
    cap._p.addnext(t._tbl)


def _itens_apos(d, intro_texto: str, itens: list[str]) -> None:
    """Alíneas a), b)… em parágrafos próprios, recuados e sem a numeração do modelo
    (formato ajustado à mão na informação do 000633/2026)."""
    intro = next(p for p in d.paragraphs if p.text.endswith(intro_texto))
    ancora = intro._p
    for texto in itens:
        novo = deepcopy(intro._p)
        ppr = novo.find(qn("w:pPr"))
        num = ppr.find(qn("w:numPr"))
        if num is not None:
            ppr.remove(num)
        ind = ppr.find(qn("w:ind"))
        if ind is None:
            ind = OxmlElement("w:ind")
            ppr.append(ind)
        ind.set(qn("w:left"), "1701")
        ind.set(qn("w:firstLine"), "0")
        runs = novo.findall(qn("w:r"))
        for r in runs[1:]:
            novo.remove(r)
        ts = runs[0].findall(qn("w:t"))
        ts[0].text = texto
        for t in ts[1:]:
            t.getparent().remove(t)
        ancora.addnext(novo)
        ancora = novo


def _data_antes_da_assinatura(d) -> None:
    assinado = next(p for p in d.paragraphs if "assinado digitalmente" in p.text)
    for texto in (DATA, ""):
        novo = deepcopy(assinado._p)
        for run_el in list(novo):
            if run_el.tag == qn("w:r"):
                novo.remove(run_el)
        assinado._p.addprevious(novo)
        if texto:
            run_el = deepcopy(assinado.runs[0]._r)
            novo.append(run_el)
            for t in run_el.findall(qn("w:t")):
                t.text = texto


def gerar(item: dict) -> Path:
    conferir(item)
    numero, ano = item["processo"].split("/")
    pasta = DESTINO / f"{numero}_{ano}"
    pasta.mkdir(exist_ok=True)
    out = pasta / f"informacao_{numero}_{ano}.docx"

    doc = DocxTemplate(TEMPLATE)
    doc.render({
        "processo": f"{item['processo']} - TC",
        "assunto": item["assunto"],
        "relator": item["relator"],
        "parágrafos": item["paragrafos"],  # sem numerar: o modelo numera
    })
    d = doc.docx  # pós-processamento (get_docx() descartaria o render)
    _tabela_apos(d, "conforme o quadro a seguir",
                 f"Quadro 1 – Situação da multa executada no Processo nº {item['processo']}-TC",
                 item.get("quadros") or [item["quadro"]])
    if item.get("itens"):
        _itens_apos(d, "sugere-se:", item["itens"])
    _data_antes_da_assinatura(d)

    # preserva a versão anterior antes de sobrescrever (skill edicao-minima)
    if out.exists():
        bak = out.with_name(f"{out.stem}_{datetime.now():%Y%m%d_%H%M%S}{out.suffix}")
        shutil.copy2(out, bak)
        print(f"  versão anterior preservada em: {bak.name}")
    doc.save(str(out))

    # check: o docx salvo traz processo, quadro, fundamento e data; sem "É a informação"
    salvo = docx.Document(str(out))
    texto = "\n".join(p.text for p in salvo.paragraphs)
    for trecho in (item["processo"], "Quadro 1 –", "art. 209, inciso III", DATA):
        assert trecho in texto, f"{item['processo']}: faltou {trecho}"
    assert "É a informação" not in texto
    assert len(salvo.tables) == 1
    for it in item.get("itens", []):
        assert any(p.text == it for p in salvo.paragraphs), f"{item['processo']}: faltou {it[:20]}"
    return out


if __name__ == "__main__":
    import sys
    # ex.: gerar_informacoes.py 000625/2026  (sem argumento, gera todos)
    for item in [p for p in PROCESSOS if not sys.argv[1:] or p["processo"] in sys.argv[1:]]:
        print(f"{item['processo']}: {gerar(item)}")
    pendentes = [p for p, (n, _) in CERTIDOES.items() if "_" in n]
    if pendentes:
        print(f"ATENÇÃO: certidão de quitação a preencher em {', '.join(pendentes)}")
