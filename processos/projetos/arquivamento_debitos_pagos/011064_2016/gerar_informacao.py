"""Informação do Processo nº 011064/2016-TC (execução do Acórdão nº 276/2015-TC,
proferido no 014484/2004 — multas a Ivis Alberto L. B. de Andrade e George
Antunes de Oliveira).

Sobrestado em 22/03/2023 (Evento 25) até a apreciação do pedido de unificação de
multas do 014502/2017, que foi julgado prejudicado por perda do objeto
(Decisão de 26/02/2024) e arquivado em 19/06/2024. As duas multas (Exe_Debito
4388 e 4389) estão quitadas e baixadas; a certidão de quitação que faltava
(4389, nº 502347/2026, Evento 38) foi emitida de ofício pela CCD (art. 26, p.ú.,
Res. 013/2015) e sugere-se apenas o arquivamento.

Rodar: .venv/Scripts/python.exe processos/projetos/arquivamento_debitos_pagos/011064_2016/gerar_informacao.py
"""
import shutil
from copy import deepcopy
from datetime import datetime
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor
from docxtpl import DocxTemplate

from ccd.config import REPO_ROOT
from ccd.db import run_query_df

TEMPLATE = str(REPO_ROOT / "scripts" / "automacao" / "templates" / "modelo_informacao.docx")
OUT = Path(__file__).parent / "informacao_011064_2016.docx"

MESES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho",
         "agosto", "setembro", "outubro", "novembro", "dezembro"]
hoje = datetime.now()
DATA = f"Natal/RN, {hoje.day} de {MESES[hoje.month - 1]} de {hoje.year}."

# certidão de quitação do débito 4389 (ato de ofício da CCD,
# art. 26, p.ú., Res. 013/2015 c/c art. 32, III, Res. 42/2024)
CERTIDAO_4389 = "502347/2026"
EVENTO_CERTIDAO = "38"

# --- conferência dos débitos no banco (o texto abaixo depende destes valores) ---
deb = run_query_df(
    "SELECT IdDebito, valorOriginalDebito, ValorPago, dataBaixa, Concluido, "
    "CodigoSituacaoDivida FROM processo.dbo.Exe_Debito WHERE IdDebito IN (4388, 4389) "
    "ORDER BY IdDebito")
assert len(deb) == 2, f"esperava 2 débitos, veio {len(deb)}"
esperado = {4388: (250.0, 333.05, "2018-07-24"), 4389: (55.0, 79.08, "2020-11-23")}
for _, r in deb.iterrows():
    orig, pago, baixa = esperado[int(r.IdDebito)]
    assert float(r.valorOriginalDebito) == orig, r
    assert round(float(r.ValorPago), 2) == pago, r
    assert r.dataBaixa.date().isoformat() == baixa, r
    assert bool(r.Concluido), r

paragrafos = [
    "Trata-se de processo de execução do Acórdão nº 276/2015-TC, de 9 de junho de 2015, "
    "proferido nos autos do Processo nº 014484/2004-TC e transitado em julgado em 13 de "
    "julho de 2015, que impôs multas ao Senhor Ivis Alberto Lourenço Bezerra de Andrade, "
    "no valor de R$ 250,00 (duzentos e cinquenta reais), e ao Senhor George Antunes de "
    "Oliveira, no valor de R$ 55,00 (cinquenta e cinco reais), ambas inscritas no "
    "Cadastro Informativo de Créditos Não Quitados em 1º de junho de 2016 (Eventos 2 e 3). "
    "Os autos foram transferidos a esta Coordenadoria em 8 de janeiro de 2025 (Evento 37), "
    "em decorrência da Lei Complementar Estadual nº 757, de 10 de junho de 2024.",

    "Por Despacho de 22 de março de 2023 (Evento 25), acolhendo a quota ministerial do "
    "Evento 22, o então Conselheiro Relator determinou o sobrestamento do feito até a "
    "apreciação do pedido de unificação de multas formulado pelo Senhor George Antunes de "
    "Oliveira no Processo nº 014502/2017-TC.",

    "O referido pedido já foi apreciado. Por Decisão de 26 de fevereiro de 2024, a "
    "Presidência desta Corte julgou prejudicado o prosseguimento do Processo nº "
    "014502/2017-TC, por perda do objeto, uma vez que a unificação produziria efeitos "
    "apenas sobre as multas pendentes de adimplemento e a maioria das sanções já constava "
    "como paga, cancelada ou negociada. Intimado, o requerente não interpôs recurso, e os "
    "autos foram encaminhados ao arquivo em 19 de junho de 2024, com fundamento no art. "
    "209, inciso II, do Regimento Interno. Cessou, portanto, a causa do sobrestamento.",

    "Quanto às multas executadas nestes autos, ambas encontram-se integralmente quitadas, "
    "conforme o quadro a seguir:",

    f"As certidões declaratórias de quitação constam dos Eventos 15 e {EVENTO_CERTIDAO}, e "
    "os nomes dos responsáveis foram excluídos do Cadastro Informativo de Créditos Não "
    "Quitados – CADINQ.",

    "Ante o exposto, cumpridas as obrigações impostas no Acórdão nº 276/2015-TC e cessada a "
    "causa do sobrestamento, sugere-se o arquivamento do feito, nos termos do art. 26 da "
    "Resolução nº 013/2015-TCE e do art. 209, inciso III, do Regimento Interno desta Corte "
    "de Contas.",
]

doc = DocxTemplate(TEMPLATE)
doc.render({
    "processo": "011064/2016 - TC",
    "assunto": "EXECUÇÃO DA DECISÃO PROFERIDA NOS AUTOS DO PROCESSO Nº 014484/2004 - TC",
    "relator": "George Montenegro Soares",
    "parágrafos": paragrafos,  # sem numerar: o modelo numera
})

# pós-processamento via doc.docx (get_docx() descartaria o render)
d = doc.docx

# o estilo Normal do modelo ainda é Calibri (linhas vazias, marcas de parágrafo)
_rf = d.styles["Normal"].element.get_or_add_rPr().get_or_add_rFonts()
for _attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
    _rf.set(qn(_attr), "Times New Roman")

# identidade visual TCE/RN (skill tce-rn-identity)
TCE = {
    "verde_principal": "2E5B3C",
    "verde_escuro": "1A3D28",
    "cinza_claro": "F2F2F2",
    "cinza_borda": "CCCCCC",
    "cinza_texto": "333333",
}


def _times(run) -> None:
    """Times New Roman como o modelo; sem isso a tabela herda Calibri do tema."""
    run.font.name = "Times New Roman"
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), "Times New Roman")


def _shade(cell, hex_cor: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_cor)
    cell._tc.get_or_add_tcPr().append(shd)


def _tabela_apos(ancora_texto: str, legenda: str, cabecalho: list[str],
                 linhas: list[list[str]]) -> None:
    """Insere legenda + tabela no padrão visual do TCE/RN após o parágrafo-âncora."""
    ancora = next(p for p in d.paragraphs if ancora_texto in p.text)
    cap = d.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_cap = cap.add_run(legenda)
    run_cap.bold = True
    run_cap.font.color.rgb = RGBColor.from_string(TCE["verde_escuro"])
    _times(run_cap)
    cap.paragraph_format.keep_with_next = True  # legenda não fica órfã da tabela
    ancora._p.addnext(cap._p)
    ancora = cap
    t = d.add_table(rows=0, cols=len(cabecalho))
    borders = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"),
               TCE["cinza_borda"] if lado.startswith("inside") else TCE["verde_principal"])
        borders.append(el)
    t._tbl.tblPr.append(borders)
    for i, linha in enumerate([cabecalho] + linhas):
        cells = t.add_row().cells
        for cell, texto in zip(cells, linha, strict=False):
            if i == 0:
                _shade(cell, TCE["verde_principal"])
            elif i % 2 == 1:
                _shade(cell, TCE["cinza_claro"])
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = par.add_run(texto)
            run.bold = i == 0
            run.font.color.rgb = RGBColor.from_string(
                "FFFFFF" if i == 0 else TCE["cinza_texto"])
            _times(run)
    ancora._p.addnext(t._tbl)


_tabela_apos(
    "conforme o quadro a seguir",
    "Quadro 1 – Situação das multas executadas no Processo nº 011064/2016-TC",
    ["Código", "Responsável", "Valor original", "Valor recolhido", "Baixa"],
    [
        ["4.388", "Ivis Alberto Lourenço Bezerra de Andrade", "R$ 250,00",
         "R$ 333,05 em 24/07/2018",
         "Certificada em 10/10/2018 (Evento 14); certidão de quitação nº 500829/2018 (Evento 15)"],
        ["4.389", "George Antunes de Oliveira", "R$ 55,00",
         "R$ 79,08 (PGE, doc. nº 5346/2020)",
         "Certificada em 13/06/2023 (Evento 29); certidão de quitação nº 502347/2026 (Evento 38)"],
    ],
)

# data antes do bloco de assinatura (o modelo não traz)
assinado = next(p for p in d.paragraphs if "assinado digitalmente" in p.text)
for texto in (DATA, ""):
    novo = deepcopy(assinado._p)
    for run_el in list(novo):
        if run_el.tag == qn("w:r"):
            novo.remove(run_el)
    assinado._p.addprevious(novo)
    if texto:
        run_el = deepcopy(assinado.runs[0]._r)
        novo.append(run_el)
        for t in run_el.findall(qn("w:t")):
            t.text = texto

# bloco de assinatura (data → nome/cargo) não se parte entre páginas
pars = d.paragraphs
ini = next(i for i, p in enumerate(pars) if p.text == DATA)
vazios = []  # vazios do modelo antes da data: deixa um só, para o bloco caber na página
for p in reversed(pars[:ini]):
    if p.text.strip():
        break
    vazios.append(p)
for p in vazios[1:]:
    p._p.getparent().remove(p._p)
pars = d.paragraphs
ini = next(i for i, p in enumerate(pars) if p.text == DATA)
for p in pars[ini:-1]:
    p.paragraph_format.keep_with_next = True

# preserva a versão anterior antes de sobrescrever (skill edicao-minima)
if OUT.exists():
    bak = OUT.with_name(f"{OUT.stem}_{datetime.now():%Y%m%d_%H%M%S}{OUT.suffix}")
    shutil.copy2(OUT, bak)
    print(f"Versão anterior preservada em: {bak.name}")
doc.save(str(OUT))

# check (ponytail): docx salvo traz os elementos essenciais e o quadro
import docx  # noqa: E402

salvo = docx.Document(str(OUT))
texto_final = "\n".join(p.text for p in salvo.paragraphs)
for trecho in ("011064/2016", "014484/2004", "014502/2017", "George Montenegro Soares",
               "Acórdão nº 276/2015-TC", "perda do objeto", "art. 26 da Resolução nº 013/2015-TCE",
               "art. 209, inciso III", "Quadro 1 – Situação das multas", "Eventos 15 e 38", "CADINQ", DATA):
    assert trecho in texto_final, f"faltou: {trecho}"
assert "É a informação" not in texto_final
assert len(salvo.tables) == 1 and len(salvo.tables[0].rows) == 3
assert all(r.font.name == "Times New Roman" for row in salvo.tables[0].rows
           for c in row.cells for par in c.paragraphs for r in par.runs)
print(f"Informação salva em: {OUT}")
