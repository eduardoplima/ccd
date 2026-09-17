"""Informações de verificação de transferência FRAP — 26 processos do rastreio.

Uma informação por processo da seção "Desconto em Folha — Verificar transferência
FRAP" do AGRUPAMENTO_PROCESSOS.md (8 já no marcador 5953 + 18 a marcar; rastreio:
scripts/analise/rastreio_verificar_frap.py). Cada documento traz:
  - "Trata-se de..." com processo, assunto, responsáveis e valor atualizado da multa;
  - conciliação: Quadro de descontos em folha (SIAI Pessoal, rubrica TCE/FRAP) e
    Quadro de repasses (extrato FRAP por CPF depositante ou baixas de parcela em
    FRAPDescontoFolhaParcela), com confronto dos totais; ou
  - parágrafo explicativo da ausência de registros, quando nada foi localizado.

Base legal: art. 25, § 1º, I (desconto com crédito ao FRAP) e §§ 3º e 4º
(comprovação do desconto e do crédito em 15 dias) da Resolução nº 013/2015-TCE.

Modelo: scripts/automacao/templates/modelo_informacao.docx; quadros no padrão
visual TCE/RN com legenda "Quadro N – título" (ref.: processos/001454_2023).

Rodar: .venv/Scripts/python.exe processos/verificacao_frap/gerar_informacoes.py
"""
import re
import shutil
from copy import deepcopy
from datetime import datetime
from pathlib import Path

import pandas as pd
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor
from docxtpl import DocxTemplate

from ccd.config import REPO_ROOT
from ccd.db import get_connection, run_query_df
from ccd.docs import docx_to_pdf

TEMPLATE = str(REPO_ROOT / "scripts" / "automacao" / "templates" / "modelo_informacao.docx")
DESTINO = Path(__file__).parent

# Desconto cessado há meses sem quitação do débito (deveriam estar quitados);
# os demais — inclusive planos-tranche do FRAPDescontoFolha concluídos, mas com
# dívida maior ainda em desconto — vão para curso/.
COMPLETO = {"005351/2017", "003272/2023"}

PROCESSOS = [  # 8 já com marcador 5953
    "700985/2012", "021694/2016", "005351/2017", "005352/2017",
    "003267/2023", "002167/2025", "003051/2025", "003667/2025",
    # 18 candidatos ao 5953 (rastreio de 18/08/2026)
    "006149/2006", "003272/2023", "002025/2024", "004871/2024", "004912/2024",
    "000868/2025", "001493/2025", "001494/2025", "002619/2025", "002739/2025",
    "003828/2025", "000068/2026", "000073/2026", "000076/2026", "000079/2026",
    "000469/2026", "000544/2026", "000625/2026",
]

MESES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho",
         "agosto", "setembro", "outubro", "novembro", "dezembro"]
hoje = datetime.now()
DATA = f"Natal/RN, {hoje.day} de {MESES[hoje.month - 1]} de {hoje.year}."

SQL_BASE = """
SELECT CONCAT(RTRIM(p.numero_processo), '/', RTRIM(p.ano_processo)) AS processo,
       RTRIM(p.assunto) AS assunto, RTRIM(r.nome) AS relator,
       d.IdDebito, processo.dbo.fn_Exe_RetornaValorAtualizado(d.IdDebito) AS valor,
       RTRIM(gp.Nome) AS responsavel, gp.Documento AS documento
FROM dbo.Processos p
LEFT JOIN dbo.Relator r ON r.codigo = p.codigo_relator
JOIN dbo.Exe_Debito d ON p.IdProcesso IN (d.IdProcessoExecucao, d.IdProcessoOrigem)
JOIN dbo.Exe_StatusDivida sd ON sd.CodigoStatusDivida = d.CodigoStatusDivida
JOIN dbo.Exe_DebitoPessoa dp ON dp.IDDebito = d.IdDebito
JOIN dbo.GenPessoa gp ON gp.IdPessoa = dp.IDPessoa
WHERE sd.StatusCancelamento IS NULL AND d.CodigoTipoDebito IN (2, 4, 5)
  AND CONCAT(RTRIM(p.numero_processo), '/', RTRIM(p.ano_processo)) IN ({procs})
"""

SQL_FOLHA = """
SELECT P.CPF, CAST(LTRIM(RTRIM(CCI.AnoReferencia)) AS INT) AS Ano,
       CAST(LTRIM(RTRIM(CCI.MesReferencia)) AS INT) AS Mes, CCI.Valor,
       MIN(R.Codigo) AS RubricaCodigo, MIN(LTRIM(RTRIM(R.Descricao))) AS Rubrica
FROM dbo.SiaiDp_ContraChequeItem CCI
JOIN dbo.SiaiDp_Rubrica R ON R.IdRubrica = CCI.IdRubrica
JOIN dbo.SiaiDp_ContraCheque CC ON CC.IdContraCheque = CCI.IdContraCheque
LEFT JOIN dbo.SiaiDp_Funcionario FU ON FU.IdFuncionario = CC.IdFuncionario
LEFT JOIN dbo.Comum_Pessoa P ON P.IdPessoa = FU.IdPessoa
WHERE P.CPF IN ({cpfs}) AND R.Tipo = '2'
  AND (UPPER(R.Descricao) LIKE '%TCE%' OR UPPER(R.Descricao) LIKE '%TRIBUNAL DE CONTAS%'
       OR UPPER(R.Descricao) LIKE '%FRAP%')
GROUP BY P.CPF, CAST(LTRIM(RTRIM(CCI.AnoReferencia)) AS INT),
         CAST(LTRIM(RTRIM(CCI.MesReferencia)) AS INT), CCI.Valor
"""

SQL_REPASSES = """
SELECT RIGHT(L.CpfCnpjDepositante, 11) AS CPF, L.DtMovimento, L.Documento, L.Valor
FROM dbo.FRAPLancamento L
JOIN dbo.FRAPCategoria cat ON cat.IdCategoria = L.IdCategoria
WHERE cat.Codigo IN ('OB_RECEBIDA', 'TRANSFERENCIA') AND L.ValorDC = 'C'
  AND RIGHT(L.CpfCnpjDepositante, 11) IN ({cpfs})
"""

# O controle FRAPDescontoFolha(Parcela) tem cadastros repetidos da mesma parcela
# (planos-tranche sobrepostos, até 5x a mesma competência/valor) — o GROUP BY
# deduplica por competência+valor antes de qualquer soma ou casamento.
SQL_BAIXAS = """
SELECT DF.CpfCnpj AS CPF, P.AnoReferencia AS Ano, P.MesReferencia AS Mes,
       P.ValorEsperado, MIN(P.DataPagamentoParcela) AS DataBaixa
FROM dbo.FRAPDescontoFolha DF
JOIN dbo.FRAPDescontoFolhaParcela P ON P.IdFRAPDescontoFolha = DF.IdFRAPDescontoFolha
WHERE DF.Ativo = 1 AND P.DataPagamentoParcela IS NOT NULL AND DF.CpfCnpj IN ({cpfs})
GROUP BY DF.CpfCnpj, P.AnoReferencia, P.MesReferencia, P.ValorEsperado
"""

FUNDAMENTO = "art. 25, §§ 3º e 4º, da Resolução nº 013/2015-TCE"


def brl(v: float) -> str:
    return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _cpf(doc: object) -> str | None:
    d = re.sub(r"\D", "", str(doc or ""))
    return d.zfill(11) if 0 < len(d) <= 11 else None


def _in_list(cpfs: set[str]) -> str:
    assert cpfs and all(re.fullmatch(r"\d{11}", c) for c in cpfs)
    return ", ".join(f"'{c}'" for c in sorted(cpfs))


def fmt_cpf(c: str) -> str:
    return f"{c[:3]}.{c[3:6]}.{c[6:9]}-{c[9:]}"


# --- identidade visual TCE/RN (skill tce-rn-identity; ref. processos/001454_2023) ---
TCE = {"verde_principal": "2E5B3C", "verde_escuro": "1A3D28",
       "cinza_claro": "F2F2F2", "cinza_borda": "CCCCCC", "cinza_texto": "333333"}


def _shade(cell, hex_cor: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_cor)
    cell._tc.get_or_add_tcPr().append(shd)


def _tabela_apos(d, ancora_texto: str, legenda: str, cabecalho: list[str],
                 linhas: list[list[str]]) -> None:
    """Insere legenda "Quadro N – título" + tabela após o parágrafo-âncora."""
    ancora = next(p for p in d.paragraphs if ancora_texto in p.text)
    cap = d.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_cap = cap.add_run(legenda)
    run_cap.bold = True
    run_cap.font.color.rgb = RGBColor.from_string(TCE["verde_escuro"])
    ancora._p.addnext(cap._p)
    t = d.add_table(rows=0, cols=len(cabecalho))
    borders = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"),
               TCE["cinza_borda"] if lado.startswith("inside") else TCE["verde_principal"])
        borders.append(el)
    t._tbl.tblPr.append(borders)
    for i, linha in enumerate([cabecalho] + linhas):
        total = linha[0].startswith("Total")
        cells = t.add_row().cells
        for cell, texto in zip(cells, linha, strict=True):
            if i == 0:
                _shade(cell, TCE["verde_principal"])
            elif total:
                _shade(cell, TCE["verde_escuro"])
            elif i % 2 == 1:
                _shade(cell, TCE["cinza_claro"])
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = par.add_run(texto)
            run.bold = i == 0 or total
            run.font.color.rgb = RGBColor.from_string(
                "FFFFFF" if (i == 0 or total) else TCE["cinza_texto"])
    cap._p.addnext(t._tbl)


def montar_paragrafos(item: dict) -> tuple[list[str], list[tuple]]:
    """Texto da informação + lista de quadros [(âncora, legenda, cabeçalho, linhas)]."""
    varios = len(item["responsaveis"]) > 1
    nomes = "; ".join(f"{n} (CPF nº {fmt_cpf(c)})" for n, c in item["responsaveis"])
    quem = ("dos responsáveis " if varios else "do responsável ") + nomes
    folha, rep, baixas = item["folha"], item["repasses"], item["baixas"]
    quadros: list[tuple] = []
    n_quadro = 0

    pars = [
        f'Trata-se do Processo nº {item["processo"]}-TC — {item["assunto"]} —, em que se '
        f"apura o adimplemento de multa em desfavor {quem}, no valor atualizado de "
        f'R$ {brl(item["valor"])}. Os autos encontram-se nesta Coordenadoria de Controle de '
        "Decisões para verificação da efetivação do desconto em folha de pagamento e da "
        "transferência dos respectivos valores à conta do Fundo de Reaparelhamento e "
        "Aperfeiçoamento (FRAP), nos termos do art. 25, § 1º, inciso I, da Resolução nº "
        "013/2015-TCE.",
    ]

    if not folha.empty or not baixas.empty:
        # conciliação SIAI × FRAP: casa cada competência descontada (SIAI Pessoal)
        # com a baixa de parcela do controle do FRAP, por responsável+competência+valor
        f = folha.assign(cents=(folha.Valor.astype(float) * 100).round().astype(int))
        b = baixas.assign(cents=(baixas.ValorEsperado.astype(float) * 100).round().astype(int))
        conc = f.merge(b[["CPF", "Ano", "Mes", "cents", "DataBaixa"]],
                       on=["CPF", "Ano", "Mes", "cents"], how="outer", indicator=True)
        conc = conc.rename(columns={"_merge": "origem"})
        conc["valor"] = conc.cents / 100
        conc = conc.sort_values(["CPF", "Ano", "Mes", "cents"])
        casadas = int((conc.origem == "both").sum())
        so_siai = int((conc.origem == "left_only").sum())
        so_frap = int((conc.origem == "right_only").sum())
        total_sem_baixa = round(float(conc.valor[conc.origem == "left_only"].sum()), 2)

        n_quadro += 1
        pars.append(
            "A verificação confrontou duas fontes: (i) os descontos lançados nas folhas de "
            "pagamento remetidas ao SIAI Pessoal, em rubrica de recolhimento em favor deste "
            "Tribunal; e (ii) as baixas de parcela registradas no controle de descontos em "
            "folha do FRAP, mantido pela DIP — desconsiderados os cadastros repetidos da "
            "mesma parcela. O casamento, por competência e valor, está discriminado no "
            "quadro a seguir:")
        cab = (["Responsável"] if varios else []) + [
            "Competência", "Valor (R$)", "Rubrica (SIAI Pessoal)",
            "Baixa no controle FRAP", "Situação"]
        linhas = []
        for r in conc.itertuples():
            if r.origem == "both":
                sit = "descontada e baixada"
            elif r.origem == "left_only":
                sit = "descontada sem baixa no FRAP"
            else:
                sit = "baixa no FRAP sem folha no SIAI"
            linhas.append(([item["nome_por_cpf"][r.CPF]] if varios else []) + [
                f"{MESES[int(r.Mes) - 1]}/{int(r.Ano)}", brl(float(r.valor)),
                f"{r.RubricaCodigo} – {r.Rubrica}" if pd.notna(r.RubricaCodigo) else "—",
                r.DataBaixa.strftime("%d/%m/%Y") if pd.notna(r.DataBaixa) else "—",
                sit])
        linha_total = [""] * len(cab)
        linha_total[0] = "Total"
        linha_total[2 if varios else 1] = brl(round(float(conc.valor.sum()), 2))
        linhas.append(linha_total)
        quadros.append(("discriminado no quadro a seguir",
                        f"Quadro {n_quadro} – Conciliação dos descontos em folha "
                        "(SIAI Pessoal) com as baixas do controle FRAP",
                        cab, linhas))

        sintese = (
            f'Em síntese: os descontos localizados no SIAI Pessoal somam R$ '
            f'{brl(item["total_folha"])} e as baixas registradas no controle do FRAP somam '
            f'R$ {brl(item["total_baixas"])}. Do casamento, {casadas} competência(s) têm '
            "desconto e baixa correspondentes")
        if so_siai:
            sintese += f"; {so_siai} foram descontadas sem baixa no controle do FRAP"
        if so_frap:
            sintese += (f"; {so_frap} têm baixa registrada sem desconto visível no SIAI "
                        "Pessoal (folha do órgão não remetida ao sistema no período)")
        pars.append(sintese + ".")

        if not rep.empty:
            n_quadro += 1
            pars.append(
                "Além disso, o extrato bancário da conta do FRAP registra os seguintes "
                "créditos vinculados ao CPF do depositante:")
            linhas2 = [[r.DtMovimento.strftime("%d/%m/%Y"), str(r.Documento or ""),
                        brl(float(r.Valor))] for r in rep.itertuples()]
            linhas2.append(["Total", "", brl(item["total_repasses"])])
            quadros.append(("créditos vinculados ao CPF do depositante",
                            f"Quadro {n_quadro} – Créditos no extrato bancário da conta "
                            "do FRAP",
                            ["Data do crédito", "Documento", "Valor (R$)"], linhas2))

        if total_sem_baixa > 0:
            pars.append(
                f"Constata-se, assim, o montante de R$ {brl(total_sem_baixa)} descontado em "
                "folha sem baixa correspondente no controle do FRAP, em desacordo com o "
                f"{FUNDAMENTO}, que impõe ao órgão responsável pela folha de pagamento a "
                "comprovação não apenas do lançamento das parcelas, mas também do crédito "
                "do valor na conta do FRAP.")
            pars.append(
                "Ante o exposto, sugere-se o encaminhamento dos autos ao Gabinete da "
                "Relatoria, com a proposta de notificação do órgão responsável pela folha "
                "de pagamento para que, no prazo de 15 (quinze) dias, comprove o repasse de "
                f"R$ {brl(total_sem_baixa)} à conta do FRAP, ou apresente as justificativas "
                "cabíveis.")
        elif item["processo"] in COMPLETO:
            # desconto interrompido sem quitação — apontar a cessação e propor retomada
            ult = conc.loc[(conc.Ano * 100 + conc.Mes).idxmax()]
            pars.append(
                "Não se identifica quantia descontada em folha sem a correspondente baixa "
                "no controle do FRAP. Contudo, o último desconto localizado refere-se à "
                f"competência {MESES[int(ult.Mes) - 1]}/{int(ult.Ano)}, e a dívida permanece "
                f'em aberto no valor atualizado de R$ {brl(item["valor"])} — ou seja, o '
                "desconto em folha foi interrompido sem a quitação do débito.")
            pars.append(
                "Ante o exposto, sugere-se o encaminhamento dos autos ao Gabinete da "
                "Relatoria, com a proposta de notificação do órgão responsável pela folha "
                "de pagamento para que, no prazo de 15 (quinze) dias, informe as razões da "
                "interrupção e comprove a retomada do desconto até a quitação integral da "
                "dívida, na forma do art. 25, § 1º, inciso I, e §§ 3º e 4º, da Resolução "
                "nº 013/2015-TCE.")
        else:
            pars.append(
                "Não se identifica, portanto, quantia descontada em folha sem a "
                "correspondente baixa no controle do FRAP. Ante o exposto, sugere-se o "
                "prosseguimento do acompanhamento dos descontos e repasses até a quitação "
                f"integral da dívida, na forma do {FUNDAMENTO}.")
    else:
        pars.append(
            "Em consulta à base de dados do SIAI Pessoal, não foram localizados lançamentos "
            "de desconto em folha de pagamento em favor deste Tribunal vinculados ao CPF "
            + ("dos responsáveis" if varios else "do responsável")
            + ", tampouco créditos correspondentes no extrato bancário da conta do FRAP ou "
            "parcelas com baixa registrada no sistema de acompanhamento dos descontos. "
            "Registre-se que a consulta alcança apenas as folhas de pagamento remetidas ao "
            "SIAI Pessoal, podendo haver descontos efetuados por órgão cuja folha não é "
            "carregada no sistema.")
        pars.append(
            "Ante o exposto, sugere-se o encaminhamento dos autos ao Gabinete da Relatoria, "
            "com a proposta de notificação do órgão responsável pela folha de pagamento para "
            "que, no prazo de 15 (quinze) dias, comprove a implantação do desconto e o "
            f"respectivo crédito à conta do FRAP, na forma do {FUNDAMENTO}, ou informe as "
            "razões da não implantação, observado o disposto no art. 27 da mesma Resolução.")

    if item["outros_processos"]:
        outros = ", ".join(f"nº {p}-TC" for p in item["outros_processos"])
        pars.append(
            "Registre-se, por oportuno, que o mesmo responsável figura em outro(s) processo(s) "
            f"em fase de verificação nesta Coordenadoria ({outros}), de modo que os "
            "lançamentos identificados pelo CPF podem abranger mais de uma obrigação, "
            "cabendo ao órgão responsável pela folha discriminar os valores por débito.")

    pars.append("É a informação.")
    return pars, quadros


def gerar(item: dict) -> Path:
    num, ano = item["processo"].split("/")
    nivel = "completo" if item["processo"] in COMPLETO else "curso"
    pasta = DESTINO / nivel / f"{num}_{ano}"
    pasta.mkdir(parents=True, exist_ok=True)
    out = pasta / f"informacao_{num}_{ano}.docx"

    pars, quadros = montar_paragrafos(item)
    doc = DocxTemplate(TEMPLATE)
    doc.render({
        "processo": f'{item["processo"]} - TC',
        "assunto": item["assunto"],
        "relator": item["relator"],
        "parágrafos": pars,
    })
    d = doc.docx  # pós-processamento (get_docx() descartaria o render)
    for ancora, legenda, cab, linhas in quadros:
        _tabela_apos(d, ancora, legenda, cab, linhas)

    # data antes do bloco de assinatura (o modelo não traz)
    assinado = next(p for p in d.paragraphs if "assinado digitalmente" in p.text)
    for texto in (DATA, ""):
        novo = deepcopy(assinado._p)
        for r in list(novo):
            if r.tag == qn("w:r"):
                novo.remove(r)
        assinado._p.addprevious(novo)
        if texto:
            r = deepcopy(assinado.runs[0]._r)
            novo.append(r)
            for t in r.findall(qn("w:t")):
                t.text = texto

    # preserva a versão anterior antes de sobrescrever (skill edicao-minima)
    if out.exists():
        bak = out.with_name(f"{out.stem}_{datetime.now():%Y%m%d_%H%M%S}{out.suffix}")
        shutil.copy2(out, bak)
        print(f"  versão anterior preservada em: {bak.name}")
    doc.save(str(out))
    docx_to_pdf(str(out), str(pasta))
    return out


def carregar() -> list[dict]:
    eng = get_connection("processo")
    procs = ", ".join(f"'{p}'" for p in PROCESSOS)
    base = run_query_df(SQL_BASE.format(procs=procs), eng)
    base["cpf"] = base.documento.map(_cpf)
    faltando = set(PROCESSOS) - set(base.processo)
    assert not faltando, f"sem débito de multa vigente/responsável: {faltando}"

    cpfs = set(base.cpf.dropna())
    engp = get_connection("BdSIAIPessoal")
    engd = get_connection("BdDIP")
    folha = run_query_df(SQL_FOLHA.format(cpfs=_in_list(cpfs)), engp)
    rep = run_query_df(SQL_REPASSES.format(cpfs=_in_list(cpfs)), engd)
    baixas = run_query_df(SQL_BAIXAS.format(cpfs=_in_list(cpfs)), engd)

    proc_por_cpf: dict[str, list[str]] = {}
    for r in base.dropna(subset=["cpf"]).itertuples():
        proc_por_cpf.setdefault(r.cpf, [])
        if r.processo not in proc_por_cpf[r.cpf]:
            proc_por_cpf[r.cpf].append(r.processo)

    itens = []
    for proc in PROCESSOS:
        g = base[base.processo == proc]
        resp = sorted({(r.responsavel, r.cpf) for r in g.dropna(subset=["cpf"]).itertuples()})
        meus = sorted(c for _, c in resp)
        f = folha[folha.CPF.isin(meus)].sort_values(["CPF", "Ano", "Mes"])
        rp = rep[rep.CPF.isin(meus)].sort_values("DtMovimento")
        bx = baixas[baixas.CPF.isin(meus)].sort_values(["Ano", "Mes"])
        outros = sorted({p for c in meus for p in proc_por_cpf[c] if p != proc})
        itens.append({
            "processo": proc,
            "assunto": g.assunto.iloc[0],
            "relator": g.relator.iloc[0] or "",
            "valor": round(float(g.drop_duplicates("IdDebito").valor.fillna(0).sum()), 2),
            "responsaveis": resp,
            "nome_por_cpf": {c: n for n, c in resp},
            "folha": f, "repasses": rp, "baixas": bx,
            "total_folha": round(float(f.Valor.sum()), 2) if not f.empty else 0.0,
            "total_repasses": round(float(rp.Valor.sum()), 2) if not rp.empty else 0.0,
            "total_baixas": round(float(bx.ValorEsperado.sum()), 2) if not bx.empty else 0.0,
            "outros_processos": outros,
        })
    return itens


if __name__ == "__main__":
    import docx

    resumo = []
    for item in carregar():
        out = gerar(item)
        tem_conc = not item["folha"].empty or not item["baixas"].empty
        caso = "conciliação" if tem_conc else "sem registros"
        resumo.append((item["processo"], caso))
        print(f'{item["processo"]}: {caso} | folha R$ {brl(item["total_folha"])} | '
              f'repasses R$ {brl(item["total_repasses"])} | baixas R$ {brl(item["total_baixas"])}')

        # check: sem placeholder solto, "Trata-se" presente e quadros coerentes com o caso
        d = docx.Document(str(out))
        texto = "\n".join([p.text for p in d.paragraphs]
                          + [c.text for t in d.tables for r in t.rows for c in r.cells])
        assert "{{" not in texto and "}}" not in texto, f'{item["processo"]}: placeholder solto'
        assert "Trata-se do Processo" in texto and "É a informação." in texto
        assert item["processo"] in texto and DATA in texto
        esperado = (1 if tem_conc else 0) + (0 if item["repasses"].empty else 1)
        assert len(d.tables) == esperado, f'{item["processo"]}: {len(d.tables)} quadros'
    print(f"\n{len(resumo)} informações em {DESTINO}")
    for caso in ("conciliação", "sem registros"):
        n = sum(1 for _, c in resumo if c == caso)
        print(f"  {caso}: {n}")
