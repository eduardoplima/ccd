"""Instrução da CCD no processo 001603/2025 (seletivo e prioritário).

Representação (remuneração dos vereadores de São Fernando, legislatura 2025-2028) julgada
pelo Acórdão nº 302/2026-TC – 2ª Câmara (sessão de 29/09/2026, Evento 99), ainda sem
trânsito em julgado: multa de R$ 6.447,92 ao ex-Presidente da Câmara (item "c") e
recomendações à Câmara (item "f"). Fase do art. 5º da Resolução nº 013/2015: registro
provisório da multa pela CCD e envio à DE para a intimação (padrão do Despacho da CCD no
600443/2019, Evento 243).

A peça afirma o registro provisório da multa: o script avisa se o débito ainda não
aparece no banco (que atrasa em relação ao e-Contas — conferir lá antes de assinar).

Rodar: .venv/Scripts/python.exe processos/seletivos/001603_2025/gerar_informacao.py
"""
import shutil
from copy import deepcopy
from datetime import datetime
from pathlib import Path

from docx.oxml.ns import qn
from docxtpl import DocxTemplate

from ccd.config import REPO_ROOT
from ccd.db import run_query_df
from ccd.docs import docx_to_pdf

PASTA = Path(__file__).resolve().parent
TEMPLATE = str(REPO_ROOT / "scripts" / "automacao" / "templates" / "modelo_informacao.docx")
OUT = PASTA / "informacao_001603_2025.docx"

MESES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho",
         "agosto", "setembro", "outubro", "novembro", "dezembro"]
hoje = datetime.now()
DATA = f"Natal/RN, {hoje.day} de {MESES[hoje.month - 1]} de {hoje.year}."

PARAGRAFOS = [
    ("Trata-se de Representação, de caráter seletivo e prioritário, acerca da remuneração "
     "dos agentes políticos do Município de São Fernando para a legislatura 2025-2028, "
     "julgada pela Segunda Câmara deste Tribunal em sessão de 29 de setembro de 2026, por "
     "meio do Acórdão nº 302/2026-TC (Evento 99), que, entre outras deliberações, declarou "
     "a irregularidade da conduta do Sr. Misael Bruno de Araujo Silva, Presidente da Câmara "
     "Municipal de São Fernando à época da aprovação da Lei Municipal nº 0909/2023, e lhe "
     "aplicou multa no valor de R$ 6.447,92 (seis mil, quatrocentos e quarenta e sete reais "
     "e noventa e dois centavos), nos termos do item “c” do dispositivo, além de expedir "
     "recomendações à Câmara Municipal de São Fernando (item “f”)."),
    ("Em cumprimento ao art. 5º, caput, da Resolução nº 013/2015-TCE, esta Coordenadoria "
     "procedeu ao registro provisório da referida multa nos sistemas deste Tribunal."),
    ("Ante o exposto, sugere-se o envio dos autos à Diretoria de Expediente (DE) para a "
     "intimação do responsável, acompanhada do voto, do inteiro teor do acórdão, da memória "
     "de cálculo da dívida provisoriamente registrada e da guia bancária para pagamento, "
     "nos termos do art. 47, parágrafo único, da Lei Complementar Estadual nº 464/2012 c/c "
     "o art. 5º, § 1º, da Resolução nº 013/2015-TCE, bem como para a cientificação da "
     "Câmara Municipal de São Fernando quanto às recomendações constantes do item “f” do "
     "Acórdão nº 302/2026-TC, observada a prioridade inerente ao caráter seletivo do feito."),
]


def carregar() -> dict:
    p = run_query_df("""
        SELECT p.IdProcesso, RTRIM(p.assunto) assunto, r.nome relator
        FROM Processos p LEFT JOIN Relator r ON r.codigo = p.codigo_relator
        WHERE p.numero_processo = :n AND p.ano_processo = :a""", n="001603", a="2025").iloc[0]
    debitos = run_query_df(
        "SELECT IdDebito FROM Exe_Debito WHERE IdProcessoOrigem = :id", id=int(p.IdProcesso))
    if debitos.empty:
        print("  ATENÇÃO: multa ainda não cadastrada no banco — registrar (ou conferir no "
              "e-Contas) antes de assinar, a peça afirma o registro provisório.")
    return {"assunto": p.assunto, "relator": p.relator}


def gerar(item: dict) -> Path:
    doc = DocxTemplate(TEMPLATE)
    doc.render({
        "processo": "001603/2025 - TC",
        "assunto": item["assunto"],
        "relator": item["relator"].title(),
        "parágrafos": PARAGRAFOS,  # sem numerar: o modelo numera
    })
    d = doc.docx  # pós-processamento (get_docx() descartaria o render)

    # data antes do bloco de assinatura (o modelo não traz); os parágrafos vazios antes
    # da assinatura saem para caber em uma página (mesmo bloco dos geradores de lote)
    assinado = next(p for p in d.paragraphs if "assinado digitalmente" in p.text)
    prev = assinado._p.getprevious()
    while prev is not None and not "".join(
            t.text or "" for t in prev.findall(f".//{qn('w:t')}")).strip():
        vazio, prev = prev, prev.getprevious()
        vazio.getparent().remove(vazio)
    novo = deepcopy(assinado._p)
    for r in list(novo):
        if r.tag == qn("w:r"):
            novo.remove(r)
    assinado._p.addprevious(novo)
    r = deepcopy(assinado.runs[0]._r)
    novo.append(r)
    for t in r.findall(qn("w:t")):
        t.text = DATA
    idx_data = next(i for i, p in enumerate(d.paragraphs) if p.text.startswith("Natal/RN"))
    for p in d.paragraphs[idx_data:-1]:
        p.paragraph_format.keep_with_next = True

    # preserva a versão anterior antes de sobrescrever (skill edicao-minima)
    if OUT.exists():
        bak = OUT.with_name(f"{OUT.stem}_{datetime.now():%Y%m%d_%H%M%S}{OUT.suffix}")
        shutil.copy2(OUT, bak)
        print(f"  versão anterior preservada em: {bak.name}")
    doc.save(str(OUT))
    docx_to_pdf(str(OUT), str(PASTA))
    return OUT


if __name__ == "__main__":
    import docx  # noqa: E402

    item = carregar()
    out = gerar(item)
    # check (ponytail): nada de placeholder solto e o texto foi para o documento
    texto = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    assert "{{" not in texto and "}}" not in texto, "placeholder não substituído"
    for trecho in ("001603/2025 - TC", item["assunto"], DATA, *PARAGRAFOS):
        assert trecho in texto, f"faltou \"{trecho[:40]}\""
    assert out.with_suffix(".pdf").is_file(), "PDF não gerado"
    print(f"  salvo: {out} (+ .pdf)")
