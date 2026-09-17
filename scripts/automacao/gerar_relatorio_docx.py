"""Converte um relatório em Markdown para `.docx` e `.pdf` no padrão da CCD.

O Markdown é a fonte única; o `.docx` herda o cabeçalho institucional, as margens e
os estilos de `templates/modelo_informacao.docx`, ganha numeração de páginas no
rodapé (que o modelo não tem) e sai com as tabelas na identidade visual do TCE/RN.
O PDF vem do `.docx`, para que as duas saídas sejam a mesma peça.

ponytail: o parser cobre um subconjunto deliberadamente pequeno de Markdown —
títulos `#`/`##`/`###`, parágrafo, lista `-`, tabela `| … |`, `**negrito**` e `---`.
Não é um conversor de Markdown; é o suficiente para os relatórios da CCD. Se um dia
precisar de imagem, nota de rodapé ou lista numerada aninhada, o caminho é adotar o
Pandoc com `--reference-doc`, não crescer isto aqui.

Uso:
    python -m scripts.automacao.gerar_relatorio_docx RELATORIO_AUDITORIA_FINANCEIRA_CCD.md
    python -m scripts.automacao.gerar_relatorio_docx arquivo.md --sem-pdf
"""
from __future__ import annotations

import argparse
import re
import sys
from datetime import datetime
from pathlib import Path

import docx
from docx.oxml.ns import qn

from ccd.config import REPO_ROOT
from ccd.docs import docx_to_pdf

TEMPLATE = REPO_ROOT / "scripts" / "automacao" / "templates" / "modelo_informacao.docx"
DESTINO = REPO_ROOT / "saidas" / "analise" / "relatorio_auditoria_financeira"

# Paleta institucional TCE/RN (skill tce-rn-identity).
TCE = {
    "verde_principal": "2E5B3C",
    "verde_escuro": "1A3D28",
    "branco": "FFFFFF",
    "cinza_texto": "333333",
    "cinza_claro": "F2F2F2",
    "cinza_borda": "CCCCCC",
}

LARGURA_UTIL = 8640  # twips disponíveis entre as margens do modelo
NUMERICO = re.compile(r"^[\(\)R\$\s\d.,%–\-]+$")


# ── parser do subconjunto de Markdown ────────────────────────────────────


def parse_markdown(texto: str) -> list[tuple]:
    """Devolve blocos ('h1'|'h2'|'h3'|'p'|'li'|'quadro'|'tabela'|'hr', conteúdo)."""
    blocos: list[tuple] = []
    linhas = texto.splitlines()
    i = 0
    while i < len(linhas):
        linha = linhas[i].rstrip()
        if not linha.strip():
            i += 1
        elif linha.startswith("|"):
            tabela, i = _le_tabela(linhas, i)
            blocos.append(("tabela", tabela))
        elif linha.startswith("### "):
            blocos.append(("h3", linha[4:].strip()))
            i += 1
        elif linha.startswith("## "):
            blocos.append(("h2", linha[3:].strip()))
            i += 1
        elif linha.startswith("# "):
            blocos.append(("h1", linha[2:].strip()))
            i += 1
        elif linha.strip() in {"---", "***", "___"}:
            blocos.append(("hr", ""))
            i += 1
        elif linha.startswith("- "):
            blocos.append(("li", linha[2:].strip()))
            i += 1
        else:
            # A legenda do quadro é o parágrafo em negrito que precede a tabela.
            tipo = "quadro" if linha.lstrip("*").startswith("Quadro ") else "p"
            blocos.append((tipo, linha.strip()))
            i += 1
    return blocos


def _le_tabela(linhas: list[str], i: int) -> tuple[list[list[str]], int]:
    linhas_tabela = []
    while i < len(linhas) and linhas[i].strip().startswith("|"):
        celulas = [c.strip() for c in linhas[i].strip().strip("|").split("|")]
        # A linha de separação (|---|---|) não é conteúdo.
        if not all(set(c) <= set("-: ") and c for c in celulas):
            linhas_tabela.append(celulas)
        i += 1
    largura = max(len(linha) for linha in linhas_tabela)
    return [linha + [""] * (largura - len(linha)) for linha in linhas_tabela], i


# ── construção do OOXML ──────────────────────────────────────────────────


def _run(texto: str, *, negrito: bool = False, cor: str, tamanho: int) -> object:
    r = docx.oxml.OxmlElement("w:r")
    rPr = docx.oxml.OxmlElement("w:rPr")
    if negrito:
        rPr.append(docx.oxml.OxmlElement("w:b"))
    el_cor = docx.oxml.OxmlElement("w:color")
    el_cor.set(qn("w:val"), cor)
    rPr.append(el_cor)
    sz = docx.oxml.OxmlElement("w:sz")
    sz.set(qn("w:val"), str(tamanho))
    rPr.append(sz)
    r.append(rPr)
    t = docx.oxml.OxmlElement("w:t")
    t.text = texto
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    return r


def _runs_com_negrito(texto: str, cor: str, tamanho: int, negrito_geral: bool = False) -> list:
    """Quebra `**negrito**` em runs distintos, preservando o resto do texto."""
    return [
        _run(parte, negrito=negrito_geral or bool(n % 2), cor=cor, tamanho=tamanho)
        for n, parte in enumerate(re.split(r"\*\*", texto))
        if parte
    ]


def make_p(texto: str, *, estilo: str = "p") -> object:
    """Um parágrafo do corpo. `estilo` controla espaçamento, cor e alinhamento."""
    formato = {
        "h1": dict(tamanho=32, negrito=True, cor=TCE["verde_escuro"], jc="center", antes=0, depois=360),
        "h2": dict(tamanho=26, negrito=True, cor=TCE["verde_principal"], jc="left", antes=360, depois=120),
        "h3": dict(tamanho=23, negrito=True, cor=TCE["verde_principal"], jc="left", antes=240, depois=100),
        "quadro": dict(tamanho=20, negrito=True, cor=TCE["verde_escuro"], jc="center", antes=240, depois=80),
        "li": dict(tamanho=22, negrito=False, cor=TCE["cinza_texto"], jc="both", antes=60, depois=60),
        "p": dict(tamanho=22, negrito=False, cor=TCE["cinza_texto"], jc="both", antes=120, depois=120),
    }[estilo]

    p = docx.oxml.OxmlElement("w:p")
    pPr = docx.oxml.OxmlElement("w:pPr")
    spacing = docx.oxml.OxmlElement("w:spacing")
    spacing.set(qn("w:before"), str(formato["antes"]))
    spacing.set(qn("w:after"), str(formato["depois"]))
    pPr.append(spacing)
    jc = docx.oxml.OxmlElement("w:jc")
    jc.set(qn("w:val"), formato["jc"])
    pPr.append(jc)
    if estilo == "li":
        ind = docx.oxml.OxmlElement("w:ind")
        ind.set(qn("w:left"), "567")
        ind.set(qn("w:hanging"), "227")
        pPr.append(ind)
    p.append(pPr)

    texto = f"•  {texto}" if estilo == "li" else texto
    for r in _runs_com_negrito(texto, formato["cor"], formato["tamanho"], formato["negrito"]):
        p.append(r)
    return p


def _larguras(linhas: list[list[str]]) -> list[int]:
    """Larguras proporcionais ao conteúdo, para a primeira coluna não sufocar."""
    pesos = [max(len(linha[c]) for linha in linhas) for c in range(len(linhas[0]))]
    pesos = [min(max(p, 6), 40) for p in pesos]
    total = sum(pesos)
    return [max(int(LARGURA_UTIL * p / total), 600) for p in pesos]


def make_table(linhas: list[list[str]]) -> object:
    larguras = _larguras(linhas)
    tbl = docx.oxml.OxmlElement("w:tbl")

    tblPr = docx.oxml.OxmlElement("w:tblPr")
    tblW = docx.oxml.OxmlElement("w:tblW")
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(sum(larguras)))
    tblPr.append(tblW)
    jc = docx.oxml.OxmlElement("w:jc")
    jc.set(qn("w:val"), "center")
    tblPr.append(jc)
    borders = docx.oxml.OxmlElement("w:tblBorders")
    for b in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = docx.oxml.OxmlElement(f"w:{b}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"),
               TCE["cinza_borda"] if b.startswith("inside") else TCE["verde_principal"])
        borders.append(el)
    tblPr.append(borders)
    tbl.append(tblPr)

    grid = docx.oxml.OxmlElement("w:tblGrid")
    for largura in larguras:
        gc = docx.oxml.OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(largura))
        grid.append(gc)
    tbl.append(grid)

    for ri, linha in enumerate(linhas):
        tr = docx.oxml.OxmlElement("w:tr")
        if ri == 0:
            # Cabeçalho repete no topo de cada página.
            trPr = docx.oxml.OxmlElement("w:trPr")
            trPr.append(docx.oxml.OxmlElement("w:tblHeader"))
            tr.append(trPr)
            fundo, cor = TCE["verde_principal"], TCE["branco"]
        elif linha[0].strip("* ").upper().startswith("TOTAL"):
            fundo, cor = TCE["verde_escuro"], TCE["branco"]
        elif ri % 2 == 1:
            fundo, cor = TCE["cinza_claro"], TCE["cinza_texto"]
        else:
            fundo, cor = TCE["branco"], TCE["cinza_texto"]

        for ci, celula in enumerate(linha):
            tc = docx.oxml.OxmlElement("w:tc")
            tcPr = docx.oxml.OxmlElement("w:tcPr")
            tcW = docx.oxml.OxmlElement("w:tcW")
            tcW.set(qn("w:type"), "dxa")
            tcW.set(qn("w:w"), str(larguras[ci]))
            tcPr.append(tcW)
            shd = docx.oxml.OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), fundo)
            tcPr.append(shd)
            tc.append(tcPr)

            p = docx.oxml.OxmlElement("w:p")
            pPr = docx.oxml.OxmlElement("w:pPr")
            spacing = docx.oxml.OxmlElement("w:spacing")
            spacing.set(qn("w:before"), "40")
            spacing.set(qn("w:after"), "40")
            pPr.append(spacing)
            alinhamento = docx.oxml.OxmlElement("w:jc")
            if ri == 0:
                alinhamento.set(qn("w:val"), "center")
            else:
                alinhamento.set(qn("w:val"),
                                "right" if NUMERICO.match(celula.replace("*", "")) else "left")
            pPr.append(alinhamento)
            p.append(pPr)
            negrito = ri == 0 or linha[0].strip("* ").upper().startswith("TOTAL")
            for r in _runs_com_negrito(celula, cor, 18, negrito):
                p.append(r)
            tc.append(p)
            tr.append(tc)
        tbl.append(tr)
    return tbl


def _campo(paragrafo, instrucao: str, cor: str) -> None:
    """Insere um campo do Word (PAGE, NUMPAGES) — python-docx não tem API para isso."""
    fld = docx.oxml.OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instrucao)
    fld.append(_run("1", cor=cor, tamanho=18))
    paragrafo._p.append(fld)


def numerar_paginas(doc) -> None:
    """O modelo vem com o rodapé vazio; aqui ele ganha 'Página X de Y'."""
    rodape = doc.sections[0].footer
    p = rodape.paragraphs[0] if rodape.paragraphs else rodape.add_paragraph()
    for filho in list(p._p):
        if filho.tag != qn("w:pPr"):
            p._p.remove(filho)
    pPr = p._p.get_or_add_pPr()
    jc = docx.oxml.OxmlElement("w:jc")
    jc.set(qn("w:val"), "center")
    pPr.append(jc)
    p._p.append(_run("Página ", cor=TCE["cinza_texto"], tamanho=18))
    _campo(p, "PAGE", TCE["cinza_texto"])
    p._p.append(_run(" de ", cor=TCE["cinza_texto"], tamanho=18))
    _campo(p, "NUMPAGES", TCE["cinza_texto"])


ASSINATURA = ("[assinado digitalmente]", "Eduardo Pereira Lima",
              "Coordenador de Controle de Decisões")


def montar(blocos: list[tuple]) -> object:
    """Abre o modelo, esvazia o corpo (preservando cabeçalho e seção) e escreve."""
    doc = docx.Document(str(TEMPLATE))
    corpo = doc.element.body
    for filho in list(corpo):
        if filho.tag in (qn("w:p"), qn("w:tbl")):
            corpo.remove(filho)

    elementos = []
    for tipo, conteudo in blocos:
        if tipo == "tabela":
            elementos.append(make_table(conteudo))
        elif tipo == "hr":
            continue  # a régua horizontal do Markdown só separa seções na leitura
        else:
            elementos.append(make_p(conteudo, estilo=tipo))

    elementos.append(make_p(""))
    elementos += [make_p(linha, estilo="quadro") for linha in ASSINATURA]

    # Insere antes do w:sectPr, que precisa continuar sendo o último filho do corpo.
    sect = corpo.find(qn("w:sectPr"))
    for el in elementos:
        (corpo.append(el) if sect is None else sect.addprevious(el))

    numerar_paginas(doc)
    return doc


def _asserts(doc) -> None:
    """Nada de placeholder do modelo pode sobreviver no documento final."""
    texto = "\n".join(p.text for p in doc.paragraphs)
    for marca in ("{{", "{%", "INFORMAÇÃO INSTRUTIVA"):
        assert marca not in texto, f"resíduo do modelo no documento final: {marca!r}"
    assert doc.tables or "Quadro" in texto, "o relatório perdeu os quadros"
    rodape = doc.sections[0].footer._element.xml
    assert "PAGE" in rodape and "NUMPAGES" in rodape, "rodapé sem numeração de páginas"


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("markdown", type=Path, help="arquivo .md de origem")
    ap.add_argument("--sem-pdf", action="store_true", help="gera só o .docx")
    args = ap.parse_args()

    origem = args.markdown if args.markdown.is_absolute() else REPO_ROOT / args.markdown
    blocos = parse_markdown(origem.read_text(encoding="utf-8"))
    doc = montar(blocos)
    _asserts(doc)

    DESTINO.mkdir(parents=True, exist_ok=True)
    saida = DESTINO / f"{origem.stem}_{datetime.now():%Y%m%d_%H%M%S}.docx"
    doc.save(str(saida))
    print(f"docx: {saida}")

    if not args.sem_pdf:
        print(f"pdf:  {docx_to_pdf(saida, DESTINO)}")


if __name__ == "__main__":
    sys.exit(main())
